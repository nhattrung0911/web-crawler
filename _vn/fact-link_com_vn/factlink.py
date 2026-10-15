"""
fact-link.com.vn - Vietnam Manufacturing Business Directory (httpx + selectolax, khong browser).

RECON (2026-10-01)
------------------
- PHP SSR co dien, utf-8. KHONG co JSON API / AJAX; sitemap.xml tra ve HTML trang chu (vo dung).
  robots.txt: User-agent * khong chan gi. Server CHAM ~3-5 s/request -> throughput = concurrency.
- !! Cookie `vlang` (ngon ngu TEN trong danh sach: session moi = jp) chi duoc set boi index.php?lang=vn, param lang=vn
  o URL search_* KHONG doi duoc -> moi client "warm" bang index.php?lang=vn, neu khong ten cong ty chi-co-ten
  lech theo nguon (jp/en/vn) va key "N:<ten>" bi nhan doi.
- Quirk session: request dau toi search_post.php / mem_profile.php co the tra body ~100 byte
  `<meta http-equiv="refresh" content="0;URL = ...">` + set PHPSESSID -> giu cookie, theo meta-refresh (toi da 2 hop).
- Trang chu `index.php?lang=vn` co banner so thanh vien (`li.flvn-mem-toltal`, ban en: "registered firms N").
  Cung chua link nguon liet ke: search_category.php?id=NNN (danh muc), search_post.php?id=N&type=prov|ctry|ie.
- LIET KE cong ty (30 dong/trang), phai UNION 4 nguon (khong nguon nao phu het):
    cat : search_category.php?id=NNN&start=N&lang=vn   (GET thuong, khong state)
    prov/ctry/ie : GET search_post.php?id=N&type=T (set search vao PHP session, tra meta-refresh toi
                   search_result.php?...start=0) roi search_result.php?id=N&type=T&start=N&lang=vn.
                   State theo PHPSESSID -> moi nguon 1 httpx.AsyncClient (cookie jar) rieng.
  Phan trang: link start=N day du; last = max(start).
- !! PHP SESSION LOCK: cac request cung PHPSESSID bi server XEP HANG (8 GET song song 1 cookie jar = 48 s, 8 jar
  rieng = 8 s). Moi request ~5 s TTFB co dinh (file tinh 0.05 s). => concurrency CHI co tac dung khi co N
  httpx.AsyncClient doc lap (N cookie jar): Net giu pool N client, moi request/cong ty/nguon muon 1 client.
  Hai layout dong khac nhau (category: co td co quoc tich + link home/<slug>; search_result: link mem_profile),
  chung `span.default_head` + 3 icon ngon ngu tpl_{jp|vn|en}_01.png (_00 = khong co ban ngon ngu do).
- Dinh danh: MID 8 so (hoi vien tra phi co link). Hoi vien chi-co-ten khong co MID -> key "N:<ten chuan hoa>".
- HO SO: mem_profile.php?id=MID&page=PAGE&lang=vn|en|jp. Gioi thieu = td dau tien trong div.panelcontent;
  bang field (hang 3 td, label o td dau, mau chu label KHAC nhau theo template) ; dia chi nam trong comment
  <!-- memaddN -->; trang con = link mem_home/mem_content.php?id=MID (loai banner cua MID khac, mem_inquiry).
- SAN PHAM: products_list.php?start=0 chi la trang tom tat theo nhom; liet ke day du qua
  products_list.php?id=CAT&start=N ; chi tiet product_detail.php?proid=N.
Output: out_factlink/ (factlink_state.db = hang doi + resume, mo_ta/*.md, factlink.xlsx, factlink.log).
"""
from __future__ import annotations

import argparse
import asyncio
import contextlib
import html as htmllib
import json
import logging
import os
import pathlib
import queue
import re
import sqlite3
import sys
import threading
import time
from urllib.parse import urljoin

import httpx
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font
from selectolax.lexbor import LexborHTMLParser

HERE = pathlib.Path(__file__).parent
BASE = "https://www.fact-link.com.vn"
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")
HEADERS = {"User-Agent": UA, "Accept-Language": "vi-VN,vi;q=0.9,en;q=0.8"}

PAGE_SIZE = 30
MAX_REDIRECTS = 5
MAX_HOPS = 2             # meta-refresh toi da
CIRCUIT_BREAK = 25       # loi LIEN TIEP -> dung
MAX_ITEMS = 20000        # tran so cong ty
MAX_PAGES = 600          # tran phan trang / nguon
MAX_CONC = 16            # tran so session song song
MAX_SUBS = 30            # tran trang con / cong ty
TIMEOUT = 60.0
RETRIES = 5
PAUSE_SECS = 60

log = logging.getLogger("factlink")

# --------------------------------------------------------------------------- cot Excel
COMPANY_COLS = [
    "MID", "Page ID", "Ten (list)", "Status", "Quoc tich", "Khu vuc", "Danh muc", "Tinh/TP", "KCN",
    "Mo ta ngan", "Ten cong ty", "Nguoi dai dien", "Mo ta kinh doanh", "San pham", "Dia chi", "TEL", "FAX",
    "Nguoi phu trach", "Email", "Website", "Ngay thanh lap", "Von dieu le", "Cong ty me", "Co dong chinh",
    "Nhan vien", "Quyet toan", "Ngan hang", "ISO", "Khach hang lon", "Cu diem", "Dat dai", "Award",
    "Thong so", "Cap nhat", "Ngon ngu", "So trang con", "File mo ta", "Anh", "URL ho so", "URL FL",
]
PRODUCT_COLS = ["ProID", "Ten", "Danh muc", "SKU", "Ton kho VN", "MID", "Cong ty", "File mo ta", "Anh", "URL"]
WRAP_COLS = {"Mo ta ngan", "Mo ta kinh doanh", "San pham", "Dia chi", "Cong ty me", "Khach hang lon",
             "Cu diem", "Dat dai", "Award", "Thong so", "Nguoi phu trach", "Danh muc", "KCN", "Ten cong ty"}

# label (casefold) VN + EN -> cot
LABELS = {
    "tên công ty": "Ten cong ty", "company name": "Ten cong ty",
    "tên người đại diện": "Nguoi dai dien", "representative name": "Nguoi dai dien",
    "mô tả kinh doanh": "Mo ta kinh doanh", "business description": "Mo ta kinh doanh",
    "sản phẩm": "San pham", "produce": "San pham", "product": "San pham", "products": "San pham",
    "tel": "TEL", "e-mail": "Email", "email": "Email", "số điện thoại": "TEL", "phone": "TEL",
    "fax": "FAX", "số fax": "FAX",
    "người phụ trách": "Nguoi phu trach", "contact person": "Nguoi phu trach",
    "trang web": "Website", "website": "Website",
    "ngày thành lập": "Ngay thanh lap", "establish date": "Ngay thanh lap",
    "vốn điều lệ": "Von dieu le", "capital": "Von dieu le",
    "công ty mẹ": "Cong ty me", "parent company": "Cong ty me",
    "cổ đông chính": "Co dong chinh", "shareholder": "Co dong chinh",
    "nhân viên": "Nhan vien", "employee": "Nhan vien",
    "quyết toán hàng năm": "Quyet toan", "account period": "Quyet toan",
    "ngân hàng giao dịch chính": "Ngan hang", "bank": "Ngan hang",
    "chứng nhận / iso": "ISO", "certification / iso": "ISO",
    "khách hàng lớn": "Khach hang lon", "value customer": "Khach hang lon",
    "cứ điểm": "Cu diem", "overseas bases": "Cu diem",
    "đất đai, nhà": "Dat dai", "land and factory": "Dat dai",
    "award": "Award",
    "会社名": "Ten cong ty", "代表者": "Nguoi dai dien", "事業内容": "Mo ta kinh doanh", "取扱品目": "San pham",
    "担当者": "Nguoi phu trach", "ウェブサイト": "Website", "創立日": "Ngay thanh lap", "資本金": "Von dieu le",
    "親会社": "Cong ty me", "主要株主": "Co dong chinh", "従業員数": "Nhan vien", "決算期": "Quyet toan",
    "取引銀行": "Ngan hang", "規格認証 / iso": "ISO", "主要取引先": "Khach hang lon", "拠点": "Cu diem",
    "土地、建物": "Dat dai", "表彰": "Award",
}
LABEL_SKIP = {"liên hệ", "contact", "問い合わせ"}
LABEL_ADDR = {"địa chỉ", "address"}
TEL_KEY = re.compile(r"(?<![a-z])tel|điện thoại|phone|電話")
FAX_KEY = re.compile(r"fax|ファックス")
MAIL_KEY = re.compile(r"e-?mail|メール")

NLRE = re.compile(r"\n+")
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
SUB_RE = re.compile(r"mem_(?:home|content)\.php\?id=(\d+)&(?:amp;)?page=(\d+)", re.I)


# =========================================================================== PURE PARSERS
def clean(s) -> str:
    return re.sub(r"\s+", " ", htmllib.unescape(str(s or "")).replace("\xa0", " ")).strip()


def norm_name(s: str) -> str:
    return clean(s).casefold()


def safe_filename(s: str, maxlen: int = 80) -> str:
    """Bo ky tu cam cua Windows, gon khoang trang, cat do dai."""
    s = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", str(s or ""))
    s = re.sub(r"\s+", " ", s).strip(" .")
    return s[:maxlen].strip(" .") or "untitled"


def html_to_text(node) -> str:
    """Node selectolax (hoac chuoi HTML) -> text doc duoc: <br>/<p>/<li> thanh xuong dong, bo &nbsp;."""
    h = node if isinstance(node, str) else (node.html or "")
    h = re.sub(r"(?is)<(script|style)\b.*?</\1>", "", h)
    h = re.sub(r"(?i)<br\s*/?>", "\n", h)
    h = re.sub(r"(?i)<li\b[^>]*>", "\n", h)
    h = re.sub(r"(?i)</?(p|div|li|ul|ol|tr|table|h[1-6]|blockquote)\b[^>]*>", "\n", h)
    h = re.sub(r"<[^>]+>", "", h)
    h = htmllib.unescape(h).replace("\xa0", " ")
    lines = [re.sub(r"[ \t\r\f\v]+", " ", ln).strip() for ln in h.split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def parse_home(html: str) -> tuple[list[tuple[str, str, str]], int]:
    """Trang chu -> ([(type, id, ten)], so thanh vien banner). type: cat|prov|ctry|ie."""
    srcs, seen = [], set()
    tree = LexborHTMLParser(html)
    for a in tree.css("a[href]"):
        href = a.attributes.get("href") or ""
        m = re.search(r"search_category\.php\?id=(\d+)", href)
        if m:
            t, i = "cat", m.group(1)
        else:
            m = re.search(r"search_post\.php\?id=(\d+)&(?:amp;)?type=(prov|ctry|ie)", href)
            if not m:
                continue
            t, i = m.group(2), m.group(1)
        if (t, i) in seen:
            continue
        seen.add((t, i))
        name = re.sub(r"\s*\(\d+\)\s*$", "", clean(a.text()))
        srcs.append((t, i, name))
    m = (re.search(r'flvn-mem-toltal">\s*([\d,.]+)', html)
         or re.search(r"registered firms\D{0,40}?([\d][\d,.]*)", html, re.I))
    banner = int(re.sub(r"\D", "", m.group(1))) if m else 0
    return srcs, banner


def parse_listing(html: str) -> tuple[list[dict], int]:
    """Trang liet ke (category HOAC search_result) -> (rows, max_start).
    Moi dong: key, mid, page, name, nat, region, short, langs, slug. Khoa theo span.default_head."""
    rows = []
    marks = [m.start() for m in re.finditer(r'<span class="default_head">', html)]
    prev = 0
    for pos in marks:
        before = html[prev:pos]
        prev = pos
        nat = re.findall(r'bg_nation_\w+\.\w+"\s+title="([^"]*)"', before)
        chunk = html[pos:]
        end = chunk.find("</table>")
        chunk = chunk[:end] if end > 0 else chunk
        m = re.match(r'<span class="default_head">(.*?)</span>', chunk, re.S)
        if not m:
            continue
        inner = m.group(1)
        name = clean(re.sub(r"<[^>]+>", "", inner))
        slug = ""
        hm = re.search(r'href="home/([^"/]+)', inner)
        if hm:
            slug = hm.group(1)
        rest = chunk[m.end():]
        desc = re.split(r'<a href="mem_profile|<img src="images/tpl_|</td>', rest, maxsplit=1)[0]
        desc = clean(re.sub(r"<[^>]+>", " ", desc))
        pm = re.search(r"mem_profile\.php\?id=(\d+)&(?:amp;)?page=(\d+)", chunk)
        mid, page = (pm.group(1), pm.group(2)) if pm else ("", "")
        langs = ",".join(l for l in ("jp", "vn", "en") if f"tpl_{l}_01" in chunk)
        region = ",".join(re.findall(r"images/sector/([a-z]+)_(?:en|vn)\.png", chunk))
        rows.append({
            "key": mid or "N:" + norm_name(name), "mid": mid, "page": page, "name": name,
            "nat": nat[-1] if nat else "", "region": region, "short": desc, "langs": langs, "slug": slug,
        })
    starts = [int(x) for x in re.findall(
        r"(?:search_category|search_result|products_list)\.php\?[^\"' >]*?start=(\d+)", html)]
    return rows, max(starts, default=0)


def _label(td) -> str:
    return clean(html_to_text(td).replace("\n", " "))


def parse_profile(html: str, mid: str) -> dict:
    """Trang ho so -> dict(ok, name, short, fields, addresses, other, emails, intro, subpages, updated, image, fl_url)."""
    # comment <!-- memaddN --> bao hieu hang dia chi -> doi thanh thuoc tinh de selectolax thay duoc
    html = re.sub(r"<!--\s*memadd\d+\s*-->\s*<tr", '<tr data-addr="1"', html)
    tree = LexborHTMLParser(html)
    out = {"ok": False, "name": "", "short": "", "fields": {}, "addresses": [], "other": [], "emails": [],
           "intro": "", "subpages": [], "updated": "", "image": "", "fl_url": ""}
    h1 = tree.css_first("h1.h1_title")
    out["name"] = clean(h1.text()) if h1 else ""
    sh = tree.css_first("td.white_n")
    out["short"] = clean(sh.text()) if sh else ""
    panel = tree.css_first("div.panelcontent")
    if panel is None:
        return out
    first = panel.css_first("td")
    out["intro"] = html_to_text(first) if first is not None else ""
    values = []
    for tr in panel.css("tr"):
        tds = [c for c in tr.iter() if c.tag == "td"]
        if len(tds) != 3 or tds[0].css_first("strong") is None:
            continue
        p, in_form = tr, False
        while p is not None:
            if p.tag == "form":
                in_form = True
                break
            p = p.parent
        if in_form:
            continue
        label = _label(tds[0]).strip(" :：　")      # bo ":" / "：" cuoi, chi lam 1 lan
        value = html_to_text(tds[2])
        key = label.casefold()
        if not label or key in LABEL_SKIP or not value:
            continue
        values.append(value)
        flags = (bool(TEL_KEY.search(key)), bool(FAX_KEY.search(key)), bool(MAIL_KEY.search(key)))

        def add(col, exact=False):
            line = value if exact else f"{label}: {value}"
            out["fields"][col] = out["fields"][col] + "\n" + line if col in out["fields"] else line

        if tr.attributes.get("data-addr") is not None:
            out["addresses"].append(NLRE.sub("\n", f"{label}: {value}"))
        elif key in LABEL_ADDR:
            out["addresses"].append(NLRE.sub("\n", value))
        elif "fact-link" in key or "ファクトリンク" in key:
            out["fl_url"] = clean(value)
        elif key in LABELS:                                   # khop CHINH XAC truoc
            col = LABELS[key]
            if col == "Website" and "fact-link.com.vn" in value:
                out["fl_url"] = clean(value)
            elif col == "Email":
                out["emails"].extend(EMAIL_RE.findall(value))
            else:
                add(col, exact=True)
        elif sum(flags) > 1:                                  # "TEL / FAX / E-mail": giu nguyen text, lay email
            out["emails"].extend(EMAIL_RE.findall(value))
            out["other"].append(f"{label}: {value}")
        elif flags[2]:
            found = EMAIL_RE.findall(value)
            out["emails"].extend(found)
            if not found:
                out["other"].append(f"{label}: {value}")
        elif flags[0]:
            add("TEL")
        elif flags[1]:
            add("FAX")
        else:
            out["other"].append(f"{label}: {value}")
    seen = set()
    for e in EMAIL_RE.findall("\n".join(values)):
        out["emails"].append(e)
    out["emails"] = [e.lower() for e in out["emails"] if not (e.lower() in seen or seen.add(e.lower()))]
    m = re.search(r"LATEST UPDATE\s*</strong>\s*([^<]+)", html)
    out["updated"] = clean(m.group(1)) if m else ""
    m = re.search(rf'home/[^"\'\s]+/{mid}-\d+-P\.jpg[^"\'\s]*', html) if mid else None
    out["image"] = urljoin(BASE + "/", m.group(0)) if m else ""
    seen_sub = set()
    for a in tree.css("a[href]"):
        href = a.attributes.get("href") or ""
        sm = SUB_RE.search(href)
        if not sm or sm.group(1) != mid or href in seen_sub:
            continue
        seen_sub.add(href)
        out["subpages"].append((clean(a.text()) or "Trang con", urljoin(BASE + "/", href)))
    out["ok"] = bool(h1 or out["fields"])
    return out


def parse_subpage(html: str) -> tuple[str, str]:
    """Trang mem_home/mem_content -> (tieu de h2, text div.panelcontent). Rong neu khong phai trang con."""
    tree = LexborHTMLParser(html)
    panel = tree.css_first("div.panelcontent")
    if panel is None:
        return "", ""
    h2 = panel.css_first("h2")
    return (clean(h2.text()) if h2 else ""), html_to_text(panel)


def parse_product(html: str) -> dict:
    """product_detail.php -> dict(title, category, sku, desc, stock, mid, company, image)."""
    tree = LexborHTMLParser(html)

    def txt(sel):
        n = tree.css_first(sel)
        return clean(n.text()) if n is not None else ""

    crumbs = [clean(a.text()) for a in tree.css("nav.breadcrumb .breadcrumb-item")]
    crumbs = [c for c in crumbs[1:] if c]
    sku = re.sub(r"^SKU\s*:\s*#?", "", txt("p.product-detail-sku"))
    stock = re.sub(r"^.*?:\s*", "", txt("p.product-detail-sta"))
    # VN: <p class="product-detail-des"><p>..</p></p> (p long p) -> lexbor dong p ngoai ngay => cat tho theo HTML
    m = re.search(r'<p class="product-detail-des">(.*?)<p class="product-detail-sta', html, re.S)
    des = m.group(1) if m else (tree.css_first("p.product-detail-des"))
    out = {"title": txt("h2.product-detail-tit"), "category": " > ".join(crumbs), "sku": sku,
           "desc": html_to_text(des) if des is not None else "", "stock": stock,
           "mid": "", "company": "", "image": ""}
    for a in tree.css("dl a[href]"):
        m = re.search(r"mem_profile\.php\?id=(\d+)", a.attributes.get("href") or "")
        if m:
            out["mid"], out["company"] = m.group(1), clean(a.text())
            break
    img = tree.css_first("#show-img")
    if img is not None:
        out["image"] = urljoin(BASE + "/", img.attributes.get("src") or "")
    return out


def parse_prod_ids(html: str) -> list[str]:
    seen: dict[str, None] = {}
    for x in re.findall(r"product_detail\.php\?proid=(\d+)", html):
        seen[x] = None
    return list(seen)


# =========================================================================== DB
class DB:
    def __init__(self, path: pathlib.Path):
        self.c = sqlite3.connect(path)
        self.c.execute("PRAGMA journal_mode=WAL")
        self.c.executescript("""
        CREATE TABLE IF NOT EXISTS meta(k TEXT PRIMARY KEY, v TEXT);
        CREATE TABLE IF NOT EXISTS src(type TEXT, id TEXT, name TEXT, rows INTEGER DEFAULT 0,
                                       done INTEGER DEFAULT 0, PRIMARY KEY(type,id));
        CREATE TABLE IF NOT EXISTS co(key TEXT PRIMARY KEY, mid TEXT, page TEXT, name TEXT, nat TEXT,
                                      region TEXT, short TEXT, langs TEXT, slug TEXT);
        CREATE TABLE IF NOT EXISTS mem(key TEXT, stype TEXT, sname TEXT, UNIQUE(key,stype,sname));
        CREATE TABLE IF NOT EXISTS det(key TEXT PRIMARY KEY, status TEXT, lang TEXT, fields TEXT,
                                       subs INTEGER, file TEXT);
        CREATE TABLE IF NOT EXISTS prod(proid TEXT PRIMARY KEY, status TEXT, data TEXT);
        """)
        self.c.commit()

    def meta(self, k, v=None):
        if v is None:
            r = self.c.execute("SELECT v FROM meta WHERE k=?", (k,)).fetchone()
            return r[0] if r else ""
        self.c.execute("INSERT OR REPLACE INTO meta VALUES(?,?)", (k, str(v)))
        self.c.commit()

    def add_rows(self, rows: list[dict], stype: str, sname: str) -> None:
        for r in rows:
            old = self.c.execute("SELECT langs,slug,region,nat,short,page FROM co WHERE key=?", (r["key"],)).fetchone()
            if old is None:
                self.c.execute("INSERT INTO co VALUES(?,?,?,?,?,?,?,?,?)",
                               (r["key"], r["mid"], r["page"], r["name"], r["nat"], r["region"],
                                r["short"], r["langs"], r["slug"]))
            else:   # gop thong tin con thieu
                langs, slug, region, nat, short, page = old
                if len(r["langs"]) > len(langs):
                    langs = r["langs"]
                self.c.execute("UPDATE co SET langs=?,slug=?,region=?,nat=?,short=?,page=? WHERE key=?",
                               (langs, slug or r["slug"], region or r["region"], nat or r["nat"],
                                short or r["short"], page or r["page"], r["key"]))
            self.c.execute("INSERT OR IGNORE INTO mem VALUES(?,?,?)", (r["key"], stype, sname))
        self.c.commit()

    def n_co(self) -> int:
        return self.c.execute("SELECT COUNT(*) FROM co").fetchone()[0]


def count_report(db: DB) -> list[list]:
    """Bang dem: moi loai nguon + hop nhat + so banner."""
    out = [["Nguon", "So nguon", "Nguon da xong", "Dong (tho)", "Key duy nhat", "Co MID", "Chi co ten"]]
    names = {"cat": "Danh muc", "prov": "Tinh/TP", "ctry": "Quoc gia", "ie": "KCN"}
    for t, label in names.items():
        n, done, rows = db.c.execute(
            "SELECT COUNT(*), COALESCE(SUM(done),0), COALESCE(SUM(rows),0) FROM src WHERE type=?", (t,)).fetchone()
        u, w = db.c.execute(
            "SELECT COUNT(DISTINCT key), COUNT(DISTINCT CASE WHEN key NOT LIKE 'N:%' THEN key END) "
            "FROM mem WHERE stype=?", (t,)).fetchone()
        out.append([label, n, done, rows, u, w, u - w])
    u = db.n_co()
    w = db.c.execute("SELECT COUNT(*) FROM co WHERE mid<>''").fetchone()[0]
    out.append(["HOP NHAT (union)", "", "", "", u, w, u - w])
    banner = int(db.meta("banner") or 0)
    out.append(["Banner website", "", "", "", banner, "", ""])
    out.append(["Chenh lech (union - banner)", "", "", "", u - banner, "", ""])
    return out


# =========================================================================== NETWORK
class FetchError(Exception):
    pass


class Net:
    """Pool N httpx.AsyncClient (moi client 1 cookie jar = 1 PHPSESSID rieng) + retry/backoff + tam dung
    khi 403/429/503 + circuit breaker + meta-refresh. PHP session lock: cac request cung PHPSESSID bi server
    xep HANG -> KHONG bao gio ban 2 request song song qua cung 1 client; moi request/cong ty/nguon muon 1 client.
    stop/pause la threading.Event de GUI bam nut tu thread khac."""

    def __init__(self, conc: int, delay: float, stop: threading.Event | None = None,
                 pause: threading.Event | None = None):
        self.conc = conc
        self.delay = delay
        self.stop = stop or threading.Event()
        self.pause = pause or threading.Event()
        self.pool: asyncio.Queue = asyncio.Queue()
        for _ in range(conc):
            self.pool.put_nowait(self.client())
        self.warmed: set[int] = set()
        self.fails = 0
        self.n_req = 0
        self.n_err = 0
        self.paused_until = 0.0
        self.total = self.done = self.ok = self.bad = self.named = 0   # thong ke cho GUI

    def client(self) -> httpx.AsyncClient:
        return httpx.AsyncClient(headers=HEADERS, timeout=TIMEOUT, follow_redirects=True,
                                 max_redirects=MAX_REDIRECTS, verify=True)

    async def aclose(self):
        while not self.pool.empty():
            await self.pool.get_nowait().aclose()

    @contextlib.asynccontextmanager
    async def session(self):
        """Muon 1 client (cookie jar rieng) cho tron 1 cong ty / 1 nguon co state."""
        cl = await self.pool.get()
        try:
            if id(cl) not in self.warmed:
                # cookie `vlang` quyet dinh ngon ngu TEN hien trong danh sach (session moi = jp, param lang=vn
                # o URL khong doi duoc) -> moi client phai GET index.php?lang=vn 1 lan truoc (cung set PHPSESSID)
                await self.get(f"{BASE}/index.php?lang=vn", cl)
                self.warmed.add(id(cl))
            yield cl
        finally:
            self.pool.put_nowait(cl)

    async def _gate(self):
        """Cho khi Pause / tam dung WAF; thoat som neu Stop."""
        while (self.pause.is_set() or time.time() < self.paused_until) and not self.stop.is_set():
            await asyncio.sleep(0.2)
        if self.stop.is_set():
            raise FetchError("stop")

    async def get(self, url: str, cl: httpx.AsyncClient | None = None, hops: int = 0) -> str:
        if cl is None:
            async with self.session() as c:
                return await self.get(url, c, hops)
        last = ""
        for i in range(RETRIES):
            await self._gate()
            try:
                await asyncio.sleep(self.delay)
                await self._gate()
                self.n_req += 1
                req = asyncio.ensure_future(cl.get(url))
                while not req.done():                      # huy request ngay khi Stop (khong cho timeout 60s)
                    await asyncio.wait({req}, timeout=0.3)
                    if self.stop.is_set():
                        req.cancel()
                        raise FetchError("stop")
                r = req.result()
                code = r.status_code
                if code in (403, 429, 503):
                    self.paused_until = time.time() + PAUSE_SECS
                    log.warning("[WARN] HTTP %s -> tam dung %ss (nghi WAF/rate-limit)", code, PAUSE_SECS)
                if code == 404:
                    raise FetchError("404")
                if code in (403, 429) or code >= 500:
                    last = f"HTTP {code}"
                    raise httpx.TransportError(last)
                text = r.text
                self.fails = 0
                m = re.search(r'(?i)http-equiv\s*=\s*"?refresh"?[^>]*URL\s*=\s*([^"\'>\s]+)', text)
                if m and len(text) < 600:
                    if hops >= MAX_HOPS:
                        raise FetchError("meta-refresh qua nhieu hop")
                    return await self.get(urljoin(str(r.url), htmllib.unescape(m.group(1))), cl, hops + 1)
                return text
            except FetchError:
                raise
            except (httpx.TimeoutException, httpx.TransportError, httpx.TooManyRedirects) as e:
                self.n_err += 1
                last = last or type(e).__name__
                await self._wait(min(2 ** (i + 1), 30))
        self.fails += 1
        if self.fails >= CIRCUIT_BREAK:
            log.error("[STOP] %d loi lien tiep -> dung (nghi bi chan IP)", self.fails)
            self.stop.set()
        raise FetchError(last or "loi mang")

    async def _wait(self, secs: float):
        end = time.time() + secs
        while time.time() < end and not self.stop.is_set():
            await asyncio.sleep(0.2)


# =========================================================================== DISCOVERY (count)
def _bad_page(rr: list[dict], seen: set, n: int, max_start: int) -> bool:
    """Trang dang ngo: khong them key moi, hoac ngan (<30 dong) trong khi chua toi trang cuoi."""
    return not ({r["key"] for r in rr} - seen) or (len(rr) < PAGE_SIZE and n < max_start)


async def crawl_source(net: Net, db: DB, st: str, sid: str, name: str):
    """Danh muc (stateless): moi trang muon 1 client bat ky -> tai song song theo dot.
    Trang ngan/khong key moi: thu lai 1 lan; chi dung sau 3 trang xau LIEN TIEP (khong dung vi 1 trang loi)."""
    url = lambda n: f"{BASE}/search_category.php?id={sid}&start={n}&lang=vn"
    first = await net.get(url(0))
    rows, max_start = parse_listing(first)
    if not rows:
        log.info("[MISS] %s %s %s: trang dau rong", st, sid, name)
    seen = {r["key"] for r in rows}
    db.add_rows(rows, st, name)
    total, streak = len(rows), 0
    starts = list(range(PAGE_SIZE, min(max_start, MAX_PAGES * PAGE_SIZE) + 1, PAGE_SIZE))
    for i in range(0, len(starts), net.conc):
        wave = starts[i:i + net.conc]
        for n, h in zip(wave, await asyncio.gather(*(net.get(url(n)) for n in wave))):
            rr, _ = parse_listing(h)
            if _bad_page(rr, seen, n, max_start):
                rr2, _ = parse_listing(await net.get(url(n)))
                rr = rr2 if len(rr2) > len(rr) else rr
            if _bad_page(rr, seen, n, max_start):
                streak += 1
                log.warning("[WARN] %s %s %s start=%d: trang xau (%d dong), lien tiep %d", st, sid, name, n, len(rr), streak)
            else:
                streak = 0
            seen |= {r["key"] for r in rr}
            db.add_rows(rr, st, name)
            total += len(rr)
        if streak >= 3:
            log.warning("[WARN] %s %s %s: 3 trang xau lien tiep -> dung phan trang", st, sid, name)
            break
    return total, len(seen)


async def crawl_stateful(net: Net, db: DB, st: str, sid: str, name: str):
    """prov/ctry/ie: GET search_post (set search vao session, tra meta-refresh toi start=0) roi search_result.
    Giu nguyen 1 client (1 PHPSESSID) cho ca nguon vi state nam trong session. Retry = post lai tren cung client."""
    post = f"{BASE}/search_post.php?id={sid}&type={st}"
    async with net.session() as cl:
        rows, max_start = parse_listing(await net.get(post, cl))
        if not rows:
            log.info("[MISS] %s %s %s: trang dau rong", st, sid, name)
        seen = {r["key"] for r in rows}
        db.add_rows(rows, st, name)
        total, streak = len(rows), 0
        for n in range(PAGE_SIZE, min(max_start, MAX_PAGES * PAGE_SIZE) + 1, PAGE_SIZE):
            page = f"{BASE}/search_result.php?id={sid}&type={st}&start={n}&lang=vn"
            rr, _ = parse_listing(await net.get(page, cl))
            if _bad_page(rr, seen, n, max_start):
                await net.get(post, cl)                       # dat lai search state roi thu lai 1 lan
                rr2, _ = parse_listing(await net.get(page, cl))
                rr = rr2 if len(rr2) > len(rr) else rr
            if _bad_page(rr, seen, n, max_start):
                streak += 1
                log.warning("[WARN] %s %s %s start=%d: trang xau (%d dong), lien tiep %d", st, sid, name, n, len(rr), streak)
                if streak >= 3:
                    log.warning("[WARN] %s %s %s: 3 trang xau lien tiep -> dung phan trang", st, sid, name)
                    seen |= {r["key"] for r in rr}
                    db.add_rows(rr, st, name)
                    total += len(rr)
                    break
            else:
                streak = 0
            seen |= {r["key"] for r in rr}
            db.add_rows(rr, st, name)
            total += len(rr)
    return total, len(seen)


async def discover(net: Net, db: DB, refresh: bool, limit: int = 0) -> list[list]:
    t0 = time.time()
    html = await net.get(f"{BASE}/index.php?lang=vn")
    srcs, banner = parse_home(html)
    db.meta("banner", banner)
    for t, i, n in srcs:
        db.c.execute("INSERT OR IGNORE INTO src(type,id,name) VALUES(?,?,?)", (t, i, n))
    if refresh:
        db.c.execute("UPDATE src SET done=0")
    db.c.commit()
    log.info("[SYS] trang chu: %d nguon (banner %s cong ty)", len(srcs), banner)
    # nguon co state truoc (tuan tu trong 1 session -> dai nhat), danh muc sau (song song theo trang)
    todo = db.c.execute("SELECT type,id,name FROM src WHERE done=0 ORDER BY type DESC, id").fetchall()
    todo = todo[:limit] if limit else todo
    log.info("[SYS] can quet %d nguon (bo qua %d da xong)", len(todo), len(srcs) - len(todo))
    net.total, net.done = len(todo), 0
    outer = asyncio.Semaphore(net.conc)

    async def run(t, i, n):
        async with outer:
            if net.stop.is_set():
                return
            try:
                total, uniq = await (crawl_source if t == "cat" else crawl_stateful)(net, db, t, i, n)
            except FetchError as e:
                if str(e) != "stop":
                    net.bad += 1
                    log.error("[ERR] nguon %s %s %s: %s", t, i, n, e)
                return
            db.c.execute("UPDATE src SET rows=?, done=1 WHERE type=? AND id=?", (total, t, i))
            db.c.commit()
            net.done += 1
            net.ok += 1
            net.named = db.c.execute("SELECT COUNT(*) FROM co WHERE mid=''").fetchone()[0]
            log.info("[OK] %s %s %s: %d dong, %d key moi/nguon, tong cong ty %d", t, i, n, total, uniq, db.n_co())

    await asyncio.gather(*(run(*s) for s in todo))
    log.info("[SYS] count %s sau %.0fs, %d request, %d loi mang",
             "DUNG" if net.stop.is_set() else "xong", time.time() - t0, net.n_req, net.n_err)
    return count_report(db)


# =========================================================================== DETAIL
def pick_lang(langs: str) -> str:
    for l in ("vn", "en", "jp"):
        if l in langs:
            return l
    return "vn"


def profile_url(mid: str, page: str, lang: str) -> str:
    return f"{BASE}/mem_profile.php?id={mid}" + (f"&page={page}" if page else "") + f"&lang={lang}"


def write_company_md(out: pathlib.Path, mid: str, url: str, p: dict, subs: list[tuple[str, str]], slug: str):
    """Ghi mo_ta/cong_ty/{MID}_{slug}.md; tra ve duong dan tuong doi hoac ''."""
    if not p["intro"] and not subs:
        return ""
    rel = f"mo_ta/cong_ty/{mid}_{safe_filename(slug or p['name'], 60)}.md"
    parts = [f"# {p['name'] or mid}", "", f"- Ho so: {url}"]
    if p["fl_url"]:
        parts.append(f"- Trang Fact-Link: {p['fl_url']}")
    if p["intro"]:
        parts += ["", "## Giới thiệu", "", p["intro"]]
    for title, text in subs:
        parts += ["", f"## {title}", "", text]
    f = out / rel
    f.parent.mkdir(parents=True, exist_ok=True)
    f.write_text("\n".join(parts).strip() + "\n", "utf-8")
    return rel


async def crawl_company(net: Net, db: DB, out: pathlib.Path, key: str, mid: str, page: str,
                        langs: str, slug: str, lang: str = ""):
    """Profile + cac trang con tren CUNG 1 client (tuan tu), ghi file .md + dong det."""
    lang = lang or pick_lang(langs)
    url = profile_url(mid, page, lang)
    status, f, rel, n_subs = "OK", {}, "", 0
    async with net.session() as cl:
        if net.stop.is_set():
            return
        try:
            p, err = None, None
            for L in [lang] + [x for x in ("vn", "en", "jp") if x != lang]:   # fallback ngon ngu
                try:
                    cand = parse_profile(await net.get(profile_url(mid, page, L), cl), mid)
                except FetchError as e:
                    if str(e) == "stop":
                        raise
                    err = e
                    continue
                if cand["ok"]:
                    p, lang, url = cand, L, profile_url(mid, page, L)
                    break
                err = FetchError("khong phai trang ho so")
            if p is None:
                raise err
            subs = []
            for title, su in p["subpages"][:MAX_SUBS]:
                try:
                    _, text = parse_subpage(await net.get(su, cl))
                except FetchError as e:
                    if str(e) == "stop":
                        raise
                    log.warning("[WARN] %s trang con %s: %s", mid, su, e)
                    continue
                if text:
                    subs.append((title, text))
            rel, n_subs = write_company_md(out, mid, url, p, subs, slug), len(subs)
            f = dict(p["fields"])
            f["Dia chi"] = "\n".join(p["addresses"])
            f["Email"] = "\n".join(p["emails"])
            f["Thong so"] = "\n".join(p["other"])
            f["Cap nhat"], f["Anh"], f["URL FL"] = p["updated"], p["image"], p["fl_url"]
            f["Ten (profile)"], f["Mo ta ngan (profile)"] = p["name"], p["short"]
            log.info("[OK] %s %s lang=%s: %d truong, %d trang con, %d dia chi", mid, p["name"], lang,
                     len(p["fields"]), n_subs, len(p["addresses"]))
        except FetchError as e:
            if str(e) == "stop":
                return          # dang dung giua chung: KHONG luu cong ty dang do -> lan sau lam lai
            status = f"Loi: {e}"
            log.error("[ERR] %s: %s", mid, e)
    db.c.execute("INSERT OR REPLACE INTO det VALUES(?,?,?,?,?,?)",
                 (key, status, lang, json.dumps(f, ensure_ascii=False), n_subs, rel))
    if db.c.execute("SELECT name FROM co WHERE key=?", (key,)).fetchone()[0] == "" and f.get("Ten (profile)"):
        db.c.execute("UPDATE co SET name=? WHERE key=?", (f["Ten (profile)"], key))
    db.c.commit()
    net.done += 1
    net.ok += status == "OK"
    net.bad += status != "OK"


async def crawl_many(net: Net, db: DB, out: pathlib.Path, rows: list[tuple]):
    """rows = (key, mid, page, langs, slug[, lang]). Gioi han song song = so client trong pool."""
    net.total, net.done, net.ok, net.bad = len(rows), 0, 0, 0
    net.named = db.c.execute("SELECT COUNT(*) FROM co WHERE mid=''").fetchone()[0]
    for r in await asyncio.gather(*(crawl_company(net, db, out, *r) for r in rows), return_exceptions=True):
        if isinstance(r, Exception):
            log.error("[ERR] crawl_company: %s: %s", type(r).__name__, r)


async def detail(net: Net, db: DB, out: pathlib.Path, limit: int, refresh: bool):
    where = "" if refresh else "AND (d.key IS NULL OR d.status LIKE 'Loi%')"
    todo = db.c.execute(
        f"SELECT c.key,c.mid,c.page,c.langs,c.slug FROM co c LEFT JOIN det d ON d.key=c.key "
        f"WHERE c.mid<>'' {where} ORDER BY c.mid").fetchall()
    total_mid = db.c.execute("SELECT COUNT(*) FROM co WHERE mid<>''").fetchone()[0]
    todo = todo[:min(limit or MAX_ITEMS, MAX_ITEMS)]
    log.info("[SYS] detail: %d cong ty se crawl (tong co MID: %d, da OK: %d)", len(todo), total_mid,
             db.c.execute("SELECT COUNT(*) FROM det WHERE status='OK'").fetchone()[0])
    await crawl_many(net, db, out, todo)


def parse_items(items: list[str]) -> list[tuple]:
    """Moi muc: URL ho so (mem_profile.php?id=..&page=..&lang=..) hoac MID tran. -> (key,mid,page,langs,slug,lang)."""
    rows, seen = [], set()
    for it in items:
        it = str(it).strip()
        m = re.search(r"id=(\d+)(?:&(?:amp;)?page=(\d+))?(?:&(?:amp;)?lang=(\w+))?", it)
        if m and "mem_" in it:
            mid, page, lang = m.group(1).zfill(8), m.group(2) or "", m.group(3) or "vn"
        elif re.fullmatch(r"\d{1,8}", it):
            mid, page, lang = it.zfill(8), "", "vn"
        else:
            if it:
                log.error("[ERR] khong nhan ra MID/URL ho so: %s", it)
            continue
        if mid not in seen:
            seen.add(mid)
            rows.append((mid, mid, page, "", "", lang))
    return rows


async def crawl_items(net: Net, db: DB, out: pathlib.Path, items: list[str]):
    """mode link / Excel: chi crawl dung cac MID / URL ho so duoc cho (them vao DB neu chua co)."""
    rows = parse_items(items)[:MAX_ITEMS]
    for key, mid, page, _, _, lang in rows:
        old = db.c.execute("SELECT page,langs,slug FROM co WHERE key=?", (key,)).fetchone()
        if old is None:
            db.c.execute("INSERT INTO co(key,mid,page,name,nat,region,short,langs,slug) VALUES(?,?,?,?,?,?,?,?,?)",
                         (key, mid, page, "", "", "", "", lang, ""))
    db.c.commit()
    todo = []
    for key, mid, page, _, _, lang in rows:
        p, langs, slug = db.c.execute("SELECT page,langs,slug FROM co WHERE key=?", (key,)).fetchone()
        todo.append((key, mid, page or p, langs, slug, lang))
    log.info("[SYS] link: %d cong ty", len(todo))
    await crawl_many(net, db, out, todo)


def read_excel_items(path: str, sheet: str, col: str) -> list[str]:
    """Doc 1 cot (MID hoac URL ho so) tu Excel bang openpyxl."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(c or "").strip() for c in next(it, [])]
    if col not in header:
        wb.close()
        raise ValueError(f"Khong thay cot '{col}' trong sheet")
    idx = header.index(col)
    out = [str(r[idx]).strip() for r in it if idx < len(r) and r[idx] not in (None, "")]
    wb.close()
    return out


# =========================================================================== PRODUCTS
def write_product_md(out: pathlib.Path, proid: str, d: dict) -> str:
    if not d["desc"]:
        return ""
    rel = f"mo_ta/san_pham/{proid}.md"
    f = out / rel
    f.parent.mkdir(parents=True, exist_ok=True)
    f.write_text(f"# {d['title']}\n\n- SKU: #{d['sku']}\n- Cong ty: {d['company']}\n\n{d['desc']}\n", "utf-8")
    return rel


async def products(net: Net, db: DB, out: pathlib.Path, limit: int, refresh: bool):
    # products_list.php?start=N = trang tom tat theo nhom (xem truoc vai SP + link SHOW MORE id=CAT)
    ids, cats, n = [], [], 0
    while n <= MAX_PAGES * 24:
        h = await net.get(f"{BASE}/products_list.php?start={n}")
        got = parse_prod_ids(h)
        if not got:
            break
        ids += got
        cats += re.findall(r"products_list\.php\?id=(\d+)&(?:amp;)?start=0", h)
        n += 24
    cats = list(dict.fromkeys(cats))
    log.info("[SYS] san pham: %d trang tom tat, %d nhom", n // 24, len(cats))

    async def one_cat(cid):
        first = await net.get(f"{BASE}/products_list.php?id={cid}&start=0")
        got = parse_prod_ids(first)
        own = [int(x) for x in re.findall(rf"products_list\.php\?id={cid}&(?:amp;)?start=(\d+)", first)]
        for h in await asyncio.gather(*(net.get(f"{BASE}/products_list.php?id={cid}&start={s}")
                                        for s in range(24, min(max(own, default=0), MAX_PAGES * 24) + 1, 24))):
            got += parse_prod_ids(h)
        return got

    for got in await asyncio.gather(*(one_cat(c) for c in cats), return_exceptions=True):
        if isinstance(got, list):
            ids += got
        elif str(got) != "stop":
            log.error("[ERR] nhom san pham: %s", got)
    for x in dict.fromkeys(ids):
        db.c.execute("INSERT OR IGNORE INTO prod(proid,status,data) VALUES(?,NULL,'{}')", (x,))
    db.c.commit()
    n_all = db.c.execute("SELECT COUNT(*) FROM prod").fetchone()[0]
    where = "" if refresh else "WHERE status IS NULL OR status LIKE 'Loi%'"
    todo = [r[0] for r in db.c.execute(f"SELECT proid FROM prod {where} ORDER BY CAST(proid AS INT)")]
    todo = todo[:limit] if limit else todo
    log.info("[SYS] san pham: %d unique, crawl %d", n_all, len(todo))
    net.total, net.done, net.ok, net.bad = len(todo), 0, 0, 0

    async def one(pid):
        try:
            d = parse_product(await net.get(f"{BASE}/product_detail.php?proid={pid}"))
            if not d["title"]:
                raise FetchError("khong phai trang san pham")
            d["file"] = write_product_md(out, pid, d)
            db.c.execute("UPDATE prod SET status='OK', data=? WHERE proid=?", (json.dumps(d, ensure_ascii=False), pid))
            net.ok += 1
            log.info("[OK] san pham %s %s", pid, d["title"])
        except FetchError as e:
            if str(e) == "stop":
                return
            db.c.execute("UPDATE prod SET status=?, data='{}' WHERE proid=?", (f"Loi: {e}", pid))
            net.bad += 1
            log.error("[ERR] san pham %s: %s", pid, e)
        db.c.commit()
        net.done += 1
    await asyncio.gather(*(one(p) for p in todo))


# =========================================================================== EXPORT
def _cell(v) -> str | int | float:
    """Moi o la scalar: list -> noi bang xuong dong, dict -> 'k: v', bo ky tu XML cam."""
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, (list, tuple, set)):
        v = "\n".join(str(_cell(x)) for x in v)
    elif isinstance(v, dict):
        v = "\n".join(f"{k}: {_cell(x)}" for k, x in v.items())
    return ILLEGAL_CHARACTERS_RE.sub("", str(v))[:32000]


def _write_sheet(wb: Workbook, title: str, header: list[str], rows: list[list], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for r in rows:
        ws.append([_cell(c) for c in r])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = 60 if h in WRAP_COLS else 18
        if h in WRAP_COLS:
            for cell in ws[get_column_letter(i)][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def export_excel(db: DB, path: pathlib.Path) -> pathlib.Path:
    mem: dict[str, dict[str, list[str]]] = {}
    for k, st, sn in db.c.execute("SELECT key,stype,sname FROM mem ORDER BY sname"):
        mem.setdefault(k, {}).setdefault(st, []).append(sn)
    det = {r[0]: r[1:] for r in db.c.execute("SELECT key,status,lang,fields,subs,file FROM det")}
    rows = []
    q = "SELECT key,mid,page,name,nat,region,short,langs,slug FROM co ORDER BY (mid=''), mid, name"
    for key, mid, page, name, nat, region, short, langs, slug in db.c.execute(q).fetchall():
        m = mem.get(key, {})
        d = det.get(key)
        f = json.loads(d[2]) if d and d[2] else {}
        if not mid:
            status = "Chi co ten"
        elif d is None:
            status = "Chua crawl"
        else:
            status = d[0]
        lang = d[1] if d else ""
        row = {
            "MID": mid, "Page ID": page, "Ten (list)": name, "Status": status,
            "Quoc tich": nat or "\n".join(m.get("ctry", [])), "Khu vuc": region,
            "Danh muc": "\n".join(m.get("cat", [])), "Tinh/TP": "\n".join(m.get("prov", [])),
            "KCN": "\n".join(m.get("ie", [])), "Mo ta ngan": short or f.get("Mo ta ngan (profile)", ""),
            "Ngon ngu": lang, "So trang con": d[3] if d else "", "File mo ta": d[4] if d else "",
            "URL ho so": profile_url(mid, page, lang or pick_lang(langs)) if mid else "",
        }
        for c in COMPANY_COLS:
            if c not in row:
                row[c] = f.get(c, "")
        rows.append([row[c] for c in COMPANY_COLS])
    prow = []
    for pid, data in db.c.execute("SELECT proid,data FROM prod ORDER BY CAST(proid AS INT)"):
        d = json.loads(data or "{}")
        if not d:
            continue
        prow.append([pid, d["title"], d["category"], d["sku"], d["stock"], d["mid"], d["company"],
                     d.get("file", ""), d["image"], f"{BASE}/product_detail.php?proid={pid}"])
    rep = count_report(db)
    wb = Workbook()
    _write_sheet(wb, "Cong ty", COMPANY_COLS, rows, first=True)
    _write_sheet(wb, "San pham", PRODUCT_COLS, prow)
    _write_sheet(wb, "Dem", rep[0], rep[1:])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    log.info("[OK] xuat Excel: %s (%d cong ty, %d san pham)", path, len(rows), len(prow))
    return path


# =========================================================================== CORE JOB (CLI + GUI dung chung)
class _SinkHandler(logging.Handler):
    """Day dong log vao queue cua GUI."""

    def __init__(self, sink):
        super().__init__()
        self.sink = sink

    def emit(self, rec):
        self.sink(self.format(rec))


def setup_logging(out: pathlib.Path, sink=None) -> None:
    out.mkdir(parents=True, exist_ok=True)
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass
    fmt = logging.Formatter("%(asctime)s %(message)s", "%H:%M:%S")
    log.setLevel(logging.INFO)
    for h in list(log.handlers):
        h.close()
    log.handlers.clear()
    for h in (logging.FileHandler(out / "factlink.log", encoding="utf-8"),
              _SinkHandler(sink) if sink else logging.StreamHandler(sys.stdout)):
        h.setFormatter(fmt)
        log.addHandler(h)


def table_lines(rows: list[list]) -> list[str]:
    w = [max(len(str(r[i])) for r in rows) for i in range(len(rows[0]))]
    return ["  ".join(str(c).ljust(w[i]) for i, c in enumerate(r)).rstrip() for r in rows]


async def run_mode(args, db: DB, out: pathlib.Path, stop, pause, on_net=None) -> None:
    net = Net(args.conc, args.delay, stop, pause)
    if on_net:
        on_net(net)
    mode = args.mode
    try:
        if mode in ("count", "all"):
            for line in table_lines(await discover(net, db, args.refresh, args.limit if mode == "count" else 0)):
                log.info(line)
        if mode in ("detail", "all") and not net.stop.is_set():
            await detail(net, db, out, args.limit, args.refresh)
        if mode in ("products", "all") and not net.stop.is_set():
            await products(net, db, out, args.limit, args.refresh)
        if mode == "link":
            items = list(args.url or [])
            if args.excel:
                items += read_excel_items(args.excel, args.sheet, args.col)
            await crawl_items(net, db, out, items)
    finally:
        log.info("[SYS] %d request, %d loi mang", net.n_req, net.n_err)
        await net.aclose()


def run_job(args, stop, pause, on_net=None) -> None:
    """1 job day du: crawl theo mode -> LUON xuat Excel trong finally (ke ca sau Stop / crash)."""
    out = pathlib.Path(args.out)
    db = DB(out / "factlink_state.db")
    try:
        if args.mode != "export":
            asyncio.run(run_mode(args, db, out, stop, pause, on_net))
    finally:
        export_excel(db, out / "factlink.xlsx")
        db.c.close()


# =========================================================================== GUI
def run_gui(root=None, mainloop: bool = True):
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    root_dir = next(p for p in (HERE, *HERE.parents) if (p / "gui_common.py").exists())
    sys.path.insert(0, str(root_dir))
    from gui_common import attach_excel_picker

    root = root or tk.Tk()
    root.title("Fact-Link Vietnam Crawler")
    root.geometry("1100x720")
    ttk.Style().theme_use("clam")

    q: queue.Queue[str] = queue.Queue()
    stop_evt, pause_evt = threading.Event(), threading.Event()
    state: dict = {"running": False, "net": None, "thread": None}

    left = ttk.Frame(root, padding=10)
    left.pack(side="left", fill="y")
    right = ttk.Frame(root, padding=(0, 10, 10, 10))
    right.pack(side="right", fill="both", expand=True)

    mode = tk.StringVar(value="count")
    ttk.Label(left, text="Pipeline", font=("Segoe UI", 10, "bold")).pack(anchor="w")
    for val, label in [("count", "1. Dem cong ty (count)"), ("all", "2. Crawl ALL"),
                       ("detail", "3. Chi tiet ho so"), ("products", "4. San pham"),
                       ("link", "5. Theo LINK / MID / Excel"), ("export", "6. Chi EXPORT Excel")]:
        ttk.Radiobutton(left, text=label, value=val, variable=mode).pack(anchor="w")

    ttk.Label(left, text="\nLink ho so / MID (moi dong 1 link):").pack(anchor="w")
    box = tk.Text(left, width=40, height=5)
    box.pack(anchor="w")

    ttk.Label(left, text="Excel tuy chon (cot MID hoac URL ho so):").pack(anchor="w", pady=(6, 0))
    xls, sheet, col, _unused = tk.StringVar(), tk.StringVar(), tk.StringVar(), tk.StringVar()
    before = set(left.winfo_children())
    attach_excel_picker(left, xls, sheet, col, _unused)
    for w in [w for w in left.winfo_children() if w not in before]:
        # gui_common them dong "Thu muc anh" + chu thich ve anh: khong dung o day (site nay khong tai anh) -> bo
        texts = [str(c.cget("text")) for c in w.winfo_children() if c.winfo_class() in ("TLabel", "Label")]
        if w.winfo_class() in ("TLabel", "Label"):
            texts.append(str(w.cget("text")))
        if any(t.startswith("Thu muc anh") or "anh dat canh" in t for t in texts):
            w.destroy()

    def guess_col(*_):
        try:
            cols = list(load_workbook(xls.get(), read_only=True).worksheets[0][1])
            names = [str(c.value or "") for c in cols]
        except Exception:
            return
        pick = next((n for n in names if re.search(r"mid|url|link|^id$", n, re.I)), "")
        if pick:
            col.set(pick)
    xls.trace_add("write", lambda *_: root.after(500, guess_col))

    outdir = tk.StringVar(value=str(HERE / "out_factlink"))
    conc, delay, limit = tk.IntVar(value=8), tk.DoubleVar(value=0.3), tk.IntVar(value=0)
    refresh = tk.BooleanVar(value=False)
    fr = ttk.Frame(left)
    fr.pack(anchor="w", pady=(8, 1))
    ttk.Label(fr, text="Thu muc ket qua:", width=15).pack(side="left")
    ttk.Entry(fr, textvariable=outdir, width=18).pack(side="left")
    ttk.Button(fr, text="...", width=3,
               command=lambda: outdir.set(filedialog.askdirectory() or outdir.get())).pack(side="left")
    for label, var in [("Concurrency (1-16):", conc), ("Delay (s):", delay), ("Limit (0=het):", limit)]:
        f = ttk.Frame(left)
        f.pack(anchor="w", pady=1)
        ttk.Label(f, text=label, width=18).pack(side="left")
        ttk.Entry(f, textvariable=var, width=8).pack(side="left")
    ttk.Checkbutton(left, text="Lam lai tu dau (refresh)", variable=refresh).pack(anchor="w")

    console = tk.Text(right, bg="#12161c", fg="#d6e2ee", insertbackground="#fff", wrap="none")
    console.pack(fill="both", expand=True)
    pb = ttk.Progressbar(right, mode="determinate")
    pb.pack(fill="x", pady=4)
    stat = ttk.Label(right, text="0/0 | OK 0 | loi 0 | chi co ten 0")
    stat.pack(anchor="w")

    def drain() -> None:
        for _ in range(300):                      # tran moi nhip de GUI khong dung hinh
            try:
                console.insert("end", q.get_nowait() + "\n")
            except queue.Empty:
                break
            console.see("end")
        if int(console.index("end-1c").split(".")[0]) > 5000:
            console.delete("1.0", "1000.0")
        n = state["net"]
        if n is not None:
            pb["maximum"] = max(n.total, 1)
            pb["value"] = n.done
            stat.config(text=f"{n.done}/{n.total} | OK {n.ok} | loi {n.bad} | chi co ten {n.named}")
        root.after(100, drain)

    def worker(args) -> None:
        xlsx = pathlib.Path(args.out) / "factlink.xlsx"
        failed = False
        try:
            run_job(args, stop_evt, pause_evt, lambda n: state.update(net=n))
        except BaseException as exc:
            failed = True
            q.put(f"[FATAL] {type(exc).__name__}: {exc}")
        finally:
            state["running"] = False
            q.put(f"[SYS] DUNG do loi - Excel phan da xong: {xlsx}" if failed else f"[SYS] XONG - Excel da ghi: {xlsx}")

    def start() -> None:
        if state["running"]:
            return
        try:
            items = [x for x in box.get("1.0", "end").split("\n") if x.strip()]
            args = argparse.Namespace(
                mode=mode.get(), limit=max(0, limit.get()), conc=max(1, min(conc.get(), MAX_CONC)),
                delay=max(0.0, delay.get()), out=outdir.get(), refresh=refresh.get(), url=items,
                excel=xls.get() if mode.get() == "link" and col.get() else "", sheet=sheet.get(), col=col.get())
        except tk.TclError:
            messagebox.showwarning("Sai du lieu", "Concurrency / Delay / Limit phai la so.")
            return
        if args.mode == "link" and not (items or args.excel):
            messagebox.showwarning("Thieu du lieu", "Dan link/MID (moi dong 1 muc) hoac chon Excel + cot.")
            return
        setup_logging(pathlib.Path(args.out), q.put)
        stop_evt.clear()
        pause_evt.clear()
        state.update(running=True, net=None)
        console.delete("1.0", "end")
        t = threading.Thread(target=worker, args=(args,), daemon=True)
        state["thread"] = t
        t.start()

    def toggle_pause() -> None:
        if pause_evt.is_set():
            pause_evt.clear()
            btn_pause.config(text="Pause")
            q.put("[SYS] Tiep tuc")
        else:
            pause_evt.set()
            btn_pause.config(text="Tiep tuc")
            q.put("[SYS] Tam dung (Pause)")

    def stop() -> None:
        stop_evt.set()
        if pause_evt.is_set():
            toggle_pause()
        q.put("[STOP] Dang dung... se xuat Excel phan da xong")

    def export() -> None:
        def go():
            try:
                db = DB(pathlib.Path(outdir.get()) / "factlink_state.db")
                p = export_excel(db, pathlib.Path(outdir.get()) / "factlink.xlsx")
                db.c.close()
                q.put(f"[OK] Da xuat: {p}")
            except Exception as e:
                q.put(f"[ERR] Export loi: {e}")
        threading.Thread(target=go, daemon=True).start()

    def open_dir() -> None:
        p = pathlib.Path(outdir.get())
        p.mkdir(parents=True, exist_ok=True)
        try:
            os.startfile(p)
        except Exception as e:
            q.put(f"[WARN] Khong mo duoc thu muc: {e}")

    fr_btn = ttk.Frame(left)
    fr_btn.pack(anchor="w", pady=8)
    ttk.Button(fr_btn, text="START", command=start).pack(side="left")
    btn_pause = ttk.Button(fr_btn, text="Pause", command=toggle_pause)
    btn_pause.pack(side="left")
    ttk.Button(fr_btn, text="Stop", command=stop).pack(side="left")
    ttk.Button(fr_btn, text="EXPORT", command=export).pack(side="left")
    ttk.Button(left, text="Mo thu muc ket qua", command=open_dir).pack(anchor="w")

    def on_close() -> None:
        if state["running"] and not messagebox.askyesno("Thoat", "Dang crawl. Thoat va tat luon?"):
            return
        stop_evt.set()
        t = state["thread"]
        if t is not None and t.is_alive():
            t.join(timeout=15)          # cho worker xuat Excel phan da xong
        try:
            root.destroy()
        finally:
            os._exit(0)                 # kill cung: khong de thread asyncio / socket treo

    root.protocol("WM_DELETE_WINDOW", on_close)
    drain()
    ui = argparse.Namespace(root=root, q=q, state=state, mode=mode, limit=limit, conc=conc, delay=delay,
                            outdir=outdir, box=box, start=start, pause=toggle_pause, stop=stop,
                            export=export, on_close=on_close, console=console, stat=stat)
    if mainloop:
        root.mainloop()
    return ui


# =========================================================================== CLI
def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="fact-link.com.vn crawler (khong tham so = mo GUI)")
    ap.add_argument("--mode", choices=["count", "all", "detail", "products", "export", "link"])
    ap.add_argument("--limit", type=int, default=0, help="count: so nguon; detail/products/all: so muc")
    ap.add_argument("--conc", type=int, default=8, help=f"so session song song (1-{MAX_CONC})")
    ap.add_argument("--delay", type=float, default=0.3)
    ap.add_argument("--out", default=str(HERE / "out_factlink"))
    ap.add_argument("--refresh", action="store_true")
    ap.add_argument("--url", nargs="*", help="URL ho so hoac MID (mode link)")
    ap.add_argument("--excel", default="", help="Excel co cot MID/URL ho so (mode link)")
    ap.add_argument("--sheet", default="")
    ap.add_argument("--col", default="")
    args = ap.parse_args(argv)
    if not args.mode:
        run_gui()
        return 0
    args.conc = max(1, min(args.conc, MAX_CONC))
    setup_logging(pathlib.Path(args.out))
    try:
        run_job(args, threading.Event(), threading.Event())
    except KeyboardInterrupt:
        log.warning("[STOP] Ctrl+C -> dung, da xuat phan da xong")
    except Exception as e:
        log.exception("[FATAL] %s", e)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
