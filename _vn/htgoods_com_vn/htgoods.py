"""
htgoods.com.vn (HT GOODS) — httpx, NO browser, NO Selenium.

RECON (2026-09-13)
------------------
Nen tang : **Sapo/Bizweb** (robots.txt "we use Sapo as our ecommerce platform",
           CDN bizweb.dktcdn.net, cookie _sapo_t). Hang Nhat (MonotaRO resale):
           TRUSCO, THK, TOHNICHI, KANON, TONE, WERA, KYOCERA... ten dang
           "<HANG> Viet Nam | <ten EN> MWxxxxxxx".
API      : CO, JSON thuan, khong can login:
    /products.json?limit=250&page=N   -> {"products":[{id, name, alias, vendor, product_type,
         price, tags[], summary, content(HTML), featured_image, images[], variants[], available}]}
    /products/<alias>.json (= .js)    -> 1 SP, cung record (dung cho link/ma/va SP sot)
    /collections.json?limit=250&page=N -> 566 collection (KHONG co products_count)
    /collections/<alias>/products.json?limit=250&page=N -> full record
    /search?q=<ma>&view=json          -> {"results_total":N,"results":[{title,url,thumb,price,id}]}
         !! CO tac dung o site nay (khac insumart) nhung CHI 4 KQ/TRANG (limit/pagesize bi bo
         qua) -> hop tra MA (1 hit), KHONG hop keyword (trusco = 13.833 KQ).
         Tim duoc theo MW (barcode), SKU hang (DTC-N500REV-A) va ca ma bo gach (DTC-N500REVA).
    /search?q=<tu khoa>&page=N        -> HTML server-side, 20 SP/trang, tong "Co N ket qua",
         link SP: <a class="thumb-img" href="/<alias>"> trong doan "title-head".."pagination".
    /sitemap.xml -> sitemap_products_1..11.xml -> **100.724 URL SP** (CHUAN VANG, ~10k/file)

    !! BAY #1 — CAP OFFSET 30.000 (giong insumart): /products.json page 121 x 250 -> HTTP 500.
       Paginate thang chi lay 30k/100.7k = MAT 70% MA KHONG BAO LOI.
       => ALL quet theo COLLECTION (do that: khong collection nao cham cap, lon nhat <30k;
       danh muc goc "htgoods-*" + "trusco-viet-nam" > 4.750 SP), uu tien "htgoods-*",
       DUNG SOM khi da phu du sitemap, con thieu -> fetch /products/<alias>.json tung SP.
    BAY #2: trang ngan giua chung khong phai het -> chi dung khi RONG / khong co id moi.
data     : 1 variant/SP (do 1.000 SP). barcode = ma he thong **MWxxxxxxx** (999/1000),
           sku = **ma hang** (DTC-N500REV-A). price = 0 het ("Lien he"), inventory = 0, unit null.
spec     : KHONG co <table>. Trong `content` (96%): sau <strong>Specifications in detail</strong>
           la cac dong "Ten: Gia tri<br>" (xen dong bullet khong co ':'; du lieu goc co dong bi
           ngat doi "2.5m Dust<br>collection"), ket thuc bang <h1 style='color: white;'>HANG SKU</h1>
           (the an -> lay lam Model). Truoc do la cac <section><h3>Features|Applications|Purpose|
           Note</h3><ul><li>... -> cot "Mo ta". Key spec tuy nganh -> gop 1 cot "Thong so" (4.6).
rate     : 570 request conc 8 trong 12s -> 0 lan 429/403. robots Disallow /search (chi dung khi
           nguoi dung tra ma/tu khoa, tan suat thap).

Run
---
  htgoods.bat                                              -> GUI
  python htgoods.py --mode all [--limit 500]               -> ~100.7k SP
  python htgoods.py --mode code --code MW1148633           (MW / SKU hang / model)
  python htgoods.py --mode keyword --keyword "THK LM" [--limit 100]
  python htgoods.py --mode link --link https://www.htgoods.com.vn/<alias>
  python htgoods.py --mode excel --excel parts.xlsx --sheet Sheet1 --col Part   (ma hoac link)
"""
from __future__ import annotations

import argparse
import asyncio
import gc
import hashlib
import html as htmlmod
import json
import os
import pathlib
import queue
import re
import threading
import time
from typing import Any, Iterable, Iterator
from urllib.parse import quote

import httpx
import pandas as pd

HERE = pathlib.Path(__file__).parent
BASE = "https://www.htgoods.com.vn"
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/126.0 Safari/537.36")
HEADERS = {"User-Agent": UA, "Accept-Language": "vi-VN,vi;q=0.9,en;q=0.8"}
LOG_FILE = HERE / "htgoods.log"

PAGE_SIZE = 250          # tran cua Sapo
OFFSET_CAP = 30000       # BAY #1: page*limit > 30000 -> HTTP 500
SEARCH_JSON_PAGES = 5    # view=json 4 KQ/trang -> toi da 20 ung vien/ma

CIRCUIT_BREAK = 25       # N loi LIEN TIEP -> nghi bi chan IP -> dung han
HIBERNATE_429 = 10       # N lan 403/429 LIEN TIEP -> ngu dai
MAX_ITEMS = 150000       # tran work-list (catalog that ~100.7k)
MAX_PAGES = 600          # tran phan trang
MAX_REDIRECTS = 5
TIMEOUT = 60.0
GC_EVERY = 5000
CACHE_MAX = 3000         # cache product json o mode ma

COLUMNS = ["Input", "Item No", "Page", "Status", "Category 1", "Category 2",
           "Ma HT", "Thuong hieu", "Model", "Gia", "Con hang", "Product ID", "Variant ID",
           "Mo ta", "Thong so", "Image", "URL"]

TAG_RE = re.compile(r"<[^>]+>")
LOC_RE = re.compile(r"<loc>([^<]+)</loc>")
SPEC_HEAD_RE = re.compile(r"<strong>\s*Specifications in detail\s*</strong>", re.I)
# data goc co the h1 HONG mat dau "<": 'h1 style="color:white;">GTG R1-5-8AJP</h1>'
H1_RE = re.compile(r"<?\bh1(?:\s[^>]*)?>(.*?)</h1>", re.S | re.I)
SECTION_RE = re.compile(r"<h3[^>]*>(.*?)</h3>(.*?)(?=<h3|</section>|$)", re.S | re.I)
LI_RE = re.compile(r"<li[^>]*>(.*?)</li>", re.S | re.I)
LINE_SPLIT_RE = re.compile(r"<br\s*/?>|</p>|</div>|</li>|</tr>", re.I)
SEARCH_ITEM_RE = re.compile(r'<a class="thumb-img" href="/([^"?#]+)"')
SEARCH_TOTAL_RE = re.compile(r"Có\s+(\d+)\s+kết quả")
SEARCH_JSON_ITEM_RE = re.compile(r'"title":"(.*?)","url":"(/[^"]*)"', re.S)


# ================================================================ CORE (khong dung GUI)
def norm_code(s: Any) -> str:
    """Chuan hoa ma de so khop. URL thi GIU NGUYEN hoa/thuong (path phan biet chu)."""
    v = re.sub(r"\s+", "", str(s or ""))
    return v if v.lower().startswith("http") else v.upper()


def alnum(s: Any) -> str:
    """Bo het dau/gach: DTC-N500REV-A == DTC-N500REVA (ten SP hay viet ma bo gach)."""
    return re.sub(r"[^A-Z0-9]", "", str(s or "").upper())


def alias_from_url(u: str) -> str:
    """https://www.htgoods.com.vn/<alias>[?variant=..] -> <alias>. Chap ca /products/<alias>."""
    a = re.sub(r"^https?://[^/]+", "", str(u or "").strip()).split("?")[0].split("#")[0]
    a = a.strip("/")
    if a.startswith("products/"):
        a = a[len("products/"):]
    return re.sub(r"\.(json|js)$", "", a)


def product_json_url(alias: str) -> str:
    return f"{BASE}/products/{quote(alias)}.json"


def text(s: str) -> str:
    return re.sub(r"\s+", " ", htmlmod.unescape(TAG_RE.sub(" ", s or "")).replace("\xa0", " ")).strip()


def parse_content(content: str) -> dict[str, Any]:
    """content HTML -> {spec{ten:gia tri}, bullets[], desc[], model}."""
    c = content or ""
    m1 = H1_RE.search(c)
    model = text(m1.group(1)) if m1 else ""
    mh = SPEC_HEAD_RE.search(c)
    desc_part = c[:mh.start()] if mh else c
    desc: list[str] = []
    for t, body in SECTION_RE.findall(desc_part):
        items = [x for x in (text(li) for li in LI_RE.findall(body)) if x] or [text(body)]
        items = [x for x in items if x]
        if items:
            desc.append(f"{text(t)}: " + "; ".join(items))

    spec: dict[str, str] = {}
    bullets: list[str] = []
    # ~4% SP khong co heading spec -> van vot cap "Ten: Gia tri" trong toan bo content
    spec_part = H1_RE.sub("", c[mh.end():] if mh else c)
    for chunk in LINE_SPLIT_RE.split(spec_part):
        s = text(chunk)
        if not s:
            continue
        k, sep, v = s.partition(":")
        k, v = k.strip(), v.strip()
        if sep and k and v and len(k) <= 80:
            if k not in spec:
                spec[k] = v
                continue
            if spec[k] == v:
                continue
        if mh:                        # bullet chi lay trong khoi spec (khong lap lai Mo ta)
            bullets.append(s)
    if not desc and not mh:
        body = H1_RE.sub("", c)
        whole = "; ".join(x for x in (text(li) for li in LI_RE.findall(body)) if x) or text(body)
        if whole and not spec:
            desc.append(whole[:2000])
    return {"spec": spec, "bullets": bullets, "desc": desc, "model": model}


def fix_img(u: Any) -> str:
    u = u.get("src") if isinstance(u, dict) else u
    u = str(u or "")
    return "https:" + u if u.startswith("//") else u


def parse_product(p: Any, input_code: str = "") -> list[dict]:
    """1 san pham -> 1 record / VARIANT (site hien 1 variant/SP, van giu vong lap cho chac).
    Moi o la scalar (4.6): spec gop san thanh chuoi "Thong so"."""
    if isinstance(p, dict) and not p.get("alias") and isinstance(p.get("product"), dict):
        p = p["product"]
    if not isinstance(p, dict) or not p.get("alias"):
        return []
    info = parse_content(p.get("content") or "")
    if not (info["spec"] and info["desc"] and info["model"]):
        # ~2% SP content rong / chi 1 dong -> mo ta + h1 nam trong summary (<ul><li>)
        s = parse_content(p.get("summary") or "")
        info["model"] = info["model"] or s["model"]
        if not info["spec"]:
            info["spec"], info["bullets"] = s["spec"], s["bullets"]
        if not info["desc"]:
            info["desc"] = s["desc"]
    vendor = str(p.get("vendor") or "")
    model = info["model"]
    if vendor and model.upper().startswith(vendor.upper() + " "):
        model = model[len(vendor) + 1:].strip()
    spec_txt = "\n".join([f"{k}: {v}" for k, v in info["spec"].items()] + info["bullets"])
    url = f"{BASE}/{p['alias']}"
    imgs = [fix_img(i) for i in (p.get("images") or []) if i]
    if not imgs and p.get("featured_image"):
        imgs = [fix_img(p["featured_image"])]

    variants = p.get("variants") or [{}]
    out: list[dict] = []
    for v in variants:
        vimg = fix_img(v.get("featured_image"))
        out.append({
            "Input": input_code,
            "Item No": v.get("sku") or model,
            "Page": p.get("name") or "",
            "Status": "OK",
            "Category 1": vendor,
            "Category 2": p.get("product_type") or "",
            "Ma HT": v.get("barcode") or "",
            "Thuong hieu": vendor,
            "Model": model,
            "Gia": (v.get("price") or p.get("price") or 0) or "Lien he",
            "Con hang": "Co" if v.get("available", p.get("available")) else "Het",
            "Product ID": p.get("id"),
            "Variant ID": v.get("id"),
            "Mo ta": "\n".join(info["desc"]),
            "Thong so": spec_txt,
            "Image": " | ".join(list(dict.fromkeys(i for i in [vimg, *imgs] if i))[:5]),
            "URL": f"{url}?variant={v['id']}" if v.get("id") and len(variants) > 1 else url,
        })
    return out


def not_found(code: str, status: str = "Not Found") -> dict:
    return {"Input": code, "Item No": "" if code.lower().startswith("http") else code,
            "Status": status}


def match_code(recs: list[dict], code: str) -> list[dict]:
    """§4.3 + BAY §10-5: khong lay mu ket qua dau. So ma nhap voi CA: SKU hang (Item No),
    ma he thong MW (Ma HT), Model. Tier: exact -> bo dau/gach -> chua trong (danh dau)."""
    c, ca = norm_code(code), alnum(code)
    keys = ("Item No", "Ma HT", "Model")
    exact = [r for r in recs if c in {norm_code(r.get(k)) for k in keys}]
    if exact:
        return exact
    loose = [r for r in recs if ca and ca in {alnum(r.get(k)) for k in keys}]
    for r in loose:
        r["Status"] = "OK (khop bo dau)"
    if loose:
        return loose
    sub = [r for r in recs if len(ca) >= 4 and any(ca in alnum(r.get(k)) for k in keys)]
    for r in sub:
        r["Status"] = "OK (khop mot phan)"
    return sub


def export_excel(rows: Iterable[dict], path: pathlib.Path) -> int:
    """Spec gop 1 COT 'Thong so' (AGENTS 4.6). Ghi openpyxl write_only = STREAM tung dong:
    crawl ALL ~100k dong khong phai nap het vao RAM/DataFrame. Tra ve so dong."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("htgoods")
    widths = {"Page": 50, "Mo ta": 50, "Thong so": 60, "Image": 40, "URL": 40}
    for i, name in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 16)
    ws.append(COLUMNS)
    wrap = Alignment(wrap_text=True, vertical="top")
    n = 0
    for r in rows:
        line = []
        for name in COLUMNS:
            v = r.get(name, "")
            if v is None:
                v = ""
            if isinstance(v, str):
                v = ILLEGAL_CHARACTERS_RE.sub("", v)[:32000]     # tran 1 o Excel 32.767
            if name in ("Mo ta", "Thong so"):
                cell = WriteOnlyCell(ws, value=v)
                cell.alignment = wrap
                v = cell
            line.append(v)
        ws.append(line)
        n += 1
    wb.save(path)
    return n


class Engine:
    """Crawler core — chay headless/CLI duoc, khong phu thuoc GUI."""

    def __init__(self, conc: int = 8, delay: float = 0.05, log=print,
                 stop_evt: threading.Event | None = None,
                 pause_evt: threading.Event | None = None) -> None:
        self.conc, self.delay, self._log = max(1, conc), max(0.0, delay), log
        self.stop_evt = stop_evt or threading.Event()
        self.pause_evt = pause_evt or threading.Event()
        self.fails = 0
        self.n_block = 0
        self.results: list[dict] = []
        self.keep = True                 # mode ALL: False -> record chi nam trong JSONL (RAM)
        self.rows = 0
        self.done = 0
        self.total = 0
        self.limit = 0
        self.key = ""
        self.done_inputs: set[str] = set()
        self._fh = None
        self._sm: set[str] = set()
        self._covered = 0
        self._ckpt_loaded = False        # BAY §10-3: khong nap checkpoint 2 lan

    def log(self, msg: str) -> None:
        self._log(msg)
        if msg.startswith(("[ERR]", "[WARN]", "[STOP]", "[FATAL]")):
            try:
                with LOG_FILE.open("a", encoding="utf-8") as f:
                    f.write(f"{time.strftime('%Y-%m-%d %H:%M:%S')} {msg}\n")
            except OSError:
                pass

    # ---- resume: JSONL sidecar (append+flush tung record; chay lai = doc lai, bo qua da xong)
    def ckpt(self) -> pathlib.Path:
        h = hashlib.md5(self.key.encode("utf-8")).hexdigest()[:8]
        return HERE / f"_raw_{re.sub(r'[^A-Za-z0-9]+', '_', self.key)[:40]}_{h}.jsonl"

    def _load_ckpt(self) -> set[str]:
        """Tra ve tap ALIAS da xong. Mode ALL (keep=False) chi dung lai seen, khong nap RAM."""
        if self._ckpt_loaded:
            return {alias_from_url(r.get("URL", "")) for r in self.results if r.get("URL")}
        self._ckpt_loaded = True
        p = self.ckpt()
        seen: set[str] = set()
        if not p.exists():
            return seen
        for line in p.open(encoding="utf-8", errors="ignore"):
            try:
                r = json.loads(line)
            except Exception:
                continue                  # dong cuoi do dang khi bi kill -> bo
            if r.get("Status") == "Not Found":
                # MISS co the do loi mang tam (search that bai) -> lan sau PHAI THU LAI,
                # khong duoc coi la "da xong" roi bao khong co hang vinh vien
                continue
            self.rows += 1
            if r.get("Input"):
                self.done_inputs.add(norm_code(r["Input"]))
            if self.keep:
                self.results.append(r)
            if r.get("URL"):
                seen.add(alias_from_url(r["URL"]))
        if self.rows:
            self.log(f"[SYS] Resume: {self.rows} dong / {len(seen)} SP da xong -> bo qua")
        return seen

    def _emit(self, recs: list[dict]) -> None:
        for r in recs:
            self._fh.write(json.dumps(r, ensure_ascii=False) + "\n")
            if r.get("Input"):
                self.done_inputs.add(norm_code(r["Input"]))
        self._fh.flush()
        self.rows += len(recs)
        if self.keep:
            self.results.extend(recs)

    def iter_rows(self) -> Iterator[dict]:
        """Nguon cho export: RAM neu co, khong thi doc stream tu JSONL (mode ALL)."""
        if self.results or self.keep:
            yield from self.results
            return
        p = self.ckpt()
        if p.exists():
            for line in p.open(encoding="utf-8", errors="ignore"):
                try:
                    yield json.loads(line)
                except Exception:
                    continue

    def cleanup(self) -> None:
        """Xong sach -> xoa checkpoint. Bi Stop/crash -> GIU de lan sau resume."""
        if not self.key:
            return
        p = self.ckpt()
        if p.exists() and not self.stop_evt.is_set():
            p.unlink()

    def _full(self) -> bool:
        return bool(self.limit) and self.rows >= self.limit

    async def _gate(self) -> bool:
        while self.pause_evt.is_set() and not self.stop_evt.is_set():
            await asyncio.sleep(0.2)
        return not self.stop_evt.is_set() and not self._full()

    async def fetch(self, cli: httpx.AsyncClient, url: str, tries: int = 4) -> str | None:
        err = ""
        for i in range(tries):
            if not await self._gate():
                return None
            try:
                r = await cli.get(url)
                if r.status_code in (403, 429, 503):          # WAF / rate-limit -> lui dan
                    self.n_block += 1
                    wait = 120 if self.n_block >= HIBERNATE_429 else min(2 ** i * 5, 60)
                    self.log(f"[WARN] {r.status_code} x{self.n_block} -> ngu {wait}s | {url[-70:]}")
                    await asyncio.sleep(wait)
                    err = str(r.status_code)
                    continue
                self.n_block = 0
                if r.status_code in (404, 500):
                    # 500 = cham CAP OFFSET 30.000 (BAY #1) hoac alias sai; khong phai loi mang
                    self.fails = 0
                    return None
                r.raise_for_status()
                self.fails = 0
                return r.text
            except (httpx.HTTPError, OSError) as e:
                err = type(e).__name__
                if i < tries - 1:
                    await asyncio.sleep(min(2 ** i, 20))
        self.fails += 1
        self.log(f"[ERR] {err} | {url[-70:]}")
        if self.fails >= CIRCUIT_BREAK:
            self.log(f"[STOP] {CIRCUIT_BREAK} loi LIEN TIEP - nghi bi chan IP. DUNG.")
            self.stop_evt.set()
        return None

    async def fetch_json(self, cli: httpx.AsyncClient, url: str) -> Any:
        t = await self.fetch(cli, url)
        if not t:
            return None
        try:
            return json.loads(t)
        except ValueError:                # redirect ra trang HTML -> coi nhu khong co
            return None

    async def _pages(self, cli: httpx.AsyncClient, path: str, field: str):
        """Async generator tung trang (khong om ca collection trong RAM).
        BAY #2: KHONG dung khi len(page) < limit. Chi dung khi rong / stuck / cham cap."""
        seen_ids: set[Any] = set()
        for page in range(1, MAX_PAGES + 1):
            if not await self._gate():
                return
            if page * PAGE_SIZE > OFFSET_CAP:
                self.log(f"[WARN] {path}: cham CAP OFFSET {OFFSET_CAP} -> sitemap se va phan con lai")
                return
            d = await self.fetch_json(cli, f"{BASE}{path}?limit={PAGE_SIZE}&page={page}")
            if d is None and page > 1:
                # loi tam KHAC het trang -> log, khong im lang bo phan con lai cua collection
                self.log(f"[WARN] dung phan trang o trang {page} vi tai loi -> CO THE SOT SP "
                         f"| {path[-50:]}")
            items = (d or {}).get(field) or []
            new = [x for x in items if x.get("id") not in seen_ids]
            if not new:                                  # rong hoac stuck-page
                return
            seen_ids.update(x.get("id") for x in new)
            yield new
            if self.delay:
                await asyncio.sleep(self.delay)

    async def _sitemap_aliases(self, cli: httpx.AsyncClient) -> set[str]:
        """sitemap.xml -> file con doc THAT tu sitemapindex (BAY §10-9, khong doan day so)."""
        idx = await self.fetch(cli, f"{BASE}/sitemap.xml")
        if not idx:
            return set()
        subs = [u for u in LOC_RE.findall(idx) if "sitemap_products" in u]
        sem = asyncio.Semaphore(self.conc)

        async def one(u: str) -> set[str]:
            async with sem:
                x = await self.fetch(cli, u)
                return {alias_from_url(v) for v in LOC_RE.findall(x)} if x else set()

        out: set[str] = set()
        for part in await asyncio.gather(*(one(s) for s in subs)):
            out |= part
        out.discard("")
        return out

    def _take(self, products: list[dict], seen: set[str], inputs: dict[str, str] | None = None) -> None:
        """Ghi record tu dict SP DA CO SAN (list JSON du du lieu -> khong goi detail)."""
        for p in products:
            a = p.get("alias") or ""
            if not a or a in seen:
                continue
            seen.add(a)
            if a in self._sm:
                self._covered += 1
            recs = parse_product(p, (inputs or {}).get(a, ""))
            if recs:
                self._emit(recs)
            self.done += 1
            if self.done % 2000 == 0:
                self.log(f"[OK ] {self.done}/{self.total or '?'} SP | {self.rows} dong")
            if self.done % GC_EVERY == 0:
                gc.collect()
            if self._full():
                return

    async def _crawl_aliases(self, cli: httpx.AsyncClient, aliases: list[str], seen: set[str],
                             inputs: dict[str, str] | None = None, nf: bool = False) -> None:
        """Fetch /products/<alias>.json tung alias (link / keyword / va SP sot).
        nf=True: alias miss van ghi dong 'Not Found' (mode link/excel)."""
        todo = [a for a in dict.fromkeys(aliases) if a and a not in seen][:MAX_ITEMS]
        self.total += len(todo)
        sem = asyncio.Semaphore(self.conc)

        async def one(a: str) -> None:
            async with sem:
                if not await self._gate():           # re-check SAU khi gianh slot
                    return
                d = await self.fetch_json(cli, product_json_url(a))
                if self.delay:
                    await asyncio.sleep(self.delay)
                if not await self._gate() and not d:
                    return                           # bi Stop giua chung -> khong ghi Not Found oan
                recs = parse_product(d, (inputs or {}).get(a, "")) if d else []
                if recs:
                    seen.add(a)
                    if a in self._sm:
                        self._covered += 1
                    self._emit(recs)
                elif nf:
                    self._emit([not_found((inputs or {}).get(a, a))])
                self.done += 1
                if not self.keep and self.done % 500:
                    return
                self.log(f"[{'OK ' if recs else 'MISS'}] {self.done}/{self.total} +{len(recs)} | {a[:58]}")

        try:
            await asyncio.gather(*(one(a) for a in todo))
        except (KeyboardInterrupt, Exception) as e:   # crash-proof: record da o trong JSONL
            self.log(f"[FATAL] {type(e).__name__}: {e}")

    async def search_json(self, cli: httpx.AsyncClient, term: str) -> list[tuple[str, str]]:
        """/search?q=&view=json (4 KQ/trang) -> [(alias, title)], toi da SEARCH_JSON_PAGES trang."""
        out: list[tuple[str, str]] = []
        seen: set[str] = set()
        for page in range(1, SEARCH_JSON_PAGES + 1):
            if not await self._gate():
                break
            raw = await self.fetch(cli, f"{BASE}/search?q={quote(term)}&view=json&page={page}") or ""
            try:
                d = json.loads(raw)
                pairs = [(r.get("title") or "", r.get("url") or "") for r in d.get("results") or []
                         if isinstance(r, dict)]
                total = int(d.get("results_total") or 0)
            except (ValueError, AttributeError):
                # BAY: theme KHONG escape dau " trong title (TRUSCO Wire Rope Sling "Aluminum Lock")
                # -> JSON hong -> json.loads chet -> ma co that bi bao Not Found oan. Vot bang regex.
                pairs = SEARCH_JSON_ITEM_RE.findall(raw)
                m = re.search(r'"results_total"\s*:\s*(\d+)', raw)
                total = int(m.group(1)) if m else 0
            new = [(alias_from_url(u), htmlmod.unescape(t)) for t, u in pairs]
            new = [(a, t) for a, t in new if a and a not in seen]
            if not new:                                   # het / stuck
                break
            seen.update(a for a, _ in new)
            out += new
            if len(out) >= total:
                break
        return out

    async def search_html(self, cli: httpx.AsyncClient, term: str, limit: int = 0) -> list[str]:
        """/search?q= HTML 20 SP/trang -> alias trong luoi ket qua (title-head..pagination)."""
        out: list[str] = []
        seen: set[str] = set()
        total = None
        for page in range(1, MAX_PAGES + 1):
            if not await self._gate():
                break
            h = await self.fetch(cli, f"{BASE}/search?q={quote(term)}&page={page}")
            if not h:
                break
            if total is None:
                m = SEARCH_TOTAL_RE.search(h)
                total = int(m.group(1)) if m else -1
                self.log(f"[SYS] '{term}': site bao {total} ket qua")
            i = h.find("title-head")
            if i < 0:
                break
            j = h.find("pagination", i)
            seg = h[i:j if j > 0 else len(h)]
            new = [a for a in dict.fromkeys(SEARCH_ITEM_RE.findall(seg)) if a not in seen]
            if not new:                                   # stuck-page = het ket qua
                break
            seen.update(new)
            out.extend(new)
            if limit and len(out) >= limit:
                return out[:limit]
            if total and total > 0 and len(out) >= total:
                break
            if self.delay:
                await asyncio.sleep(self.delay)
        return out[:MAX_ITEMS]

    async def run(self, mode: str, target: str = "", excel: str = "", sheet: str = "",
                  col: str = "", limit: int = 0) -> None:
        self.key = f"{mode}_{target or excel}_{sheet}_{col}"
        if mode == "all":
            self.keep, self.limit = False, limit
        async with httpx.AsyncClient(headers=HEADERS, timeout=TIMEOUT, follow_redirects=True,
                                     max_redirects=MAX_REDIRECTS,
                                     limits=httpx.Limits(max_connections=self.conc + 4)) as cli:
            seen = self._load_ckpt()
            self._fh = self.ckpt().open("a", encoding="utf-8")
            try:
                if mode == "all":
                    await self._run_all(cli, seen)

                elif mode in ("code", "excel"):
                    raw = [norm_code(target)] if mode == "code" else read_excel_col(excel, sheet, col)
                    raw = [c for c in dict.fromkeys(raw) if c][:MAX_ITEMS]
                    links = [c for c in raw if c.lower().startswith("http")]
                    codes = [c for c in raw if c not in links]
                    # BAY §10-2: LINK crawl RIENG, KHONG ap bo loc `codes`
                    if links:
                        todo = [u for u in links if norm_code(u) not in self.done_inputs]
                        al = [alias_from_url(u) for u in todo]
                        await self._crawl_aliases(cli, al, set(), inputs=dict(zip(al, todo)), nf=True)
                    if codes:
                        await self._run_codes(cli, codes)

                elif mode == "keyword":
                    al = await self.search_html(cli, target, limit)
                    self.log(f"[SYS] '{target}': lay {len(al)} san pham")
                    await self._crawl_aliases(cli, al, seen, inputs={a: target for a in al})

                elif mode == "link":
                    links = [u.strip() for u in target.splitlines() if u.strip().lower().startswith("http")]
                    todo = [u for u in dict.fromkeys(links) if norm_code(u) not in self.done_inputs]
                    al = [alias_from_url(u) for u in todo]
                    await self._crawl_aliases(cli, al, set(), inputs=dict(zip(al, todo)), nf=True)

                else:
                    self.log(f"[ERR] mode khong hop le: {mode}")
            finally:
                self._fh.close()

    async def _run_all(self, cli: httpx.AsyncClient, seen: set[str]) -> None:
        """CAP OFFSET 30.000 (BAY #1) => KHONG paginate /products.json.
        sitemap (chuan vang) -> quet collection (uu tien danh muc goc htgoods-*) -> DUNG SOM khi
        da phu du sitemap -> alias con thieu fetch le /products/<alias>.json."""
        self._sm = await self._sitemap_aliases(cli)
        self._covered = len(seen & self._sm)
        self.total = len(self._sm)
        self.done = len(seen)
        self.log(f"[SYS] Sitemap: {len(self._sm)} SP | da co {self._covered} (resume)")

        def covered() -> bool:
            return bool(self._sm) and self._covered >= len(self._sm)

        cols: list[str] = []
        async for items in self._pages(cli, "/collections.json", "collections"):
            cols += [c["alias"] for c in items if c.get("alias")]
        cols.sort(key=lambda a: (not a.startswith("htgoods-"), a))
        self.log(f"[SYS] {len(cols)} collection -> quet (pha CAP {OFFSET_CAP})")
        sem = asyncio.Semaphore(self.conc)
        ncol = 0

        async def one(alias: str) -> None:
            nonlocal ncol
            async with sem:
                if covered() or not await self._gate():
                    return
                async for items in self._pages(cli, f"/collections/{quote(alias)}/products.json", "products"):
                    self._take(items, seen)          # sync, khong await -> khong tranh chap
                    if covered() or self._full():
                        break
                ncol += 1
                self.log(f"[SYS] collection {ncol}/{len(cols)} {alias[:45]} | "
                         f"phu {self._covered}/{len(self._sm)} | {self.rows} dong")

        if not covered():
            try:
                await asyncio.gather(*(one(a) for a in cols))
            except (KeyboardInterrupt, Exception) as e:
                self.log(f"[FATAL] {type(e).__name__}: {e}")
        if not await self._gate():
            return
        missing = sorted(a for a in self._sm if a not in seen)
        self.log(f"[SYS] Collection xong ({ncol} quet): {len(seen)} SP -> con thieu {len(missing)} "
                 f"so voi sitemap -> va truc tiep")
        if missing:
            await self._crawl_aliases(cli, missing, seen)
        self.log(f"[SYS] ALL xong: {len(seen)} SP | {self.rows} dong")

    async def _run_codes(self, cli: httpx.AsyncClient, codes: list[str]) -> None:
        """1 ma -> search view=json -> uu tien ung vien co ma trong title -> detail JSON
        -> KHOP MA CHINH XAC (§4.3). Ma miss van co dong 'Not Found'.
        Dedup/resume theo MA NHAP (BAY §10-6): nhieu ma (MW + SKU) cung tro 1 SP -> moi ma
        van co dong rieng; product JSON cache 1 lan."""
        todo = [c for c in codes if norm_code(c) not in self.done_inputs]
        if len(todo) < len(codes):
            self.log(f"[SYS] Resume: bo qua {len(codes) - len(todo)} ma da xong")
        self.total += len(todo)
        sem = asyncio.Semaphore(self.conc)
        cache: dict[str, Any] = {}

        async def get_product(a: str) -> Any:
            if a in cache:
                return cache[a]
            d = await self.fetch_json(cli, product_json_url(a))
            if len(cache) >= CACHE_MAX:
                cache.clear()
            cache[a] = d
            return d

        async def one(c: str) -> None:
            async with sem:
                if not await self._gate():
                    return
                hits = await self.search_json(cli, c)
                if self.delay:
                    await asyncio.sleep(self.delay)
                ca = alnum(c)
                pri = [a for a, t in hits if ca and (ca in alnum(t) or ca in alnum(a))]
                recs: list[dict] = []
                for group in (pri, [a for a, _ in hits if a not in pri]):
                    for a in group:
                        recs += match_code(parse_product(await get_product(a), c), c)
                    if any(r["Status"] == "OK" for r in recs):
                        break                          # da khop exact -> khoi mo nhom con lai
                if self.stop_evt.is_set():
                    return                             # Stop giua chung -> khong ghi Not Found oan
                if any(r["Status"] == "OK" for r in recs):
                    recs = [r for r in recs if r["Status"] == "OK"]
                recs = list({r.get("URL"): r for r in recs}.values()) or [not_found(c)]
                self._emit(recs)
                self.done += 1
                ok = recs[0].get("Status") != "Not Found"
                self.log(f"[{'OK ' if ok else 'MISS'}] {self.done}/{self.total} "
                         f"+{len(recs) if ok else 0} | {c} ({len(hits)} hit)")

        try:
            await asyncio.gather(*(one(c) for c in todo))
        except (KeyboardInterrupt, Exception) as e:   # crash-proof: record da o trong JSONL
            self.log(f"[FATAL] {type(e).__name__}: {e}")


def read_excel_col(path: str, sheet: str, col: str) -> list[str]:
    """Loi -> ValueError (KHONG SystemExit: worker GUI se khong bat duoc -> treo)."""
    if not path:
        raise ValueError("Chua chon file Excel")
    df = pd.read_excel(path, sheet_name=sheet or 0, dtype=str)
    df.columns = df.columns.str.strip()
    if col not in df.columns:
        raise ValueError(f"Khong thay cot '{col}'. Cot co san: {list(df.columns)}")
    return [norm_code(v) for v in df[col].dropna().tolist()]


# ================================================================ GUI
def run_gui() -> None:
    import sys
    import tkinter as tk
    from tkinter import messagebox, ttk

    # ha tang GUI chung o root (di nguoc len cay thu muc -> song sot khi doi nhom folder)
    root_dir = next(p for p in (HERE, *HERE.parents) if (p / "gui_common.py").exists())
    sys.path.insert(0, str(root_dir))
    from gui_common import ask_export_path, attach_excel_picker

    root = tk.Tk()
    root.title("HT GOODS Crawler — htgoods.com.vn")
    root.geometry("1080x700")
    ttk.Style().theme_use("clam")

    q: queue.Queue[str] = queue.Queue()
    stop_evt, pause_evt = threading.Event(), threading.Event()
    state: dict[str, Any] = {"running": False, "eng": None}

    left = ttk.Frame(root, padding=10)
    left.pack(side="left", fill="y")
    right = ttk.Frame(root, padding=(0, 10, 10, 10))
    right.pack(side="right", fill="both", expand=True)

    mode = tk.StringVar(value="code")
    ttk.Label(left, text="Pipeline", font=("Segoe UI", 10, "bold")).pack(anchor="w")
    for val, txt in [("code", "1. Theo MA (MW / ma hang / model)"),
                     ("keyword", "2. Theo TU KHOA"),
                     ("link", "3. Theo LINK (moi dong 1 link)"),
                     ("excel", "4. Theo cot EXCEL (ma hoac link)"),
                     ("all", "5. Crawl ALL (~100.7k SP)")]:
        ttk.Radiobutton(left, text=txt, value=val, variable=mode).pack(anchor="w")

    ttk.Label(left, text="\nMa / Tu khoa:").pack(anchor="w")
    ent = ttk.Entry(left, width=36)
    ent.pack(anchor="w")
    ttk.Label(left, text="Link (moi dong 1 link):").pack(anchor="w")
    box = tk.Text(left, width=36, height=5)
    box.pack(anchor="w")

    xls, sheet, col = tk.StringVar(), tk.StringVar(), tk.StringVar()
    outdir = tk.StringVar(value="images")
    attach_excel_picker(left, xls, sheet, col, outdir)

    conc, delay, limit = tk.IntVar(value=8), tk.DoubleVar(value=0.05), tk.IntVar(value=0)
    for lbl, var in [("Concurrency:", conc), ("Delay (s):", delay), ("Limit (0=all):", limit)]:
        f3 = ttk.Frame(left)
        f3.pack(anchor="w", pady=1)
        ttk.Label(f3, text=lbl, width=13).pack(side="left")
        ttk.Entry(f3, textvariable=var, width=8).pack(side="left")

    console = tk.Text(right, bg="#10151b", fg="#d6e2ee", insertbackground="#fff", wrap="none")
    console.pack(fill="both", expand=True)
    pb = ttk.Progressbar(right, mode="determinate")
    pb.pack(fill="x", pady=4)
    stat = ttk.Label(right, text="0/0")
    stat.pack(anchor="w")

    def drain() -> None:
        for _ in range(500):
            if q.empty():
                break
            console.insert("end", q.get() + "\n")
        console.see("end")
        if int(console.index("end-1c").split(".")[0]) > 5000:      # console khong phinh RAM
            console.delete("1.0", "2000.0")
        eng: Engine | None = state["eng"]
        if eng and eng.total:
            pb["maximum"], pb["value"] = eng.total, min(eng.done, eng.total)
            stat.config(text=f"{eng.done}/{eng.total} SP | {eng.rows} dong")
        root.after(100, drain)

    def worker() -> None:
        m = mode.get()
        target = box.get("1.0", "end").strip() if m == "link" else ent.get().strip()
        eng = Engine(conc.get(), delay.get(), q.put, stop_evt, pause_evt)
        state["eng"] = eng
        try:
            asyncio.run(eng.run(m, target, xls.get(), sheet.get(), col.get(), limit.get()))
        except BaseException as e:      # BaseException: SystemExit/KeyboardInterrupt cung phai bat
            q.put(f"[FATAL] {type(e).__name__}: {e}")
        finally:                        # finally: LUON nha co, khong khoa START vinh vien
            state["running"] = False
            q.put(f"[SYS] XONG — {eng.rows} dong. Bam EXPORT de xuat Excel.")

    def start() -> None:
        if state["running"]:
            messagebox.showinfo("Dang chay", "Dang crawl. Bam Stop truoc.")
            return
        m = mode.get()
        if m == "excel" and not (xls.get() and col.get()):
            messagebox.showwarning("Thieu du lieu", "Chon file Excel va cot truoc.")
            return
        if m in ("code", "keyword") and not ent.get().strip():
            messagebox.showwarning("Thieu du lieu", "Nhap ma / tu khoa truoc.")
            return
        if m == "link" and not box.get("1.0", "end").strip():
            messagebox.showwarning("Thieu du lieu", "Dan link vao o (moi dong 1 link).")
            return
        stop_evt.clear()
        pause_evt.clear()
        state["running"] = True
        console.delete("1.0", "end")
        threading.Thread(target=worker, daemon=True).start()

    def export() -> None:
        eng: Engine | None = state["eng"]
        if state["running"]:
            messagebox.showinfo("Export", "Dang crawl. Bam Stop (hoac doi xong) roi EXPORT.")
            return
        if not eng or not eng.rows:
            messagebox.showinfo("Export", "Chua co du lieu")
            return
        p = ask_export_path(xls.get(), HERE, "htgoods.xlsx")
        if p:
            n = export_excel(eng.iter_rows(), pathlib.Path(p))
            eng.cleanup()               # da xuat xong -> xoa checkpoint (bi Stop thi GIU)
            messagebox.showinfo("Export", f"Da luu {n} dong:\n{p}")

    fb = ttk.Frame(left)
    fb.pack(anchor="w", pady=8)
    ttk.Button(fb, text="START", command=start).pack(side="left")
    ttk.Button(fb, text="Pause", command=lambda: (pause_evt.clear() if pause_evt.is_set()
                                                  else pause_evt.set())).pack(side="left")
    ttk.Button(fb, text="Stop", command=stop_evt.set).pack(side="left")
    ttk.Button(fb, text="EXPORT", command=export).pack(side="left")

    def on_close() -> None:
        if state["running"] and not messagebox.askyesno("Thoat", "Dang crawl. Thoat va tat luon?"):
            return
        stop_evt.set()
        try:
            root.destroy()
        finally:
            os._exit(0)      # kill cung: khong de thread asyncio/console treo

    root.protocol("WM_DELETE_WINDOW", on_close)
    drain()
    root.mainloop()


# ================================================================ CLI
def main() -> None:
    ap = argparse.ArgumentParser(description="HT GOODS crawler (htgoods.com.vn)")
    ap.add_argument("--mode", choices=["code", "keyword", "link", "excel", "all"])
    ap.add_argument("--code")
    ap.add_argument("--keyword")
    ap.add_argument("--link")
    ap.add_argument("--excel")
    ap.add_argument("--sheet", default="")
    ap.add_argument("--col", default="")
    ap.add_argument("--out", default="htgoods.xlsx")
    ap.add_argument("--conc", type=int, default=8)
    ap.add_argument("--delay", type=float, default=0.05)
    ap.add_argument("--limit", type=int, default=0)
    a = ap.parse_args()

    if not a.mode:
        run_gui()
        return

    target = a.code or a.keyword or a.link or ""
    eng = Engine(a.conc, a.delay)
    t0 = time.time()
    try:
        asyncio.run(eng.run(a.mode, target, a.excel or "", a.sheet, a.col, a.limit))
    except KeyboardInterrupt:
        eng.stop_evt.set()               # Ctrl+C -> GIU checkpoint de chay lai la resume
        print("[STOP] Ctrl+C -> giu checkpoint, chay lai de resume")
    finally:
        if eng.rows:
            n = export_excel(eng.iter_rows(), pathlib.Path(a.out))
            print(f"[SYS] {n} dong -> {a.out} ({time.time() - t0:.1f}s)")
        eng.cleanup()


if __name__ == "__main__":
    main()
