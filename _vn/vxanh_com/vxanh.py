"""
vxanh.com - custom SSR crawler (httpx + selectolax, no browser).

RECON (2026-08-11)
------------------
- Nen tang: custom SSR app kieu Laravel session/XSRF, khong phai WordPress/Sapo/Haravan.
- Khong thay public JSON product API; search va PDP deu tra HTML SSR.
- URL san pham: /san-pham/<slug>-<token>.
- Search/list: /tim-san-pham?q=<term>&page=N, 24 card/trang; khong q hien catalog live.
- Live list hien 24,738 URL qua 1,031 trang; sitemap hien 24,734 URL -> mode all phai union list+sitemap.
- PDP field o .product-title, .product-codes .code-chip, .product-facts .product-fact,
  table.spec-table, .product-description, .product-gallery img.
- Exact code phai so ca Ma V Xanh va Ma tham khao; khong duoc an hit dau tien.
"""
from __future__ import annotations

import argparse
import asyncio
import html
import json
import math
import os
import pathlib
import queue
import re
import sqlite3
import sys
import threading
import time
from typing import Any
from urllib.parse import quote, urljoin, urlsplit, urlunsplit

import httpx
import pandas as pd
from selectolax.lexbor import LexborHTMLParser

HERE = pathlib.Path(__file__).parent
BASE = "https://vxanh.com"
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")
HEADERS = {"User-Agent": UA, "Accept-Language": "vi-VN,vi;q=0.9,en;q=0.8"}

SEARCH_PER_PAGE = 24
CIRCUIT_BREAK = 25
MAX_ITEMS = 30000
MAX_PAGES = 1200
MAX_REDIRECTS = 5
TIMEOUT = 45.0

LOC_RE = re.compile(r"<loc>([^<]+)</loc>", re.I)


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", html.unescape(str(value or "")).replace("\xa0", " ")).strip()


def norm_code(value: Any) -> str:
    v = re.sub(r"\s+", "", str(value or ""))
    return v if v.lower().startswith("http") else v.upper()


def canonical_product_url(url: str) -> str:
    if not url:
        return ""
    abs_url = urljoin(BASE + "/", str(url).strip())
    parts = urlsplit(abs_url)
    path = re.sub(r"/+", "/", parts.path or "/").rstrip("/")
    if not path.startswith("/san-pham/"):
        return ""
    return urlunsplit(("https", "vxanh.com", path, "", ""))


def search_url(term: str, page: int = 1) -> str:
    if term:
        return f"{BASE}/tim-san-pham?q={quote(term)}&page={page}"
    return f"{BASE}/tim-san-pham?page={page}"


def exact_code_match(product_row: dict[str, Any], code: str) -> bool:
    code_norm = norm_code(code)
    patt = re.compile(rf"(?<![A-Z0-9]){re.escape(code_norm)}(?![A-Z0-9])")
    for key in ("Ma V Xanh", "Ma tham khao"):
        if patt.search(norm_code(product_row.get(key, ""))):
            return True
    return False


def status_from_product(row: dict[str, Any]) -> str:
    title = clean_text(row.get("Page"))
    code_a = norm_code(row.get("Ma V Xanh"))
    code_b = norm_code(row.get("Ma tham khao"))
    valid_codes = [c for c in (code_a, code_b) if c and c != "N/A"]
    fact_hits = [clean_text(row.get("Thuong hieu")), clean_text(row.get("Ton kho")), clean_text(row.get("Bo sung hang"))]
    fact_hits = [x for x in fact_hits if x and x != "N/A"]
    if not title:
        return "Error"
    if not valid_codes:
        return "Incomplete"
    if not fact_hits and not row.get("_spec") and not clean_text(row.get("Mo ta")):
        return "Incomplete"
    return "OK"


def parse_search_page(html_text: str, input_value: str = "") -> tuple[list[dict[str, str]], int]:
    tree = LexborHTMLParser(html_text or "")
    total_text = clean_text(tree.css_first(".result-count strong").text()) if tree.css_first(".result-count strong") else ""
    total_count = int(re.sub(r"[^\d]", "", total_text) or "0")
    total_pages = math.ceil(total_count / SEARCH_PER_PAGE) if total_count else 0
    rows: list[dict[str, str]] = []
    seen: set[str] = set()
    for node in tree.css('a[href*="/san-pham/"]'):
        url = canonical_product_url(node.attributes.get("href", ""))
        if not url or url in seen:
            continue
        seen.add(url)
        rows.append({"Input": input_value, "URL": url})
    return rows, total_pages


def parse_product_page(html_text: str, url: str, input_value: str = "") -> dict[str, Any]:
    tree = LexborHTMLParser(html_text or "")
    title = clean_text(tree.css_first(".product-title, h1").text()) if tree.css_first(".product-title, h1") else ""

    codes: dict[str, str] = {}
    for node in tree.css(".product-codes .code-chip"):
        raw = clean_text(node.text(separator=" ", strip=True))
        code_node = node.css_first("strong")
        val = clean_text(code_node.text()) if code_node else clean_text(raw.split()[-1] if raw else "")
        key = " ".join(raw.replace(val, "").split()).rstrip(":")
        if "V Xanh" in key:
            codes["Ma V Xanh"] = val
        elif "tham khảo" in key or "tham khao" in key:
            codes["Ma tham khao"] = val

    facts: dict[str, str] = {}
    for node in tree.css(".product-facts .product-fact"):
        key_node = node.css_first(".fact-label")
        val_node = node.css_first(".fact-value")
        key = clean_text(key_node.text()) if key_node else ""
        val = clean_text(val_node.text()) if val_node else ""
        if key and val:
            facts[key] = val

    spec: dict[str, str] = {}
    for tr in tree.css("table.spec-table tr"):
        th = tr.css_first("th")
        td = tr.css_first("td")
        key = clean_text(th.text()) if th else ""
        val = clean_text(td.text(separator=" ", strip=True)) if td else ""
        if key and val and key not in spec:
            spec[key] = val

    cats = []
    for node in tree.css(".breadcrumb-item"):
        txt = clean_text(node.text())
        href = node.attributes.get("href", "")
        if not txt or txt.lower() == "trang chủ" or href == "/":
            continue
        cats.append(txt)
    cats = list(dict.fromkeys(cats))

    images = []
    for node in tree.css(".product-gallery img"):
        src = clean_text(node.attributes.get("src") or node.attributes.get("data-src"))
        if not src:
            continue
        images.append(urljoin(BASE + "/", src))
    images = list(dict.fromkeys(images))

    desc_node = tree.css_first(".product-description")
    desc = clean_text(desc_node.text(separator=" ", strip=True)) if desc_node else ""

    row = {
        "Input": input_value,
        "Page": title,
        "Item No": clean_text(codes.get("Ma V Xanh") or codes.get("Ma tham khao")),
        "Status": "",
        "Ma V Xanh": clean_text(codes.get("Ma V Xanh")),
        "Ma tham khao": clean_text(codes.get("Ma tham khao")),
        "Thuong hieu": clean_text(facts.get("Thương hiệu") or facts.get("Thuong hieu")),
        "Gia tham khao": clean_text(tree.css_first(".quote-price").text()) if tree.css_first(".quote-price") else "",
        "Gia niem yet": clean_text(tree.css_first(".quote-price-old").text()) if tree.css_first(".quote-price-old") else "",
        "Giam gia": clean_text(tree.css_first(".quote-discount").text()) if tree.css_first(".quote-discount") else "",
        "Ton kho": clean_text(facts.get("Tồn kho") or facts.get("Ton kho")),
        "Bo sung hang": clean_text(facts.get("Bổ sung hàng") or facts.get("Bo sung hang")),
        "_cats": cats,
        "_spec": spec,
        "Mo ta": desc,
        "Image": " | ".join(images),
        "URL": canonical_product_url(url),
    }
    row["Status"] = status_from_product(row)
    return row


def not_found(input_value: str, status: str = "Not Found") -> dict[str, Any]:
    return {
        "Input": input_value,
        "Page": "",
        "Item No": input_value,
        "Status": status,
        "Ma V Xanh": "",
        "Ma tham khao": "",
        "Thuong hieu": "",
        "Gia tham khao": "",
        "Gia niem yet": "",
        "Giam gia": "",
        "Ton kho": "",
        "Bo sung hang": "",
        "_cats": [],
        "_spec": {},
        "Mo ta": "",
        "Image": "",
        "URL": "",
    }


def export_excel(rows: list[dict[str, Any]], path: pathlib.Path) -> pathlib.Path:
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    fixed = [
        "Input", "Page", "Item No", "Status", "Ma V Xanh", "Ma tham khao", "Thuong hieu",
        "Gia tham khao", "Gia niem yet", "Giam gia", "Ton kho", "Bo sung hang",
    ]
    max_cat = max((len(row.get("_cats") or []) for row in rows), default=0)
    data = []
    for row in rows:
        item = {k: row.get(k, "") for k in fixed}
        cats = row.get("_cats") or []
        for i in range(max_cat):
            item[f"Category {i + 1}"] = cats[i] if i < len(cats) else ""
        item["Thong so"] = "\n".join(f"{k}: {v}" for k, v in (row.get("_spec") or {}).items())
        item["Mo ta"] = row.get("Mo ta", "")
        item["Image"] = row.get("Image", "")
        item["URL"] = row.get("URL", "")
        data.append(item)
    df = pd.DataFrame(data)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False)
        ws = writer.sheets["Sheet1"]
        align = Alignment(wrap_text=True, vertical="top")
        for name, width in [("Thong so", 60), ("Mo ta", 80), ("Image", 60), ("URL", 50)]:
            if name not in df.columns:
                continue
            col = get_column_letter(df.columns.get_loc(name) + 1)
            ws.column_dimensions[col].width = width
            for cell in ws[col][1:]:
                cell.alignment = align
    return path


def read_excel_col(path: str, sheet: str, col: str) -> list[str]:
    if not path:
        raise ValueError("Chua chon file Excel")
    df = pd.read_excel(path, sheet_name=sheet or 0, dtype=str)
    df.columns = df.columns.map(str)
    if col not in df.columns:
        raise ValueError(f"Khong thay cot '{col}'")
    return [str(v).strip() for v in df[col].dropna().tolist() if str(v).strip()]


class Engine:
    def __init__(
        self,
        conc: int = 6,
        delay: float = 0.05,
        log=print,
        stop_evt: threading.Event | None = None,
        pause_evt: threading.Event | None = None,
    ) -> None:
        self.conc = max(1, conc)
        self.delay = max(0.0, delay)
        self.log = log
        self.stop_evt = stop_evt or threading.Event()
        self.pause_evt = pause_evt or threading.Event()
        self.results: list[dict[str, Any]] = []
        self.done = 0
        self.total = 0
        self.fails = 0
        self.key = ""
        self.reused_inputs: set[str] = set()
        self.reused_urls: set[str] = set()

    def db_path(self, key: str) -> pathlib.Path:
        safe = re.sub(r"[^A-Za-z0-9]+", "_", key)[:60]
        return HERE / f"_resume_{safe}.sqlite3"

    def init_db(self, key: str) -> pathlib.Path:
        path = self.db_path(key)
        conn = sqlite3.connect(path)
        try:
            conn.execute("PRAGMA journal_mode=WAL")
            conn.execute("PRAGMA synchronous=OFF")
            conn.execute(
                "CREATE TABLE IF NOT EXISTS work (url TEXT PRIMARY KEY, input_value TEXT, source TEXT, crawled INTEGER DEFAULT 0, last_error TEXT)"
            )
            conn.execute(
                "CREATE TABLE IF NOT EXISTS products (url TEXT PRIMARY KEY, payload_json TEXT NOT NULL, status TEXT NOT NULL)"
            )
            conn.execute("CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY, value TEXT)")
            conn.commit()
        finally:
            conn.close()
        return path

    def save_row(self, key: str, row: dict[str, Any]) -> None:
        path = self.init_db(key)
        row_key = canonical_product_url(row.get("URL", "")) or f"input:{norm_code(row.get('Input') or row.get('Item No') or time.time())}"
        payload = dict(row)
        payload["URL"] = canonical_product_url(payload.get("URL", "")) or payload.get("URL", "")
        conn = sqlite3.connect(path)
        try:
            conn.execute(
                "INSERT OR REPLACE INTO products(url, payload_json, status) VALUES (?, ?, ?)",
                (row_key, json.dumps(payload, ensure_ascii=False), payload.get("Status", "")),
            )
            conn.execute(
                "INSERT OR REPLACE INTO work(url, input_value, source, crawled, last_error) VALUES (?, ?, COALESCE((SELECT source FROM work WHERE url=?), ''), ?, ?)",
                (row_key, payload.get("Input", ""), row_key, 1 if payload.get("Status") in {"OK", "Incomplete"} else 0,
                 "" if payload.get("Status") in {"OK", "Incomplete"} else payload.get("Status", "")),
            )
            conn.commit()
        finally:
            conn.close()

    def load_resume_rows(self, key: str) -> list[dict[str, Any]]:
        path = self.init_db(key)
        conn = sqlite3.connect(path)
        try:
            rows = []
            for payload_json, status in conn.execute(
                "SELECT payload_json, status FROM products WHERE status IN ('OK', 'Incomplete') ORDER BY json_extract(payload_json, '$.Input')"
            ):
                row = json.loads(payload_json)
                rows.append(row)
            return rows
        finally:
            conn.close()

    def cleanup(self, exported: bool = False) -> None:
        path = self.db_path(self.key) if self.key else None
        if path and path.exists() and exported and not self.stop_evt.is_set():
            path.unlink()

    async def _gate(self) -> bool:
        while self.pause_evt.is_set() and not self.stop_evt.is_set():
            await asyncio.sleep(0.2)
        return not self.stop_evt.is_set()

    async def fetch_text(self, client: httpx.AsyncClient, url: str) -> tuple[str | None, str]:
        for i in range(4):
            if not await self._gate():
                return None, url
            try:
                resp = await client.get(url)
                final_url = str(resp.url)
                if resp.status_code in (403, 429, 503, 520, 521, 522):
                    wait = min(5 * (2 ** i), 60)
                    self.log(f"[WARN] {resp.status_code} -> ngu {wait}s | {url[-80:]}")
                    await asyncio.sleep(wait)
                    continue
                if resp.status_code == 404:
                    return None, final_url
                resp.raise_for_status()
                self.fails = 0
                return resp.text, final_url
            except Exception as exc:
                if i < 3:
                    await asyncio.sleep(min(2 ** i, 20))
                    continue
                self.fails += 1
                self.log(f"[ERR] {type(exc).__name__} | {url[-80:]}")
                if self.fails >= CIRCUIT_BREAK:
                    self.log(f"[STOP] {CIRCUIT_BREAK} loi lien tiep - dung.")
                    self.stop_evt.set()
                return None, url
        return None, url

    async def discover_list_urls(self, client: httpx.AsyncClient, limit: int = 0) -> list[str]:
        urls: list[str] = []
        seen: set[str] = set()
        page = 1
        while page <= MAX_PAGES:
            html_text, _ = await self.fetch_text(client, search_url("", page))
            if not html_text:
                break
            rows, _total_pages = parse_search_page(html_text)
            if not rows:
                break
            for row in rows:
                url = canonical_product_url(row["URL"])
                if url and url not in seen:
                    seen.add(url)
                    urls.append(url)
            if limit and len(urls) >= limit:
                return urls[:limit]
            page += 1
            if self.delay:
                await asyncio.sleep(self.delay)
        return urls

    async def discover_sitemap_urls(self, client: httpx.AsyncClient) -> list[str]:
        xml, _ = await self.fetch_text(client, f"{BASE}/sitemap.xml")
        if not xml:
            return []
        submaps = [u for u in LOC_RE.findall(xml) if "sitemap-products" in u]
        urls: list[str] = []
        seen: set[str] = set()
        for sm in submaps:
            text, _ = await self.fetch_text(client, sm)
            if not text:
                continue
            for raw in LOC_RE.findall(text):
                url = canonical_product_url(raw)
                if url and url not in seen:
                    seen.add(url)
                    urls.append(url)
        return urls

    async def search_candidates(self, client: httpx.AsyncClient, code: str) -> list[str]:
        urls: list[str] = []
        seen: set[str] = set()
        page = 1
        total_pages = 0  # 0 = site khong bao so trang (thieu .result-count) -> dung stuck-page
        while page <= MAX_PAGES and (not total_pages or page <= total_pages):
            html_text, _ = await self.fetch_text(client, search_url(code, page))
            if not html_text:
                break
            rows, total_pages = parse_search_page(html_text, code)
            new_urls = 0
            for row in rows:
                url = canonical_product_url(row["URL"])
                if url and url not in seen:
                    seen.add(url)
                    urls.append(url)
                    new_urls += 1
            if not rows or (not total_pages and not new_urls):  # het trang / stuck-page: khong co id moi
                break
            page += 1
            if self.delay:
                await asyncio.sleep(self.delay)
        return urls

    async def crawl_one_url(self, client: httpx.AsyncClient, url: str, input_value: str = "") -> dict[str, Any]:
        html_text, final_url = await self.fetch_text(client, url)
        if not html_text and self.stop_evt.is_set():
            # Bi Stop giua chung -> URL nay CHUA duoc kiem tra that, khong duoc bao Not Found gia
            return not_found(input_value or url, "Unverified - catalog scan interrupted")
        final_canon = canonical_product_url(final_url)
        if not html_text or not final_canon:
            row = not_found(input_value or url)
            row["URL"] = final_canon or canonical_product_url(url)
            return row
        row = parse_product_page(html_text, final_canon, input_value or url)
        return row

    async def crawl_urls(self, client: httpx.AsyncClient, urls: list[str], source: str = "") -> None:
        sem = asyncio.Semaphore(self.conc)
        # Loc bo URL da co trong checkpoint (resume) - khong crawl lai, khong dup row khi export
        urls = [u for u in urls if canonical_product_url(u) not in self.reused_urls]
        self.total += len(urls)

        async def one(url: str) -> None:
            async with sem:
                row = await self.crawl_one_url(client, url, source)
                self.results.append(row)
                self.save_row(self.key, row)
                self.done += 1
                self.log(f"[{row['Status']}] {self.done}/{self.total} | {url[-70:]}")

        await asyncio.gather(*(one(url) for url in urls[:MAX_ITEMS]))

    async def _run_links(self, client: httpx.AsyncClient, links: list[str]) -> None:
        todo = []
        for raw in links:
            canon = canonical_product_url(raw)
            if not canon or canon in self.reused_urls:
                continue
            todo.append(canon)
        await self.crawl_urls(client, todo, source="")

    async def _run_codes(self, client: httpx.AsyncClient, codes: list[str]) -> None:
        todo = [code for code in codes if norm_code(code) not in self.reused_inputs][:MAX_ITEMS]
        self.total += len(todo)
        for code in todo:
            urls = await self.search_candidates(client, code)
            matched: list[dict[str, Any]] = []
            for url in urls:
                row = await self.crawl_one_url(client, url, code)
                if row["Status"] == "OK" and exact_code_match(row, code):
                    matched.append(row)
            if matched:
                for row in matched:
                    self.results.append(row)
                    self.save_row(self.key, row)
            else:
                # Bi Stop giua chung -> search_candidates chua quet het trang, khong duoc bao Not Found that
                status = "Unverified - catalog scan interrupted" if self.stop_evt.is_set() else "Not Found"
                row = not_found(code, status)
                self.results.append(row)
                self.save_row(self.key, row)
            self.done += 1
            self.log(f"[{row['Status']}] {self.done}/{self.total} | {code}")
            if self.delay:
                await asyncio.sleep(self.delay)

    async def _run_keyword(self, client: httpx.AsyncClient, keyword: str, limit: int = 0) -> None:
        urls = await self.search_candidates(client, keyword)
        if limit:
            urls = urls[:limit]
        await self.crawl_urls(client, urls, source=keyword)

    async def _run_all(self, client: httpx.AsyncClient, limit: int = 0) -> None:
        urls = await self.discover_list_urls(client, limit=limit)
        if limit and len(urls) >= limit:
            await self.crawl_urls(client, urls[:limit], source="")
            return
        sitemap_urls = await self.discover_sitemap_urls(client)
        seen = set(urls)
        for url in sitemap_urls:
            if url not in seen:
                seen.add(url)
                urls.append(url)
        if limit:
            urls = urls[:limit]
        await self.crawl_urls(client, urls, source="")

    async def run(
        self,
        mode: str,
        target: str = "",
        excel: str = "",
        sheet: str = "",
        col: str = "",
        limit: int = 0,
    ) -> list[dict[str, Any]]:
        self.key = f"{mode}_{target or excel}_{col or sheet}"
        resumed = self.load_resume_rows(self.key)
        if resumed:
            self.results.extend(resumed)
            self.reused_inputs = {norm_code(row.get("Input")) for row in resumed if row.get("Input")}
            self.reused_urls = {canonical_product_url(row.get("URL", "")) for row in resumed if row.get("URL")}
        async with httpx.AsyncClient(
            headers=HEADERS,
            timeout=TIMEOUT,
            follow_redirects=True,
            max_redirects=MAX_REDIRECTS,
            limits=httpx.Limits(max_connections=self.conc + 4),
        ) as client:
            if mode == "all":
                await self._run_all(client, limit)
            elif mode in {"code", "excel"}:
                raw = [target] if mode == "code" else read_excel_col(excel, sheet, col)
                links = [v for v in raw if str(v).lower().startswith("http")]
                codes = [norm_code(v) for v in raw if not str(v).lower().startswith("http")]
                if links:
                    await self._run_links(client, links)
                if codes:
                    await self._run_codes(client, codes)
            elif mode == "keyword":
                await self._run_keyword(client, target, limit)
            elif mode == "link":
                links = [line.strip() for line in target.splitlines() if line.strip()]
                await self._run_links(client, links)
            else:
                raise ValueError(f"Mode khong hop le: {mode}")
        return self.results


def run_gui() -> None:
    import tkinter as tk
    from tkinter import messagebox, ttk

    root_dir = next(p for p in (HERE, *HERE.parents) if (p / "gui_common.py").exists())
    sys.path.insert(0, str(root_dir))
    from gui_common import ask_export_path, attach_excel_picker

    root = tk.Tk()
    root.title("V Xanh Crawler")
    root.geometry("1080x700")
    ttk.Style().theme_use("clam")

    q: queue.Queue[str] = queue.Queue()
    stop_evt = threading.Event()
    pause_evt = threading.Event()
    state: dict[str, Any] = {"running": False, "rows": [], "eng": None}

    left = ttk.Frame(root, padding=10)
    left.pack(side="left", fill="y")
    right = ttk.Frame(root, padding=(0, 10, 10, 10))
    right.pack(side="right", fill="both", expand=True)

    mode = tk.StringVar(value="code")
    ttk.Label(left, text="Pipeline", font=("Segoe UI", 10, "bold")).pack(anchor="w")
    for val, label in [
        ("code", "1. Theo MA"),
        ("keyword", "2. Theo TU KHOA"),
        ("link", "3. Theo LINK"),
        ("excel", "4. Theo EXCEL"),
        ("all", "5. Crawl ALL"),
    ]:
        ttk.Radiobutton(left, text=label, value=val, variable=mode).pack(anchor="w")

    ttk.Label(left, text="\nMa / Tu khoa:").pack(anchor="w")
    ent = ttk.Entry(left, width=36)
    ent.pack(anchor="w")
    ttk.Label(left, text="Link (moi dong 1 link):").pack(anchor="w")
    box = tk.Text(left, width=36, height=5)
    box.pack(anchor="w")

    xls, sheet, col = tk.StringVar(), tk.StringVar(), tk.StringVar()
    outdir = tk.StringVar(value="images")
    attach_excel_picker(left, xls, sheet, col, outdir)

    conc, delay, limit = tk.IntVar(value=6), tk.DoubleVar(value=0.05), tk.IntVar(value=0)
    for label, var in [("Concurrency:", conc), ("Delay (s):", delay), ("Limit (0=all):", limit)]:
        fr = ttk.Frame(left)
        fr.pack(anchor="w", pady=1)
        ttk.Label(fr, text=label, width=13).pack(side="left")
        ttk.Entry(fr, textvariable=var, width=8).pack(side="left")

    console = tk.Text(right, bg="#12161c", fg="#d6e2ee", insertbackground="#fff", wrap="none")
    console.pack(fill="both", expand=True)
    pb = ttk.Progressbar(right, mode="determinate")
    pb.pack(fill="x", pady=4)
    stat = ttk.Label(right, text="0/0")
    stat.pack(anchor="w")

    def drain() -> None:
        while not q.empty():
            console.insert("end", q.get() + "\n")
            console.see("end")
        eng: Engine | None = state["eng"]
        if eng and eng.total:
            pb["maximum"] = eng.total
            pb["value"] = eng.done
            stat.config(text=f"{eng.done}/{eng.total} | {len(eng.results)} dong")
        root.after(100, drain)

    def worker() -> None:
        target = box.get("1.0", "end").strip() if mode.get() == "link" else ent.get().strip()
        eng = Engine(conc.get(), delay.get(), q.put, stop_evt, pause_evt)
        state["eng"] = eng
        exported = False
        try:
            state["rows"] = asyncio.run(eng.run(mode.get(), target, xls.get(), sheet.get(), col.get(), limit.get()))
        except BaseException as exc:
            q.put(f"[FATAL] {type(exc).__name__}: {exc}")
        finally:
            eng.cleanup(exported)
            state["running"] = False
            q.put(f"[SYS] XONG - {len(state['rows'])} dong. Bam EXPORT de xuat Excel.")

    def start() -> None:
        if state["running"]:
            return
        if mode.get() in {"code", "keyword"} and not ent.get().strip():
            messagebox.showwarning("Thieu du lieu", "Nhap ma hoac tu khoa truoc.")
            return
        if mode.get() == "link" and not box.get("1.0", "end").strip():
            messagebox.showwarning("Thieu du lieu", "Dan link vao o moi dong 1 link.")
            return
        if mode.get() == "excel" and not (xls.get() and col.get()):
            messagebox.showwarning("Thieu du lieu", "Chon file Excel va cot.")
            return
        stop_evt.clear()
        pause_evt.clear()
        state["running"] = True
        console.delete("1.0", "end")
        threading.Thread(target=worker, daemon=True).start()

    def export() -> None:
        eng: Engine | None = state["eng"]
        if not state["rows"]:
            messagebox.showinfo("Export", "Chua co du lieu")
            return
        path = ask_export_path(xls.get(), HERE, "vxanh.xlsx")
        if path:
            export_excel(state["rows"], pathlib.Path(path))
            if eng:
                eng.cleanup(True)
            messagebox.showinfo("Export", f"Da luu:\n{path}")

    fr_btn = ttk.Frame(left)
    fr_btn.pack(anchor="w", pady=8)
    ttk.Button(fr_btn, text="START", command=start).pack(side="left")
    ttk.Button(fr_btn, text="Pause", command=lambda: (pause_evt.clear() if pause_evt.is_set() else pause_evt.set())).pack(side="left")
    ttk.Button(fr_btn, text="Stop", command=stop_evt.set).pack(side="left")
    ttk.Button(fr_btn, text="EXPORT", command=export).pack(side="left")

    def on_close() -> None:
        if state["running"] and not messagebox.askyesno("Thoat", "Dang crawl. Thoat va tat luon?"):
            return
        stop_evt.set()
        try:
            root.destroy()
        finally:
            os._exit(0)

    root.protocol("WM_DELETE_WINDOW", on_close)
    drain()
    root.mainloop()


def main() -> None:
    ap = argparse.ArgumentParser(description="Crawler vxanh.com")
    ap.add_argument("--mode", choices=["code", "keyword", "link", "excel", "all"])
    ap.add_argument("--code")
    ap.add_argument("--keyword")
    ap.add_argument("--link")
    ap.add_argument("--excel")
    ap.add_argument("--sheet", default="")
    ap.add_argument("--col", default="")
    ap.add_argument("--out", default="vxanh.xlsx")
    ap.add_argument("--conc", type=int, default=6)
    ap.add_argument("--delay", type=float, default=0.05)
    ap.add_argument("--limit", type=int, default=0)
    args = ap.parse_args()

    if not args.mode:
        run_gui()
        return

    target = args.code or args.keyword or args.link or ""
    engine = Engine(args.conc, args.delay)
    t0 = time.time()
    exported = False
    try:
        rows = asyncio.run(engine.run(args.mode, target, args.excel or "", args.sheet, args.col, args.limit))
        export_excel(rows, pathlib.Path(args.out))
        exported = True
    finally:
        engine.cleanup(exported)
    print(f"[SYS] {len(rows)} dong -> {args.out} ({time.time() - t0:.1f}s)")


if __name__ == "__main__":
    main()
