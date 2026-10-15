"""
idmarket.vn — crawler WooCommerce (httpx + selectolax, KHONG Selenium, KHONG browser).

RECON (2026-07-13)
------------------
CMS      : WordPress + WooCommerce (theme Electro).
API      : CO — WooCommerce **Store API** (public, khong can login):
             /wp-json/wc/store/v1/products?per_page=100&page=N   -> header X-WP-Total = 37,577 SP
             /wp-json/wc/store/v1/products?search=<SKU>          -> khop ma rat sat
           LUU Y search chi khop TU DAU cua ten/SKU: '6212'->42 KQ, 'vong bi'->100 KQ,
           nhung 'NSK' (nam cuoi ten) -> 0 KQ. Tra ma thi dung; tim theo hang thi khong.
           Field: id, sku, name, permalink, prices{price,regular_price,currency_code},
                  images[{src}], categories[{name}], attributes[{name,terms[{name}]}] (= THONG SO),
                  description, is_in_stock, stock_availability.
robots   : Disallow /wp-json/ va /?s= (chan bot). => uu tien API nhung TU DONG CHUYEN sang
           duong HTML (sitemap + trang san pham, deu duoc phep) khi API 403/404/chan.
sitemap  : /sitemap_index.xml -> product-sitemap1..N.xml (duong fallback discovery)
HTML     : /san-pham/<slug>/  -> h1.product_title | p.price | table.shop_attributes (thong so)
           | ld+json Product (anh, sku). Slug = SKU chuan hoa => tra ma bang slug-direct,
           KHONG can /?s= (endpoint bi robots cam).
gia      : CO gia VND cong khai (API `prices.price`, HTML `p.price`).

Run
---
  idmarket.bat                                    -> GUI
  python idmarket.py --mode code --code 6212-RS1Z/C3GJN-SKF
  python idmarket.py --mode keyword --keyword "vong bi SKF" --limit 50
  python idmarket.py --mode excel --excel parts.xlsx --col "Part ID"
  python idmarket.py --mode link --link https://idmarket.vn/san-pham/<slug>/
  python idmarket.py --mode all [--limit 500]     -> quet sach (~37.6k SP)
  python idmarket.py --mode images --excel parts.xlsx --col "Part ID" --outdir out
  them --html de ep di duong HTML (bo qua Store API)
"""
from __future__ import annotations

import argparse
import asyncio
import gc
import json
import os
import pathlib
import queue
import re
import threading
import time
from typing import Any
from urllib.parse import quote, urljoin, urlparse

import httpx
import pandas as pd
from selectolax.lexbor import LexborHTMLParser

HERE = pathlib.Path(__file__).parent
BASE = "https://idmarket.vn"
API = f"{BASE}/wp-json/wc/store/v1/products"
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/122.0 Safari/537.36")

PER_PAGE = 100            # tran cua Store API
CIRCUIT_BREAK = 25        # N loi LIEN TIEP -> nghi bi chan IP -> dung han
MAX_ITEMS = 60000         # site co ~37.6k SP
MAX_PAGES = 600
MAX_REDIRECTS = 5
TIMEOUT = 30.0
GC_EVERY = 1000


# ================================================================ CORE (khong dung GUI)
def norm_code(s: Any) -> str:
    """Chuan hoa ma de so khop. URL thi GIU NGUYEN hoa/thuong (path phan biet chu)."""
    v = re.sub(r"\s+", "", str(s or ""))
    return v if v.lower().startswith("http") else v.upper()


def slugify_sku(sku: str) -> str:
    """SKU -> slug WooCommerce: '6212-RS1Z/C3GJN-SKF' -> '6212-rs1z-c3gjn-skf'."""
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", sku.lower())).strip("-")


def api_search_url(sku: str) -> str:
    return f"{API}?search={quote(sku)}&per_page=10"


def slug_urls(sku: str) -> list[str]:
    """Ung vien URL truc tiep (khong dung /?s= vi robots cam)."""
    s = slugify_sku(sku)
    return [f"{BASE}/san-pham/{s}/", f"{BASE}/san-pham/vong-bi-{s}/"]


def price_of(p: dict) -> str:
    pr = p.get("prices") or {}
    v = pr.get("price") or ""
    return f"{int(v):,}".replace(",", ".") if str(v).isdigit() else str(v)


def rec_from_api(p: dict) -> dict:
    spec = {}
    for a in p.get("attributes") or []:
        name = (a.get("name") or "").strip()
        terms = ", ".join(t.get("name", "") for t in (a.get("terms") or []))
        if name and terms:
            spec[name] = terms
    return {"Input": "", "SKU": p.get("sku") or "", "Ten": p.get("name") or "",
            "Gia (VND)": price_of(p),
            "Ton kho": "Con hang" if p.get("is_in_stock") else "Het hang", "Status": "OK",
            "_cats": [c.get("name", "") for c in (p.get("categories") or [])],
            "_spec": spec,
            "Image": " | ".join(i.get("src", "") for i in (p.get("images") or [])[:5]),
            "URL": p.get("permalink") or ""}


def parse_html(html: str, url: str) -> dict | None:
    """Fallback khi API bi chan: parse trang san pham WooCommerce/Electro."""
    t = LexborHTMLParser(html)
    h1 = t.css_first("h1.product_title")
    if h1 is None:
        return None
    name = re.sub(r"\s+", " ", h1.text(strip=True))

    sku, images = "", []
    for sc in t.css('script[type="application/ld+json"]'):
        try:
            data = json.loads(sc.text())
        except Exception:
            continue
        for blk in (data.get("@graph", []) if isinstance(data, dict) else
                    (data if isinstance(data, list) else [data])):
            if not isinstance(blk, dict) or blk.get("@type") not in ("Product", ["Product"]):
                continue
            sku = sku or blk.get("sku") or blk.get("mpn") or ""
            img = blk.get("image")
            for u in (img if isinstance(img, list) else [img]):
                u = u.get("url") if isinstance(u, dict) else u
                if u and u not in images:
                    images.append(u)
    if not images:
        for im in t.css("figure.woocommerce-product-gallery__image img, "
                        "figure.electro-wc-product-gallery__image img"):
            u = im.attributes.get("data-large_image") or im.attributes.get("src") or ""
            if u and u not in images:
                images.append(urljoin(url, u))

    spec: dict[str, str] = {}
    tb = t.css_first("table.shop_attributes")
    if tb:
        for tr in tb.css("tr"):
            cells = tr.css("th, td")
            if len(cells) == 2:
                k = re.sub(r"\s+", " ", cells[0].text(strip=True))
                v = re.sub(r"\s+", " ", cells[1].text(strip=True))
                if k and v:
                    spec[k] = v

    pr = t.css_first("p.price")
    cats = [a.text(strip=True) for a in t.css("nav.woocommerce-breadcrumb a")][1:]
    return {"Input": "", "SKU": sku, "Ten": name,
            "Gia (VND)": re.sub(r"\s+", " ", pr.text(strip=True)) if pr else "",
            "Ton kho": "", "Status": "OK", "_cats": cats, "_spec": spec,
            "Image": " | ".join(images[:5]), "URL": url}


def export_excel(results: list[dict], path: pathlib.Path) -> pathlib.Path:
    """Spec gop 1 COT 'Thong so' (moi dong 'Ten: Gia tri'). Xem AGENTS.md 4.6."""
    from openpyxl.utils import get_column_letter

    max_cat = max((len(r.get("_cats") or []) for r in results), default=0)
    rows = []
    for r in results:
        row = {k: r.get(k, "") for k in ("Input", "SKU", "Ten", "Gia (VND)", "Ton kho", "Status")}
        cats = r.get("_cats") or []
        for i in range(max_cat):
            row[f"Category {i + 1}"] = cats[i] if i < len(cats) else ""
        row["Thong so"] = "\n".join(f"{k}: {v}" for k, v in (r.get("_spec") or {}).items())
        row["Image"] = r.get("Image", "")
        row["URL"] = r.get("URL", "")
        rows.append(row)
    df = pd.DataFrame(rows)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, index=False)
        ws = xw.sheets["Sheet1"]
        col = get_column_letter(df.columns.get_loc("Thong so") + 1)
        ws.column_dimensions[col].width = 60
        for c in ws[col][1:]:
            c.alignment = c.alignment.copy(wrap_text=True, vertical="top")
    return path


class Engine:
    """Store API la duong chinh; API chan/loi -> tu dong roi sang HTML (sitemap + slug-direct)."""

    def __init__(self, conc: int = 6, delay: float = 0.1, log=print, force_html: bool = False,
                 stop_evt: threading.Event | None = None,
                 pause_evt: threading.Event | None = None) -> None:
        self.conc, self.delay, self.log = max(1, conc), max(0.0, delay), log
        self.stop_evt = stop_evt or threading.Event()
        self.pause_evt = pause_evt or threading.Event()
        self.api_ok = not force_html
        self.fails = 0
        self.results: list[dict] = []
        self.seen_urls: set[tuple[str, str]] = set()  # key (Input, URL): site doi thu tu SP giua
                                            # cac trang API -> phai dedup, nhung 2 ma KHAC nhau
                                            # tro chung 1 URL van phai co dong rieng (BAY §10-6)
        self.done = 0
        self.total = 0
        self.key = ""
        self._ckpt_loaded = False

    def add(self, recs: list[dict], key: str) -> list[dict]:
        """Them record moi (bo trung theo (Input, URL)) + ghi checkpoint."""
        fresh = []
        for r in recs:
            u = r.get("URL") or ""
            k = (r.get("Input") or "", u)
            if u and k in self.seen_urls:
                continue
            if u:
                self.seen_urls.add(k)
            fresh.append(r)
        if fresh:
            self.results += fresh
            self._save(key, fresh)
        return fresh

    # ---- resume: JSONL sidecar
    def ckpt(self, key: str) -> pathlib.Path:
        return HERE / f"_raw_{re.sub(r'[^A-Za-z0-9]+', '_', key)[:40]}.jsonl"

    def _load_ckpt(self, key: str) -> set[str]:
        p, seen = self.ckpt(key), set()
        if self._ckpt_loaded:             # goi 2 lan (API roi fallback HTML) -> KHONG nap lai
            return set(u for _, u in self.seen_urls)
        self._ckpt_loaded = True
        if not p.exists():
            return seen
        for line in p.read_text("utf-8", errors="ignore").splitlines():
            try:
                r = json.loads(line)
            except Exception:
                continue
            self.results.append(r)
            seen.add(r.get("URL", ""))
            self.seen_urls.add((r.get("Input", ""), r.get("URL", "")))
        if seen:
            self.log(f"[SYS] Resume: {len(self.results)} record da co -> bo qua")
        return seen

    def _save(self, key: str, recs: list[dict]) -> None:
        with self.ckpt(key).open("a", encoding="utf-8") as f:
            for r in recs:
                f.write(json.dumps(r, ensure_ascii=False) + "\n")
            f.flush()

    def cleanup(self) -> None:
        p = self.ckpt(self.key) if self.key else None
        if p and p.exists() and not self.stop_evt.is_set():
            p.unlink()

    async def _gate(self) -> bool:
        while self.pause_evt.is_set() and not self.stop_evt.is_set():
            await asyncio.sleep(0.2)
        return not self.stop_evt.is_set()

    async def fetch(self, cli: httpx.AsyncClient, url: str, tries: int = 4) -> httpx.Response | None:
        for i in range(tries):
            if not await self._gate():
                return None
            try:
                r = await cli.get(url)
                if r.status_code in (403, 429, 503):
                    if API in url and self.api_ok:      # API bi chan -> chuyen huong HTML
                        self.log(f"[WARN] Store API tra {r.status_code} -> CHUYEN sang duong HTML")
                        self.api_ok = False
                        return None
                    wait = min(2 ** i * 5, 60)
                    self.log(f"[WARN] {r.status_code} -> ngu {wait}s | {url[-55:]}")
                    await asyncio.sleep(wait)
                    continue
                if r.status_code == 404:
                    return None
                r.raise_for_status()
                self.fails = 0
                return r
            except Exception as e:
                if i < tries - 1:
                    await asyncio.sleep(min(2 ** i, 20))
                    continue
                self.fails += 1
                self.log(f"[ERR] {type(e).__name__} | {url[-55:]}")
                if self.fails >= CIRCUIT_BREAK:
                    self.log(f"[STOP] {CIRCUIT_BREAK} loi LIEN TIEP - nghi bi chan IP. DUNG.")
                    self.stop_evt.set()
                return None
        return None

    async def _json(self, cli: httpx.AsyncClient, url: str) -> Any:
        r = await self.fetch(cli, url)
        if r is None:
            return None
        try:
            return r.json()
        except Exception:
            self.log("[WARN] API khong tra JSON -> CHUYEN sang duong HTML")
            self.api_ok = False
            return None

    # ---- resolve 1 ma -> record (API truoc, HTML slug-direct sau)
    async def _one_code(self, cli: httpx.AsyncClient, code: str) -> dict | None:
        if self.api_ok:
            data = await self._json(cli, api_search_url(code))
            if data:
                want = norm_code(code)
                hit = next((p for p in data if norm_code(p.get("sku")) == want), None)   # khop CHINH XAC
                partial = hit is None
                if partial:
                    hit = next((p for p in data if want in norm_code(p.get("sku"))), None)
                if hit:
                    r = rec_from_api(hit)
                    r["Input"] = code
                    if partial:                     # BAY §10-5: khong khop chinh xac -> phai danh dau
                        r["Status"] = "OK (khop mot phan)"
                    return r
        for u in slug_urls(code):                    # fallback: slug-direct (robots cam /?s=)
            resp = await self.fetch(cli, u)
            if resp is None:
                continue
            rec = parse_html(resp.text, str(resp.url))
            if rec:
                rec["Input"] = code
                return rec
        return None

    async def _html_pages(self, cli: httpx.AsyncClient, urls: list[str], key: str) -> None:
        seen = self._load_ckpt(key)
        urls = [u for u in dict.fromkeys(urls) if u not in seen][:MAX_ITEMS]
        self.total, self.done = len(urls) + len(seen), len(seen)
        sem = asyncio.Semaphore(self.conc)

        async def one(u: str) -> None:
            async with sem:
                if not await self._gate():
                    return
                r = await self.fetch(cli, u)
                if self.delay:
                    await asyncio.sleep(self.delay)
                rec = parse_html(r.text, str(r.url)) if r else None
                if rec:
                    rec["Input"] = rec["Input"] or u
                    self.add([rec], key)
                self.done += 1
                self.log(f"[{'OK ' if rec else 'MISS'}] {self.done}/{self.total} | {u[-55:]}")
                if self.done % GC_EVERY == 0:
                    gc.collect()

        try:
            await asyncio.gather(*(one(u) for u in urls))
        except (KeyboardInterrupt, Exception) as e:
            self.log(f"[FATAL] {type(e).__name__}: {e}")

    async def _sitemap_product_urls(self, cli: httpx.AsyncClient) -> list[str]:
        idx = await self.fetch(cli, f"{BASE}/sitemap_index.xml")
        if idx is None:
            return []
        sms = [u for u in re.findall(r"<loc>([^<]+)</loc>", idx.text) if "product-sitemap" in u]
        sem = asyncio.Semaphore(self.conc)

        async def one(sm: str) -> list[str]:
            async with sem:
                r = await self.fetch(cli, sm)
                got = [u for u in re.findall(r"<loc>([^<]+)</loc>", r.text)
                       if "/san-pham/" in u] if r else []
                self.log(f"[SYS] {sm.rsplit('/', 1)[-1]} -> {len(got)} URL")
                return got

        return [u for part in await asyncio.gather(*(one(s) for s in sms)) for u in part]

    async def run(self, mode: str, target: str = "", excel: str = "", sheet: str = "",
                  col: str = "", limit: int = 0, outdir: str = "out") -> list[dict]:
        key = self.key = f"{mode}_{target or excel}_{col}"
        headers = {"User-Agent": UA, "Accept-Language": "vi-VN,vi;q=0.9,en;q=0.8"}
        async with httpx.AsyncClient(headers=headers, timeout=TIMEOUT, follow_redirects=True,
                                     max_redirects=MAX_REDIRECTS,
                                     limits=httpx.Limits(max_connections=self.conc + 4)) as cli:
            if mode == "all":
                if self.api_ok:
                    await self._all_api(cli, key, limit)
                # API chan -> LUON chay fallback (KHONG check self.results: khi resume, results da
                # co san tu checkpoint -> check do se nuot fallback va ket thuc voi data thieu)
                if not self.api_ok and not self.stop_evt.is_set():
                    urls = (await self._sitemap_product_urls(cli))[:limit or MAX_ITEMS]
                    self.log(f"[SYS] Store API bi chan -> Fallback HTML: {len(urls)} trang")
                    await self._html_pages(cli, urls, key)

            elif mode in ("code", "excel"):
                raw = [norm_code(target)] if mode == "code" else read_excel_col(excel, sheet, col)
                raw = [c for c in dict.fromkeys(raw) if c][:MAX_ITEMS]
                links = [c for c in raw if c.lower().startswith("http")]
                codes = [c for c in raw if c not in links]
                if links:
                    await self._html_pages(cli, links, key)
                self._load_ckpt(key)
                self.total = len(codes) + self.done
                sem = asyncio.Semaphore(self.conc)

                async def one(c: str) -> None:
                    async with sem:
                        if not await self._gate():
                            return
                        rec = await self._one_code(cli, c)
                        await asyncio.sleep(self.delay)
                        self.done += 1
                        if rec:
                            self.add([rec], key)
                            self.log(f"[OK ] {self.done}/{self.total} {c} -> {rec['Ten'][:45]}")
                        else:
                            self.results.append({"Input": c, "SKU": c, "Ten": "", "Gia (VND)": "",
                                                 "Ton kho": "", "Status": "Not Found",
                                                 "_cats": [], "_spec": {}, "Image": "", "URL": ""})
                            self.log(f"[MISS] {self.done}/{self.total} {c}")

                await asyncio.gather(*(one(c) for c in codes
                                       if c not in {r.get("Input") for r in self.results}))

            elif mode == "keyword":
                if self.api_ok:
                    page = 1
                    while page <= MAX_PAGES and await self._gate():
                        data = await self._json(
                            cli, f"{API}?search={quote(target)}&per_page={PER_PAGE}&page={page}")
                        if not data:
                            break
                        self.add([rec_from_api(p) for p in data], key)
                        self.log(f"[SYS] '{target}' trang {page} -> {len(self.results)} SP")
                        if len(data) < PER_PAGE or (limit and len(self.results) >= limit):
                            break
                        page += 1
                        await asyncio.sleep(self.delay)
                if not self.results:
                    self.log("[MISS] 0 ket qua. Store API search chi khop TU DAU cua ten/SKU "
                             "('6212' OK, 'vong bi' OK, nhung 'NSK' -> 0). Doi tu khoa, "
                             "hoac dung mode 'code' (slug-direct) / 'all'.")
                if limit:
                    self.results = self.results[:limit]

            elif mode == "link":
                await self._html_pages(
                    cli, [u.strip() for u in target.splitlines() if u.strip().startswith("http")], key)

            elif mode == "images":
                await self._download_images(cli, excel, sheet, col, outdir)

            else:
                self.log(f"[ERR] mode khong hop le: {mode}")
        return self.results

    async def _all_api(self, cli: httpx.AsyncClient, key: str, limit: int) -> None:
        seen = self._load_ckpt(key)
        r = await self.fetch(cli, f"{API}?per_page=1")
        if r is None:
            return
        total = int(r.headers.get("X-WP-Total") or 0)
        self.total = min(total, limit or total)
        self.log(f"[SYS] Store API: {total} SP -> lay {self.total}")
        pages = (self.total + PER_PAGE - 1) // PER_PAGE
        sem = asyncio.Semaphore(self.conc)

        async def one(pg: int) -> None:
            async with sem:
                if not await self._gate() or not self.api_ok:
                    return
                data = await self._json(cli, f"{API}?per_page={PER_PAGE}&page={pg}")
                await asyncio.sleep(self.delay)
                self.add([rec_from_api(p) for p in (data or [])], key)
                self.done = min(self.done + len(data or []), self.total)
                self.log(f"[OK ] trang {pg}/{pages} -> tong {len(self.results)} SP")
                if self.done % GC_EVERY == 0:
                    gc.collect()

        try:
            await asyncio.gather(*(one(p) for p in range(1, pages + 1)))
        except (KeyboardInterrupt, Exception) as e:
            self.log(f"[FATAL] {type(e).__name__}: {e}")
        if limit:
            self.results = self.results[:limit]

    async def _download_images(self, cli: httpx.AsyncClient, excel: str, sheet: str,
                               col: str, outdir: str) -> None:
        """Excel (cot ma hoac URL) -> anh -> outdir/<ma>/<ma>_N.jpg. Ho tro ca vuavongbi.com."""
        root = pathlib.Path(outdir) if os.path.isabs(outdir) else HERE / outdir
        root.mkdir(parents=True, exist_ok=True)
        keys = read_excel_col(excel, sheet, col)[:MAX_ITEMS]
        self.total = len(keys)
        sem = asyncio.Semaphore(self.conc)

        async def one(k: str) -> None:
            async with sem:
                if not await self._gate():
                    return
                rec = None
                if k.lower().startswith("http"):
                    r = await self.fetch(cli, k)
                    rec = parse_html(r.text, str(r.url)) if r else None
                else:
                    rec = await self._one_code(cli, k)
                self.done += 1
                if not rec or not rec.get("Image"):
                    self.log(f"[MISS] {self.done}/{self.total} khong ra anh: {k}")
                    return
                name = re.sub(r'[\\/*?:"<>|]', "_", (rec.get("SKU") or k))[:60]
                folder = root / name
                folder.mkdir(parents=True, exist_ok=True)
                ok = 0
                for i, iu in enumerate(rec["Image"].split(" | "), 1):
                    try:
                        rr = await cli.get(iu, headers={"Referer": rec.get("URL") or BASE})
                        rr.raise_for_status()
                        (folder / f"{name}_{i}{pathlib.Path(urlparse(iu).path).suffix or '.jpg'}"
                         ).write_bytes(rr.content)
                        ok += 1
                    except Exception as e:
                        self.log(f"[ERR] tai anh {iu[-45:]}: {type(e).__name__}")
                self.log(f"[OK ] {self.done}/{self.total} {name} -> {ok} anh")

        await asyncio.gather(*(one(k) for k in keys))


def read_excel_col(path: str, sheet: str, col: str) -> list[str]:
    if not path:
        raise ValueError("Chua chon file Excel")
    df = pd.read_excel(path, sheet_name=sheet or 0, dtype=str)
    df.columns = df.columns.str.strip()
    if col not in df.columns:
        # ValueError chu KHONG SystemExit: SystemExit la BaseException, worker GUI bat
        # `except Exception` se KHONG bat duoc -> thread chet cam, GUI treo running=True
        raise ValueError(f"Khong thay cot '{col}'. Cot co san: {list(df.columns)}")
    return [norm_code(v) for v in df[col].dropna().tolist()]


# ================================================================ GUI
def run_gui() -> None:
    import sys
    import tkinter as tk
    from tkinter import messagebox, ttk

    # ha tang GUI chung o root (di nguoc len cay thu muc -> song sot khi doi nhom folder)
    root_dir = next(p for p in (HERE, *HERE.parents) if (p / "gui_common.py").exists())
    sys.path.insert(0, str(root_dir))
    from gui_common import ask_export_path, attach_excel_picker

    root = tk.Tk()
    root.title("IDMarket Crawler — idmarket.vn (Store API + HTML fallback)")
    root.geometry("1080x680")
    ttk.Style().theme_use("clam")

    q: queue.Queue[str] = queue.Queue()
    stop_evt, pause_evt = threading.Event(), threading.Event()
    state: dict[str, Any] = {"running": False, "results": [], "eng": None}

    left = ttk.Frame(root, padding=10)
    left.pack(side="left", fill="y")
    right = ttk.Frame(root, padding=(0, 10, 10, 10))
    right.pack(side="right", fill="both", expand=True)

    mode = tk.StringVar(value="code")
    ttk.Label(left, text="Pipeline", font=("Segoe UI", 10, "bold")).pack(anchor="w")
    for val, txt in [("code", "1. Theo MA"), ("keyword", "2. Theo TU KHOA"),
                     ("link", "3. Theo LINK (moi dong 1 link)"), ("excel", "4. Theo cot EXCEL"),
                     ("all", "5. Crawl ALL (~37.6k SP)"), ("images", "6. Tai ANH")]:
        ttk.Radiobutton(left, text=txt, value=val, variable=mode).pack(anchor="w")

    ttk.Label(left, text="\nMa / Tu khoa:").pack(anchor="w")
    ent = ttk.Entry(left, width=36)
    ent.pack(anchor="w")
    ttk.Label(left, text="Link (moi dong 1 link):").pack(anchor="w")
    box = tk.Text(left, width=36, height=5)
    box.pack(anchor="w")

    xls, sheet, col = tk.StringVar(), tk.StringVar(), tk.StringVar()
    outdir = tk.StringVar(value="out")
    attach_excel_picker(left, xls, sheet, col, outdir)

    conc, delay, limit = tk.IntVar(value=6), tk.DoubleVar(value=0.1), tk.IntVar(value=0)
    for lbl, var in [("Concurrency:", conc), ("Delay (s):", delay), ("Limit (0=all):", limit)]:
        f3 = ttk.Frame(left)
        f3.pack(anchor="w", pady=1)
        ttk.Label(f3, text=lbl, width=13).pack(side="left")
        ttk.Entry(f3, textvariable=var, width=8).pack(side="left")
    force_html = tk.BooleanVar(value=False)
    ttk.Checkbutton(left, text="Ep di duong HTML (bo Store API)", variable=force_html).pack(anchor="w")

    console = tk.Text(right, bg="#12161c", fg="#d6e2ee", insertbackground="#fff", wrap="none")
    console.pack(fill="both", expand=True)
    pb = ttk.Progressbar(right, mode="determinate")
    pb.pack(fill="x", pady=4)
    stat = ttk.Label(right, text="0/0")
    stat.pack(anchor="w")

    def drain() -> None:
        while not q.empty():
            console.insert("end", q.get() + "\n")
            console.see("end")
        eng: Engine | None = state["eng"]
        if eng and eng.total:
            pb["maximum"], pb["value"] = eng.total, eng.done
            stat.config(text=f"{eng.done}/{eng.total} | {len(eng.results)} record")
        root.after(100, drain)

    def worker() -> None:
        m = mode.get()
        target = box.get("1.0", "end").strip() if m == "link" else ent.get().strip()
        eng = Engine(conc.get(), delay.get(), q.put, force_html.get(), stop_evt, pause_evt)
        state["eng"] = eng
        try:
            asyncio.run(eng.run(m, target, xls.get(), sheet.get(), col.get(),
                                limit.get(), outdir.get()))
            eng.cleanup()
        except BaseException as e:     # ke ca SystemExit -> khong de thread chet cam
            q.put(f"[FATAL] {type(e).__name__}: {e}")
        finally:                       # LUON nha co, khong khoa START vinh vien
            state["results"] = eng.results
            state["running"] = False
            q.put(f"[SYS] XONG — {len(eng.results)} record. Bam EXPORT de xuat Excel.")

    def start() -> None:
        if state["running"]:
            return
        stop_evt.clear()
        pause_evt.clear()
        state["running"] = True
        console.delete("1.0", "end")
        threading.Thread(target=worker, daemon=True).start()

    def export() -> None:
        if not state["results"]:
            messagebox.showinfo("Export", "Chua co du lieu")
            return
        p = ask_export_path(xls.get(), HERE, "idmarket.xlsx")   # mac dinh canh file Excel dau vao
        if p:
            export_excel(state["results"], pathlib.Path(p))
            messagebox.showinfo("Export", f"Da luu:\n{p}")

    fb = ttk.Frame(left)
    fb.pack(anchor="w", pady=8)
    ttk.Button(fb, text="START", command=start).pack(side="left")
    ttk.Button(fb, text="Pause", command=lambda: (pause_evt.clear() if pause_evt.is_set()
                                                  else pause_evt.set())).pack(side="left")
    ttk.Button(fb, text="Stop", command=stop_evt.set).pack(side="left")
    ttk.Button(fb, text="EXPORT", command=export).pack(side="left")

    def on_close() -> None:
        if state["running"] and not messagebox.askyesno("Thoat", "Dang crawl. Thoat va tat luon?"):
            return
        stop_evt.set()
        try:
            root.destroy()
        finally:
            os._exit(0)

    root.protocol("WM_DELETE_WINDOW", on_close)
    drain()
    root.mainloop()


# ================================================================ CLI
def main() -> None:
    ap = argparse.ArgumentParser(description="idmarket.vn crawler (Store API + HTML fallback)")
    ap.add_argument("--mode", choices=["code", "keyword", "link", "excel", "all", "images"])
    ap.add_argument("--code")
    ap.add_argument("--keyword")
    ap.add_argument("--link")
    ap.add_argument("--excel")
    ap.add_argument("--sheet", default="")
    ap.add_argument("--col", default="Part ID")
    ap.add_argument("--outdir", default="out")
    ap.add_argument("--out", default="idmarket.xlsx")
    ap.add_argument("--html", action="store_true", help="ep di duong HTML, bo qua Store API")
    ap.add_argument("--conc", type=int, default=6)
    ap.add_argument("--delay", type=float, default=0.1)
    ap.add_argument("--limit", type=int, default=0)
    a = ap.parse_args()

    if not a.mode:
        run_gui()
        return

    eng = Engine(a.conc, a.delay, print, a.html)
    t0 = time.time()
    try:
        asyncio.run(eng.run(a.mode, a.code or a.keyword or a.link or "", a.excel or "",
                            a.sheet, a.col, a.limit, a.outdir))
    finally:
        if eng.results and a.mode != "images":
            p = export_excel(eng.results, pathlib.Path(a.out))
            print(f"[SYS] {len(eng.results)} record -> {p} ({time.time() - t0:.1f}s)")
        eng.cleanup()


if __name__ == "__main__":
    main()
