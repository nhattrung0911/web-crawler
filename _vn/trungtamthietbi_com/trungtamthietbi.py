"""
trungtamthietbi.com (Trung Tam Thiet Bi va Dung Cu Truc Tuyen) — httpx, NO browser, NO Selenium.

RECON (2026-09-28)
------------------
Nen tang : CMS PHP TU VIET (Apache/2, PHP 5.3.29, cookie PHPSESSID, template
           /generator/pages/mods/*). KHONG phai Shopify/Haravan/Sapo/WordPress.
API      : **KHONG CO** JSON API. Da thu va 404 het: /products.json, /wp-json/wp/v2/product,
           /api/search, /search?q=, /tim-kiem. Chi co HTML render server-side => parse HTML.
TLS      : !! BAY #1 — server chi gui CERT LEAF, THIEU trung gian "GoGetSSL RSA DV CA"
           => httpx/requests bao CERTIFICATE_VERIFY_FAILED (curl chay vi bundle rieng).
           KHONG dung verify=False (AGENTS 10). Cach chua: nhung san cert trung gian
           (CA_EXTRA duoi day, tai tu AIA http://crt.usertrust.com/GoGetSSLRSADVCA.crt)
           vao ssl context cung certifi => verify VAN BAT.
URL SP   : /<slug>-<id>.html  (vd /0-0-16mm-dong-ho-so-teclock-tm-1200-6336.html)
    BAY #2: sitemap con chua /lien-he,<slug>-<id>.html (form lien he, KHONG phai SP) va ban
            sao co duoi rac /<slug>-<id>.html/<gi-do>/ => loc bo, dedup theo path chuan.
sitemap  : /sitemap.xml -> sitemap1.xml (40.000 loc) + sitemap2.xml (6.052 loc) = 46.052 loc
           -> 16.113 SP that (sau khi loc lien-he + duoi rac).
    BAY #3: **SITEMAP CU** (lastmod 2016). Trang danh muc co SP KHONG nam trong sitemap
            (vd /215x180mm-dem-dau-goi-ingco-hkpt0101-24905.html, id 24905 > max sitemap).
            => discovery chinh la BFS danh muc; sitemap chi de VA them.
cay muc  : /viewall.html -> 11 x /danh-muc/<slug>-<id>.html -> 57 x /danh-sach/<slug>-<id>.html
           -> phan trang "/danh-sach/<slug>-<id>,<N>.html" (base = trang 1, ,1 = trang 2),
           16 SP/trang, item: <div class="sanpham"> sp_pic / sp_name / sp_model (Ma hang) / sp_price.
           /list/<quoc-gia|hang>,<N>.html = loc theo xuat xu (50 SP/trang) - nguon phu.
search   : /search.html?keywords=<kw> — HTML server-side.
    BAY #4: ma khop duy nhat -> **REDIRECT THANG sang trang SP** (302 -> /<slug>-<id>.html).
            Phai check URL cuoi: la trang SP thi parse luon (1 request/ma), khong thi doc luoi.
           Luoi ket qua: 15 SP/trang, block <div class="one_productsearch"> co san MA HANG
           (one_product_ma1) => khop ma TRUOC roi moi fetch detail (tiet kiem tren server cham).
           Phan trang: &p=<N>.html, p=0 == trang goc, het khi trang 0 item (stuck-page).
           Tong: "Sản phẩm tìm được (46 Sản phẩm)".
detail   : ten = <title>; ma he thong = "Ma san pham: <span>TCK-120-035</span>";
           gia = #giasanpham + #giasanphamvat (co SP "Xin lien he" = khong gia);
           hang = link /nha-san-xuat/<slug>.html; category = breadcrumb "Trang chu / c1 / c2 / c3";
           anh = ../products/<slug>.jpg (tuong doi -> phai noi BASE, co duoi "/" rac).
spec     : KHONG co <table>. Nam trong mo ta dang heading "THONG SO KI THUAT" roi cac dong
           "- Ten: Gia tri". ~5/6 SP co, con lai chi van mo ta.
    BAY #5: vung mo ta BI DINH carousel "san pham lien quan" (<div class="n-box">) -> neu khong
            cat se hut ca "Gia: 8,099,000 VND" cua SP KHAC vao spec. Cat tai class="n-box".
    Ma hang thuong nam trong spec ("Ma san pham: KPT-14ME") KHAC ma he thong (KPT-...) => so CA HAI.
toc do   : SERVER CHAM + PHP SESSION LOCK. PHP 5.3 set PHPSESSID va khoa session => request chung
           cookie bi XEP HANG. 1 client dung chung (ban cu) cho 16 SP: per-request 0,8 -> 13,4s
           (cau thang hang doi), wall 13,4s; 16 client cookie rieng: wall 3,6s. robots.txt tinh 0,25s.
           => dung ClientPool: N=conc AsyncClient, moi cai 1 cookie jar, moi request MUON 1 client.
           Do that 48 SP thu (2026-10-02, 2 mau): 1 client chung conc 8 = 1,50-1,56 req/s | conc 16
           = 2,19-2,46 ; POOL conc 8 = 3,37-4,02 | conc 16 = 4,43-4,70 req/s. 0 loi, 0 lan 429/403.
           Mac dinh **conc 16**, timeout 90s. CRAWL ALL THAT (2026-10-02, pool, conc 16): work-list
           26.612 SP -> 26.596 OK, 96 phut (~4,6 SP/s gom ca discovery). Ban cu 1 client chung
           ~1,4-2,4 req/s (lan chay 09-28 chi duoc 4.853 SP / 57 phut = 1,4/s).
robots   : Allow tat ca, co khai sitemap. Charset utf-8 (HTML dung entity &agrave; -> unescape).
Pause/Stop: server cham nen do that: Pause chi chan request MOI, cac request DANG BAY (toi conc)
           van hoan tat va ghi tiep (~5 dong) — khong mat du lieu da tai. Stop dung sau
           ~1 nhip request (do duoc 13,6s voi conc 16) chu khong tuc thi, KHONG treo.

Output   : out_trungtamthietbi/ (gitignore out_*/): checkpoint _raw_*.jsonl, trungtamthietbi.log
           (chi ghi ERR/WARN/STOP/FATAL), Excel CLI mac dinh. Checkpoint cu o folder site tu dong
           DOI (move) sang day khi chay.

Run
---
  trungtamthietbi.bat                                        -> GUI
  python trungtamthietbi.py --mode all [--limit 500]         -> BFS danh muc + va sitemap
  python trungtamthietbi.py --mode code --code TCK-120-035    (ma he thong hoac ma hang)
  python trungtamthietbi.py --mode keyword --keyword teclock [--limit 100]
  python trungtamthietbi.py --mode link --link https://www.trungtamthietbi.com/<slug>-<id>.html
  python trungtamthietbi.py --mode excel --excel parts.xlsx --sheet Sheet1 --col Part
"""
from __future__ import annotations

import argparse
import asyncio
import gc
import hashlib
import html as htmlmod
import json
import os
import pathlib
import queue
import re
import shutil
import ssl
import sys
import threading
import time
import unicodedata
from typing import Any, Iterable, Iterator
from urllib.parse import quote

import certifi
import httpx
import pandas as pd

HERE = pathlib.Path(__file__).parent
BASE = "https://www.trungtamthietbi.com"
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/126.0 Safari/537.36")
HEADERS = {"User-Agent": UA, "Accept-Language": "vi-VN,vi;q=0.9,en;q=0.8"}
OUT_DIR = HERE / "out_trungtamthietbi"      # checkpoint + log + Excel mac dinh (gitignore out_*/)
LOG_FILE = OUT_DIR / "trungtamthietbi.log"

# BAY #1: server thieu cert trung gian -> nhung vao day de VAN verify duoc (khong verify=False)
CA_EXTRA = """-----BEGIN CERTIFICATE-----
MIIF1zCCA7+gAwIBAgIRAJOLsI5imHtPdfmMtqUEXJYwDQYJKoZIhvcNAQEMBQAw
gYgxCzAJBgNVBAYTAlVTMRMwEQYDVQQIEwpOZXcgSmVyc2V5MRQwEgYDVQQHEwtK
ZXJzZXkgQ2l0eTEeMBwGA1UEChMVVGhlIFVTRVJUUlVTVCBOZXR3b3JrMS4wLAYD
VQQDEyVVU0VSVHJ1c3QgUlNBIENlcnRpZmljYXRpb24gQXV0aG9yaXR5MB4XDTE4
MDkwNjAwMDAwMFoXDTI4MDkwNTIzNTk1OVowTDELMAkGA1UEBhMCTFYxDTALBgNV
BAcTBFJpZ2ExETAPBgNVBAoTCEdvR2V0U1NMMRswGQYDVQQDExJHb0dldFNTTCBS
U0EgRFYgQ0EwggEiMA0GCSqGSIb3DQEBAQUAA4IBDwAwggEKAoIBAQCfwF4hD6E1
kLglXs1n2fH5vMQukCGyyD4LqLsc3pSzeh8we7njU4TB85BH5YXqcfwiH1Sf78aB
hk1FgXoAZ3EQrF49We8mnTtTPFRnMwEHLJRpY9I/+peKeAZNL0MJG5zM+9gmcSpI
OTI6p7MPela72g0pBQjwcExYLqFFVsnroEPTRRlmfTBTRi9r7rYcXwIct2VUCRmj
jR1GX13op370YjYwgGv/TeYqUWkNiEjWNskFDEfxSc0YfoBwwKdPNfp6t/5+RsFn
lgQKstmFLQbbENsdUEpzWEvZUpDC4qPvRrxEKcF0uLoZhEnxhskwXSTC64BNtc+l
VEk7/g/be8svAgMBAAGjggF1MIIBcTAfBgNVHSMEGDAWgBRTeb9aqitKz1SA4dib
wJ3ysgNmyzAdBgNVHQ4EFgQU+ftQxItnu2dk/oMhpqnOP1WEk5kwDgYDVR0PAQH/
BAQDAgGGMBIGA1UdEwEB/wQIMAYBAf8CAQAwHQYDVR0lBBYwFAYIKwYBBQUHAwEG
CCsGAQUFBwMCMCIGA1UdIAQbMBkwDQYLKwYBBAGyMQECAkAwCAYGZ4EMAQIBMFAG
A1UdHwRJMEcwRaBDoEGGP2h0dHA6Ly9jcmwudXNlcnRydXN0LmNvbS9VU0VSVHJ1
c3RSU0FDZXJ0aWZpY2F0aW9uQXV0aG9yaXR5LmNybDB2BggrBgEFBQcBAQRqMGgw
PwYIKwYBBQUHMAKGM2h0dHA6Ly9jcnQudXNlcnRydXN0LmNvbS9VU0VSVHJ1c3RS
U0FBZGRUcnVzdENBLmNydDAlBggrBgEFBQcwAYYZaHR0cDovL29jc3AudXNlcnRy
dXN0LmNvbTANBgkqhkiG9w0BAQwFAAOCAgEAXXRDKHiA5DOhNKsztwayc8qtlK4q
Vt2XNdlzXn4RyZIsC9+SBi0Xd4vGDhFx6XX4N/fnxlUjdzNN/BYY1gS1xK66Uy3p
rw9qI8X12J4er9lNNhrsvOcjB8CT8FyvFu94j3Bs427uxcSukhYbERBAIN7MpWKl
VWxT3q8GIqiEYVKa/tfWAvnOMDDSKgRwMUtggr/IE77hekQm20p7e1BuJODf1Q7c
FPt7T74m3chg+qu0xheLI6HsUFuOxc7R5SQlkFvaVY5tmswfWpY+rwhyJW+FWNbT
uNXkxR4v5KOQPWrY100/QN68/j17paKuSXNcsr56snuB/Dx+MACLBdsF35HxPadx
78vkfQ37WcVmKZtHrHJQ/QUyjxdG8fezMsh0f+puUln/O+NlsFtipve8qYa9h/K5
yD0oZN93ChWve78XrV4vCpjO75Nk5B8O9CWQqGTHbhkgvjyb9v/B+sYJqB22/NLl
R4RPvbmqDJGeEI+4u6NJ5YiLIVVsX+dyfFP8zUbSsj6J34RyCYKBbQ4L+r7k8Srs
LY51WUFP292wkFDPSDmV7XsUNTDOZoQcBh2Fycf7xFfxeA+6ERx2d8MpPPND7yS2
1dkf+SY5SdpSbAKtYmbqb9q8cZUDEImNWJFUVHBLDOrnYhGwJudE3OBXRTxNhMDm
IXnjEeWrFvAZQhk=
-----END CERTIFICATE-----
"""

SEARCH_PER_PAGE = 15
LIST_PER_PAGE = 16
PREFIX_MAX_PAGES = 4     # tran khi vot ma thieu gach bang tien to (4 x 15 = 60 ung vien)

CIRCUIT_BREAK = 25       # N loi LIEN TIEP -> nghi bi chan IP -> dung han
HIBERNATE_429 = 10       # N lan 403/429 LIEN TIEP -> ngu dai
MAX_ITEMS = 40000        # tran work-list (catalog that ~26.6k)
MAX_PAGES = 600          # tran phan trang
MAX_BFS_PAGES = 1500     # tran trang danh sach khi BFS
MAX_REDIRECTS = 5
TIMEOUT = 90.0           # server cham: latency do duoc toi 14s
GC_EVERY = 2000

COLUMNS = ["Input", "Item No", "Page", "Status",
           "Category 1", "Category 2", "Category 3", "Category 4",
           "Thuong hieu", "Model", "Gia", "Gia VAT", "Mo ta", "Thong so", "Image", "URL"]

TAG_RE = re.compile(r"<[^>]+>")
LOC_RE = re.compile(r"<loc>([^<]+)</loc>")
PROD_HREF_RE = re.compile(r'href="([^"?#]*?-\d+\.html)"', re.I)
SUB_HREF_RE = re.compile(r'href="[^"]*?(danh-sach/[^"?#]+\.html)"', re.I)
CAT_HREF_RE = re.compile(r'href="[^"]*?(danh-muc/[^"?#]+\.html)"', re.I)
# block SP: chan tai class ket thuc HOAC tai block ke tiep — neu chi chan bang sp_price/
# gia_sanpham_moi thi markup thieu o 1 SP se lam regex tran sang SP SAU -> ghep sai ma hang
LIST_BLOCK_RE = re.compile(r'class="sanpham"(.*?)(?=class="sp_price"|class="sanpham"|$)',
                           re.S | re.I)
SEARCH_BLOCK_RE = re.compile(r'class="one_productsearch"(.*?)'
                             r'(?=class="gia_sanpham_moi"|class="one_productsearch"|$)',
                             re.S | re.I)
HREF_IN_RE = re.compile(r'href="([^"?#]+)"')
MODEL_IN_LIST_RE = re.compile(r'class="sp_model">([^<]*)<', re.I)
MA_IN_SEARCH_RE = re.compile(r"Mã hàng\s*:\s*([^<]*)", re.I)
SEARCH_TOTAL_RE = re.compile(r"tìm được[^(]*\((\d+)")
TITLE_RE = re.compile(r"<title>(.*?)</title>", re.S | re.I)
CODE_RE = re.compile(r"Mã sản phẩm\s*:\s*<span[^>]*>([^<]+)", re.I)
PRICE_RE = re.compile(r'id="giasanpham">([^<]*)', re.I)
PRICE_VAT_RE = re.compile(r'id="giasanphamvat">([^<]*)', re.I)
BRAND_RE = re.compile(r'href="[^"]*(?:nha-san-xuat|nsx)/[^"]+\.html"[^>]*>([^<]{1,50})<', re.I)
CRUMB_RE = re.compile(r"(Trang chủ.*?)</div>", re.S)
IMG_RE = re.compile(r'<img[^>]+src="([^"]*products/[^"]+)"', re.I)
LINE_SPLIT_RE = re.compile(r"<br\s*/?>|</p>|</div>|</li>|</h\d>|</tr>", re.I)
KV_RE = re.compile(r"^[-–•*+]?\s*([^:]{2,45}?)\s*:\s*(\S.*)$")
SPEC_HEAD_RE = re.compile(r"THÔNG SỐ|Thông số|THONG SO", re.I)
# dong rac trong mo ta (banner/chinh sach/gia SP khac) — khong phai spec
SPEC_SKIP_RE = re.compile(r"^(giá|gia|liên hệ|lien he|hotline|email|website|địa chỉ|điện thoại|"
                          r"zalo|tel|fax|showroom|khuyến mãi)\b", re.I)


def ssl_ctx() -> ssl.SSLContext:
    ctx = ssl.create_default_context(cafile=certifi.where())
    ctx.load_verify_locations(cadata=CA_EXTRA)     # BAY #1: bu cert trung gian server thieu
    return ctx


# ================================================================ CORE (khong dung GUI)
def norm_code(s: Any) -> str:
    """Chuan hoa ma de so khop. URL thi GIU NGUYEN hoa/thuong (path phan biet chu)."""
    v = re.sub(r"\s+", "", str(s or ""))
    return v if v.lower().startswith("http") else v.upper()


def alnum(s: Any) -> str:
    """Bo dau/gach: TCK-120-035 == TCK120035 (mo ta hay viet thieu gach)."""
    return re.sub(r"[^A-Z0-9]", "", str(s or "").upper())


def text(s: str) -> str:
    return re.sub(r"\s+", " ", htmlmod.unescape(TAG_RE.sub(" ", s or "")).replace("\xa0", " ")).strip()


def clean_path(href: str) -> str:
    """href (co the '../x-1.html', '/x-1.html/junk/', URL day du) -> '/x-1.html'.
    BAY #2: bo duoi rac sau .html + loai /lien-he,... (form lien he, khong phai SP)."""
    p = str(href or "").strip()
    p = re.sub(r"^https?://[^/]+", "", p).split("?")[0].split("#")[0]
    p = re.sub(r"^(?:\.\./)+", "/", p)
    if not p.startswith("/"):
        p = "/" + p
    m = re.match(r"(/[^/]*?-\d+\.html)", p, re.I)
    return m.group(1) if m else p


def is_product_path(p: str) -> bool:
    return bool(re.fullmatch(r"/[^/]+-\d+\.html", p or "", re.I)) and not p.startswith("/lien-he,")


def product_url(p: str) -> str:
    return BASE + quote(clean_path(p))


OTHER_PRODUCTS_RE = re.compile(r"^[-*+\s]*(?:cac )?san pham (?:khac|lien quan|tuong tu)\s*(?:[:;-].*)?$")


def strip_other_products(txt: str) -> str:
    """Cat 'Thong so' tai dong tieu de gioi thieu SP KHAC ('Cac san pham khac:', 'San pham lien
    quan'...; khong phan biet dau/hoa thuong, co/khong dau ':' / ';' / '-', hoac 'San pham khac: <ten SP>'). Phan sau do la cross-sell cua
    shop, khong phai thong so cua SP nay. Dong 'Mau sac: khac' KHONG bi cat (phai la ca dong)."""
    lines = str(txt or "").split("\n")
    for i, ln in enumerate(lines):
        flat = unicodedata.normalize("NFD", ln.lower().replace("đ", "d"))
        flat = "".join(c for c in flat if not unicodedata.combining(c))
        if OTHER_PRODUCTS_RE.match(flat.strip()):
            return "\n".join(lines[:i]).rstrip()
    return str(txt or "")


def parse_desc(html: str) -> dict[str, Any]:
    """Vung mo ta -> {spec{}, bullets[], desc}. BAY #5: cat truoc carousel 'SP lien quan'
    (class="n-box"), khong thi hut gia cua SP khac vao spec."""
    i = html.find('class="mota-sp')
    if i < 0:
        return {"spec": {}, "bullets": [], "desc": ""}
    seg = html[html.find(">", i) + 1:]                 # bo phan attribute cua chinh the div
    for stop in ('class="n-box"', 'class="box_ct', 'class="product_row',
                 "Sản phẩm liên quan", "SẢN PHẨM LIÊN QUAN"):
        j = seg.find(stop)
        if j > 0:
            seg = seg[:j]
    seg = re.sub(r"<[^>]*$", "", seg)                  # bo the con do dang o cuoi doan cat
    lines = [text(x) for x in LINE_SPLIT_RE.split(seg)]
    lines = [x for x in lines if x]
    if lines and lines[0].startswith("Mô tả sản phẩm"):
        lines = lines[1:]
    k = next((n for n, x in enumerate(lines) if SPEC_HEAD_RE.search(x) and len(x) < 60), -1)
    spec_lines = lines[k + 1:] if k >= 0 else lines
    desc_lines = lines[:k] if k >= 0 else lines

    spec: dict[str, str] = {}
    bullets: list[str] = []
    for s in spec_lines:
        m = KV_RE.match(s)
        if m and not SPEC_SKIP_RE.match(m.group(1)):
            spec.setdefault(m.group(1).strip(), m.group(2).strip())
        elif k >= 0 and len(s) < 200 and not SPEC_SKIP_RE.match(s):
            bullets.append(s)
    return {"spec": spec, "bullets": bullets, "desc": "\n".join(desc_lines)[:4000]}


def parse_product(html: str, url: str, input_code: str = "") -> list[dict]:
    """1 trang SP -> 1 record. Moi o la scalar; spec gop 1 cot 'Thong so' (AGENTS 4.6)."""
    if not html or not is_product_path(clean_path(url)):
        return []
    body = re.sub(r"(?s)<(script|style).*?</\1>", " ", html)
    name = text(TITLE_RE.search(body).group(1)) if TITLE_RE.search(body) else ""
    m = CODE_RE.search(body)
    code = text(m.group(1)) if m else ""
    if not name and not code:
        return []
    info = parse_desc(body)
    crumbs = []
    mc = CRUMB_RE.search(body)
    if mc:
        crumbs = [x.strip() for x in text(mc.group(1)).split("/")]
        crumbs = [x for x in crumbs[1:] if x and x != name][:4]
    mb = BRAND_RE.search(body)
    brand = text(mb.group(1)) if mb else ""
    mi = IMG_RE.search(body)
    img = ""
    if mi:
        img = clean_path(mi.group(1)).rstrip("/")
        img = BASE + img if img.startswith("/") else img
    mp, mv = PRICE_RE.search(body), PRICE_VAT_RE.search(body)
    spec = info["spec"]
    rec = {
        "Input": input_code,
        "Item No": code,
        "Page": name,
        "Status": "OK",
        "Thuong hieu": brand,
        "Model": spec.get("Mã sản phẩm") or spec.get("Mã hàng") or "",
        "Gia": text(mp.group(1)) if mp else "Xin lien he",
        "Gia VAT": text(mv.group(1)) if mv else "",
        "Mo ta": info["desc"],
        "Thong so": strip_other_products(
            "\n".join([f"{k}: {v}" for k, v in spec.items()] + info["bullets"])),
        "Image": img,
        "URL": BASE + clean_path(url),
    }
    for n in range(4):
        rec[f"Category {n + 1}"] = crumbs[n] if n < len(crumbs) else ""
    return [rec]


def not_found(code: str, status: str = "Not Found") -> dict:
    return {"Input": code, "Item No": "" if code.lower().startswith("http") else code,
            "Status": status}


def match_code(recs: list[dict], code: str) -> list[dict]:
    """§4.3 + BAY §10-5: khong lay mu ket qua dau. So ma nhap voi ma he thong (Item No),
    ma hang trong spec (Model) va ten SP. Tier: exact -> bo dau/gach -> chua trong (danh dau)."""
    c, ca = norm_code(code), alnum(code)
    exact = [r for r in recs if c in {norm_code(r.get("Item No")), norm_code(r.get("Model"))}]
    if exact:
        return exact
    loose = [r for r in recs if ca and ca in {alnum(r.get("Item No")), alnum(r.get("Model"))}]
    for r in loose:
        r["Status"] = "OK (khop bo dau)"
    if loose:
        return loose
    sub = [r for r in recs if len(ca) >= 4
           and any(ca in alnum(r.get(k)) for k in ("Item No", "Model", "Page"))]
    for r in sub:
        r["Status"] = "OK (khop mot phan)"
    return sub


def export_excel(rows: Iterable[dict], path: pathlib.Path) -> int:
    """Spec gop 1 COT 'Thong so' (AGENTS 4.6). openpyxl write_only = stream tung dong."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("trungtamthietbi")
    widths = {"Page": 50, "Mo ta": 50, "Thong so": 60, "Image": 40, "URL": 40}
    for i, name in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 16)
    ws.append(COLUMNS)
    wrap = Alignment(wrap_text=True, vertical="top")
    n = 0
    for r in rows:
        line = []
        for name in COLUMNS:
            v = r.get(name, "")
            if v is None:
                v = ""
            if isinstance(v, str):
                v = ILLEGAL_CHARACTERS_RE.sub("", v)[:32000]      # tran 1 o Excel 32.767
            if name in ("Mo ta", "Thong so"):
                cell = WriteOnlyCell(ws, value=v)
                cell.alignment = wrap
                v = cell
            line.append(v)
        ws.append(line)
        n += 1
    wb.save(path)
    return n


class ClientPool:
    """Pool N httpx.AsyncClient, moi client 1 cookie jar (= 1 PHPSESSID rieng). Server PHP 5.3 khoa
    session: moi request chung PHPSESSID bi XEP HANG => KHONG BAO GIO ban 2 request song song qua
    cung 1 client. Moi get() muon 1 client tu hang doi, tra lai khi xong. Dung y nhu 1 client:
    `await pool.get(url)`."""

    def __init__(self, n: int) -> None:
        self.q: asyncio.Queue[httpx.AsyncClient] = asyncio.Queue()
        self.clients = [httpx.AsyncClient(verify=ssl_ctx(), headers=HEADERS, timeout=TIMEOUT,
                                          follow_redirects=True, max_redirects=MAX_REDIRECTS)
                        for _ in range(max(1, n))]
        for c in self.clients:
            self.q.put_nowait(c)

    async def get(self, url: str) -> httpx.Response:
        c = await self.q.get()
        try:
            return await c.get(url)
        finally:
            self.q.put_nowait(c)

    async def aclose(self) -> None:
        for c in self.clients:
            await c.aclose()


def migrate_legacy(legacy: pathlib.Path, new: pathlib.Path) -> bool:
    """Checkpoint cu nam trong folder site -> DOI (khong copy, khong xoa) sang out dir."""
    if legacy.exists() and not new.exists():
        new.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(legacy), str(new))
        return True
    return False


class Engine:
    """Crawler core — chay headless/CLI duoc, khong phu thuoc GUI."""

    def __init__(self, conc: int = 16, delay: float = 0.05, log=print,
                 stop_evt: threading.Event | None = None,
                 pause_evt: threading.Event | None = None) -> None:
        self.conc, self.delay, self._log = max(1, conc), max(0.0, delay), log
        self.stop_evt = stop_evt or threading.Event()
        self.pause_evt = pause_evt or threading.Event()
        self.fails = 0
        self.n_block = 0
        self.results: list[dict] = []
        self.keep = True                 # mode ALL: False -> record chi nam trong JSONL (RAM)
        self.rows = 0
        self.done = 0
        self.total = 0
        self.phase = ""                  # discovery (total con 0): GUI/CLI hien dang tim gi
        self.pages_seen = 0
        self.found = 0
        self.limit = 0
        self.key = ""
        self.done_inputs: set[str] = set()
        self._fh = None
        self._ckpt_loaded = False        # BAY §10-3: khong nap checkpoint 2 lan

    def log(self, msg: str) -> None:
        self._log(msg)
        if msg.startswith(("[ERR]", "[WARN]", "[STOP]", "[FATAL]")):
            try:
                LOG_FILE.parent.mkdir(parents=True, exist_ok=True)
                with LOG_FILE.open("a", encoding="utf-8") as f:
                    f.write(f"{time.strftime('%Y-%m-%d %H:%M:%S')} {msg}\n")
            except OSError:
                pass

    # ---- resume: JSONL sidecar (append+flush tung record; chay lai = doc lai, bo qua da xong)
    def ckpt(self, base: pathlib.Path | None = None) -> pathlib.Path:
        h = hashlib.md5(self.key.encode("utf-8")).hexdigest()[:8]
        return (base or OUT_DIR) / f"_raw_{re.sub(r'[^A-Za-z0-9]+', '_', self.key)[:40]}_{h}.jsonl"

    def _load_ckpt(self) -> set[str]:
        """Tra ve tap PATH SP da xong. Mode ALL (keep=False) khong nap record vao RAM."""
        if self._ckpt_loaded:
            return {clean_path(r.get("URL", "")) for r in self.results if r.get("URL")}
        self._ckpt_loaded = True
        p = self.ckpt()
        if migrate_legacy(self.ckpt(HERE), p):          # checkpoint cu o folder site -> out dir
            self.log(f"[SYS] migrate checkpoint {self.ckpt(HERE).name} -> {p}")
        seen: set[str] = set()
        if not p.exists():
            return seen
        for line in p.open(encoding="utf-8", errors="ignore"):
            try:
                r = json.loads(line)
            except Exception:
                continue                  # dong cuoi do dang khi bi kill -> bo
            if r.get("Status") == "Not Found":
                # MISS co the do loi mang tam (search that bai) -> lan sau PHAI THU LAI,
                # khong duoc coi la "da xong" roi bao khong co hang vinh vien
                continue
            self.rows += 1
            if r.get("Input"):
                self.done_inputs.add(norm_code(r["Input"]))
            if self.keep:
                self.results.append(r)
            if r.get("URL"):
                seen.add(clean_path(r["URL"]))
        if self.rows:
            self.log(f"[SYS] Resume: {self.rows} dong / {len(seen)} SP da xong -> bo qua")
        return seen

    def _emit(self, recs: list[dict]) -> None:
        for r in recs:
            self._fh.write(json.dumps(r, ensure_ascii=False) + "\n")
            if r.get("Input"):
                self.done_inputs.add(norm_code(r["Input"]))
        self._fh.flush()
        self.rows += len(recs)
        if self.keep:
            self.results.extend(recs)

    def iter_rows(self) -> Iterator[dict]:
        """Nguon cho export: RAM neu co, khong thi doc stream tu JSONL (mode ALL)."""
        # record cu trong checkpoint co the con cross-sell -> lam sach luc DOC (khong ghi lai file)
        def clean(r: dict) -> dict:
            if r.get("Thong so"):
                r["Thong so"] = strip_other_products(r["Thong so"])
            return r

        if self.results or self.keep:
            for r in self.results:
                yield clean(r)
            return
        p = self.ckpt()
        if p.exists():
            for line in p.open(encoding="utf-8", errors="ignore"):
                try:
                    yield clean(json.loads(line))
                except Exception:
                    continue

    def cleanup(self) -> None:
        """Xong sach -> xoa checkpoint. Bi Stop/crash -> GIU de lan sau resume."""
        if not self.key:
            return
        p = self.ckpt()
        if self._full():                 # --limit cham tran = crawl DANG DO -> giu de resume
            return
        if p.exists() and not self.stop_evt.is_set():
            p.unlink()

    def _cap(self, items: list) -> list:
        """AGENTS 9.3: vuot MAX_ITEMS -> cat + LOG (khong cat im lang)."""
        if len(items) > MAX_ITEMS:
            self.log(f"[WARN] work-list {len(items)} > MAX_ITEMS={MAX_ITEMS} -> cat con "
                     f"{MAX_ITEMS}, chay lai de lay tiep")
        return items[:MAX_ITEMS]

    def _stopped(self) -> bool:
        """Stop da bam -> log [STOP] (kem pha dang chay neu con o discovery) va bao True."""
        if not self.stop_evt.is_set():
            return False
        where = (f"o pha {self.phase} (trang {self.pages_seen}, {self.found} SP tim thay)"
                 if self.phase and not self.total else f"({self.done}/{self.total} SP)")
        self.log(f"[STOP] Da dung {where} — chay lai de tiep tuc")
        return True

    def _full(self) -> bool:
        return bool(self.limit) and self.rows >= self.limit

    async def _gate(self) -> bool:
        while self.pause_evt.is_set() and not self.stop_evt.is_set():
            await asyncio.sleep(0.2)
        return not self.stop_evt.is_set() and not self._full()

    async def fetch(self, cli: ClientPool, url: str,
                    tries: int = 4) -> tuple[str, str] | None:
        """-> (html, url cuoi sau redirect) hoac None. URL cuoi de bat BAY #4 (search redirect)."""
        err = ""
        for i in range(tries):
            if not await self._gate():
                return None
            try:
                r = await cli.get(url)
                if r.status_code in (403, 429, 503):          # WAF / rate-limit -> lui dan
                    self.n_block += 1
                    wait = 120 if self.n_block >= HIBERNATE_429 else min(2 ** i * 5, 60)
                    self.log(f"[WARN] {r.status_code} x{self.n_block} -> ngu {wait}s | {url[-70:]}")
                    await asyncio.sleep(wait)
                    err = str(r.status_code)
                    continue
                self.n_block = 0
                if r.status_code == 404:
                    self.fails = 0
                    return None
                r.raise_for_status()
                self.fails = 0
                return r.text, str(r.url)
            except (httpx.HTTPError, OSError) as e:
                err = type(e).__name__
                if i < tries - 1:
                    await asyncio.sleep(min(2 ** i, 20))
        self.fails += 1
        self.log(f"[ERR] {err} | {url[-70:]}")
        if self.fails >= CIRCUIT_BREAK:
            self.log(f"[STOP] {CIRCUIT_BREAK} loi LIEN TIEP - nghi bi chan IP. DUNG.")
            self.stop_evt.set()
        return None

    # ---- discovery
    async def _sitemap(self, cli: ClientPool) -> tuple[set[str], set[str]]:
        """sitemap.xml -> file con doc THAT tu sitemapindex (BAY §10-9).
        Tra ve (SP, trang danh-sach). Sitemap CU (BAY #3) nen SP chi de VA, nhung danh sach
        trang /danh-sach/ trong sitemap LA NGUON TOT NHAT: 261 trang, nhieu hon nhieu so voi
        BFS tu 11 trang danh muc (chi thay 57)."""
        got = await self.fetch(cli, f"{BASE}/sitemap.xml")
        if not got:
            return set(), set()
        subs = [u for u in LOC_RE.findall(got[0]) if "sitemap" in u.rsplit("/", 1)[-1]]
        sem = asyncio.Semaphore(self.conc)

        async def one(u: str) -> tuple[set[str], set[str]]:
            async with sem:
                g = await self.fetch(cli, u)
                if not g:
                    return set(), set()
                locs = LOC_RE.findall(g[0])
                prods = {p for p in (clean_path(x) for x in locs) if is_product_path(p)}
                # /danh-sach/<slug>-<id>[,<trang>].html -> gom ve trang goc
                lists = {re.sub(r",\d+\.html$", ".html", m.group(1))
                         for m in (re.search(r"(danh-sach/[^/?#]+\.html)", x) for x in locs) if m}
                return prods, lists

        prods: set[str] = set()
        lists: set[str] = set()
        for p, l in await asyncio.gather(*(one(s) for s in subs)):
            prods |= p
            lists |= l
        return prods, lists

    def _list_items(self, html: str) -> list[tuple[str, str]]:
        """Trang /danh-sach/ hoac /list/ -> [(path, ma hang)] trong LUOI SP (class=sanpham),
        khong quet ca trang (tranh menu/carousel)."""
        out = []
        for blk in LIST_BLOCK_RE.findall(html):
            mh = HREF_IN_RE.search(blk)
            if not mh:
                continue
            p = clean_path(mh.group(1))
            if not is_product_path(p):
                continue
            mm = MODEL_IN_LIST_RE.search(blk)
            out.append((p, text(mm.group(1)) if mm else ""))
        return out

    def _search_items(self, html: str) -> list[tuple[str, str]]:
        """Trang /search.html -> [(path, ma hang)] tu block one_productsearch."""
        out = []
        for blk in SEARCH_BLOCK_RE.findall(html):
            mh = HREF_IN_RE.search(blk)
            if not mh:
                continue
            p = clean_path(mh.group(1))
            if not is_product_path(p):
                continue
            mm = MA_IN_SEARCH_RE.search(blk)
            out.append((p, text(mm.group(1)) if mm else ""))
        return out

    async def _paged(self, cli: ClientPool, base_path: str, kind: str):
        """Async generator cac trang danh sach. base = trang 1, ',N.html' = trang N+1.
        Dung khi trang KHONG CO ITEM MOI (stuck-page) hoac rong."""
        seen: set[str] = set()
        for page in range(0, MAX_PAGES):
            if not await self._gate():
                return
            if kind == "search":
                url = f"{base_path}&p={page}.html" if page else base_path
            else:
                url = f"{base_path[:-5]},{page}.html" if page else base_path
            got = await self.fetch(cli, url)
            if not got:
                # loi tam (server cham/5xx) KHAC voi het trang -> phai log, khong im lang bo
                self.log(f"[WARN] dung phan trang o trang {page} vi tai loi -> CO THE SOT SP "
                         f"| {url[-60:]}")
                return
            items = self._search_items(got[0]) if kind == "search" else self._list_items(got[0])
            new = [(p, m) for p, m in items if p not in seen]
            if not new:
                return
            seen.update(p for p, _ in new)
            self.pages_seen += 1
            self.found += len(new)
            yield new
            if self.delay:
                await asyncio.sleep(self.delay)
        else:
            self.log(f"[WARN] phan trang dat MAX_PAGES={MAX_PAGES} -> dung, CO THE SOT SP | {base_path[-60:]}")

    async def _crawl_paths(self, cli: ClientPool, paths: list[str], seen: set[str],
                           inputs: dict[str, str] | None = None, nf: bool = False) -> None:
        """Fetch tung trang SP -> record. nf=True: miss van ghi dong 'Not Found' (link/excel)."""
        todo = self._cap([p for p in dict.fromkeys(paths) if p and p not in seen])
        self.total += len(todo)
        sem = asyncio.Semaphore(self.conc)

        async def one(p: str) -> None:
            async with sem:
                if not await self._gate():           # re-check SAU khi gianh slot
                    return
                got = await self.fetch(cli, product_url(p))
                if self.delay:
                    await asyncio.sleep(self.delay)
                if got is None and self.stop_evt.is_set():
                    return                           # Stop giua chung -> khong ghi oan
                # redirect guard: URL cuoi phai con la trang SP
                recs = parse_product(got[0], got[1], (inputs or {}).get(p, "")) if got else []
                if recs:
                    seen.add(p)
                    self._emit(recs)
                elif nf:
                    self._emit([not_found((inputs or {}).get(p, p))])
                self.done += 1
                if self.keep or self.done % 200 == 0:
                    self.log(f"[{'OK ' if recs else 'MISS'}] {self.done}/{self.total} "
                             f"+{len(recs)} | {p[:58]}")
                if self.done % GC_EVERY == 0:
                    gc.collect()

        try:
            await asyncio.gather(*(one(p) for p in todo))
        except (KeyboardInterrupt, Exception) as e:   # crash-proof: record da o trong JSONL
            self.log(f"[FATAL] {type(e).__name__}: {e}")

    async def run(self, mode: str, target: str = "", excel: str = "", sheet: str = "",
                  col: str = "", limit: int = 0) -> None:
        self.key = f"{mode}_{target or excel}_{sheet}_{col}"
        if mode == "all":
            self.keep, self.limit = False, limit
        cli = ClientPool(self.conc)          # N client cookie rieng (PHP session lock)
        try:
            seen = self._load_ckpt()
            self.ckpt().parent.mkdir(parents=True, exist_ok=True)     # cai moi: out dir chua co
            self._fh = self.ckpt().open("a", encoding="utf-8")
            try:
                if mode == "all":
                    await self._run_all(cli, seen, limit)

                elif mode in ("code", "excel"):
                    raw = [norm_code(target)] if mode == "code" else read_excel_col(excel, sheet, col)
                    raw = self._cap([c for c in dict.fromkeys(raw) if c])
                    links = [c for c in raw if c.lower().startswith("http")]
                    codes = [c for c in raw if c not in links]
                    # BAY §10-2: LINK crawl RIENG, KHONG ap bo loc `codes`
                    if links:
                        todo = [u for u in links if norm_code(u) not in self.done_inputs]
                        ps = [clean_path(u) for u in todo]
                        await self._crawl_paths(cli, ps, set(), inputs=dict(zip(ps, todo)), nf=True)
                    if codes:
                        await self._run_codes(cli, codes)

                elif mode == "keyword":
                    items: list[tuple[str, str]] = []
                    self.phase, self.pages_seen, self.found = "tim kiem", 0, 0
                    async for part in self._paged(cli, f"{BASE}/search.html?keywords={quote(target)}",
                                                  "search"):
                        items += part
                        self.log(f"[SYS] tim '{target}': trang {self.pages_seen}, {len(items)} SP")
                        if limit and len(items) >= limit:
                            break
                    ps = [p for p, _ in items][:limit or None]
                    self.log(f"[SYS] '{target}': lay {len(ps)} san pham")
                    await self._crawl_paths(cli, ps, seen, inputs={p: target for p in ps})

                elif mode == "link":
                    links = [u.strip() for u in target.splitlines()
                             if u.strip().lower().startswith("http")]
                    todo = [u for u in dict.fromkeys(links) if norm_code(u) not in self.done_inputs]
                    ps = [clean_path(u) for u in todo]
                    await self._crawl_paths(cli, ps, set(), inputs=dict(zip(ps, todo)), nf=True)

                else:
                    self.log(f"[ERR] mode khong hop le: {mode}")
            finally:
                self._fh.close()
                self.phase = ""
        finally:
            await cli.aclose()

    async def _run_all(self, cli: ClientPool, seen: set[str], limit: int) -> None:
        """Sitemap CU (BAY #3) => discovery chinh la BFS /viewall.html -> danh-muc -> danh-sach
        (+ phan trang), roi HOP voi sitemap de khong sot SP cu."""
        self.phase, self.pages_seen, self.found = "sitemap", 0, 0
        sm_prod, sm_lists = await self._sitemap(cli)
        self.phase = "danh muc"
        got = await self.fetch(cli, f"{BASE}/viewall.html")
        cats = sorted(set(CAT_HREF_RE.findall(got[0]))) if got else []
        sem = asyncio.Semaphore(self.conc)

        async def subs_of(cat: str) -> set[str]:
            async with sem:
                g = await self.fetch(cli, f"{BASE}/{cat}")
                return set(SUB_HREF_RE.findall(g[0])) if g else set()

        subs: set[str] = set(sm_lists)
        bfs_subs: set[str] = set()
        for part in await asyncio.gather(*(subs_of(c) for c in cats)):
            bfs_subs |= part
        subs |= bfs_subs
        self.log(f"[SYS] Sitemap: {len(sm_prod)} SP + {len(sm_lists)} trang danh sach | "
                 f"{len(cats)} danh muc -> {len(bfs_subs)} danh sach | tong quet {len(subs)}")

        paths: set[str] = set()
        npage = 0
        self.phase, self.pages_seen = "BFS danh sach", 0

        async def walk(sub: str) -> None:
            nonlocal npage
            async with sem:
                async for items in self._paged(cli, f"{BASE}/{sub}", "list"):
                    npage += 1
                    paths.update(p for p, _ in items)
                    self.found = len(paths)
                    if npage % 50 == 0:
                        self.log(f"[SYS] BFS: {npage} trang, {len(paths)} SP")
                    # --limit: du SP roi thi thoi quet tiep (khoi cho het discovery ~30 phut)
                    if limit and len(paths) >= limit:
                        return
                    if npage >= MAX_BFS_PAGES:
                        self.log(f"[WARN] BFS dat MAX_BFS_PAGES={MAX_BFS_PAGES} -> dung quet, CO THE SOT SP")
                        return
                    if not await self._gate():
                        return

        await asyncio.gather(*(walk(s) for s in sorted(subs)))
        if self._stopped():
            return                                   # Stop trong discovery: khong hop, khong crawl
        self.log(f"[SYS] BFS: {npage} trang danh sach -> {len(paths)} SP")
        if not (limit and len(paths) >= limit):
            self.log(f"[SYS] Hop voi sitemap: chi co trong sitemap {len(sm_prod - paths)} "
                     f"| chi co trong BFS {len(paths - sm_prod)}")
            paths |= sm_prod
        todo = sorted(paths)
        if limit:
            todo = todo[:limit]
        self.log(f"[SYS] Tong work-list: {len(todo)} SP (da co {len(seen)})")
        await self._crawl_paths(cli, todo, seen)
        if not self._stopped():
            self.log(f"[SYS] ALL xong: {len(seen)} SP | {self.rows} dong")

    async def _by_prefix(self, cli: ClientPool, code: str) -> list[dict]:
        """Vot ma viet THIEU GACH: site search khong tu bo gach ('TCK120035' -> 0 ket qua) nhung
        tra tien to chu dau ('TCK') thi ra. Luoi ket qua CO SAN ma hang -> so alnum tai day,
        chi fetch detail cua SP khop => chi ton them vai request va chi khi da miss."""
        pre = re.match(r"[A-Za-z]{2,6}", str(code or "").strip())
        ca = alnum(code)
        if not pre or not ca:
            return []
        page = 0
        async for items in self._paged(cli, f"{BASE}/search.html?keywords={quote(pre.group(0))}",
                                       "search"):
            page += 1
            for p, ma in items:
                if ca in (alnum(ma), alnum(p)):
                    g = await self.fetch(cli, product_url(p))
                    if g:
                        recs = match_code(parse_product(g[0], g[1], code), code)
                        if recs:
                            self.log(f"[HIT] vot bang tien to '{pre.group(0)}' | {code}")
                            return recs
            if page >= PREFIX_MAX_PAGES:
                break
        return []

    async def _run_codes(self, cli: ClientPool, codes: list[str]) -> None:
        """1 ma -> /search.html?keywords=<ma>. BAY #4: khop duy nhat thi site REDIRECT thang
        sang trang SP -> parse luon (1 request). Nhieu ket qua -> luoi da co MA HANG san,
        KHOP MA TRUOC (§4.3) roi moi fetch detail (server cham, tiet kiem request).
        Ma miss van co dong 'Not Found'. Dedup theo MA NHAP (BAY §10-6)."""
        todo = [c for c in codes if norm_code(c) not in self.done_inputs]
        if len(todo) < len(codes):
            self.log(f"[SYS] Resume: bo qua {len(codes) - len(todo)} ma da xong")
        self.total += len(todo)
        sem = asyncio.Semaphore(self.conc)

        async def one(c: str) -> None:
            async with sem:
                if not await self._gate():
                    return
                got = await self.fetch(cli, f"{BASE}/search.html?keywords={quote(c)}")
                recs: list[dict] = []
                nhit = 0
                if got:
                    html, final = got
                    if is_product_path(clean_path(final)):        # redirect thang sang SP
                        nhit = 1
                        recs = match_code(parse_product(html, final, c), c)
                    else:
                        items = self._search_items(html)
                        nhit = len(items)
                        ca = alnum(c)
                        pri = [p for p, ma in items if ca and (ca in alnum(ma) or ca in alnum(p))]
                        for p in (pri or [p for p, _ in items])[:SEARCH_PER_PAGE]:
                            g = await self.fetch(cli, product_url(p))
                            if g:
                                recs += match_code(parse_product(g[0], g[1], c), c)
                            if any(r["Status"] == "OK" for r in recs):
                                break
                if not any(r["Status"] == "OK" for r in recs):
                    recs += await self._by_prefix(cli, c)
                if self.stop_evt.is_set():
                    return                       # Stop giua chung -> khong ghi Not Found oan
                if any(r["Status"] == "OK" for r in recs):
                    recs = [r for r in recs if r["Status"] == "OK"]
                recs = list({r.get("URL"): r for r in recs}.values()) or [not_found(c)]
                self._emit(recs)
                self.done += 1
                ok = recs[0].get("Status") != "Not Found"
                self.log(f"[{'OK ' if ok else 'MISS'}] {self.done}/{self.total} "
                         f"+{len(recs) if ok else 0} | {c} ({nhit} hit)")
                if self.delay:
                    await asyncio.sleep(self.delay)

        try:
            await asyncio.gather(*(one(c) for c in todo))
        except (KeyboardInterrupt, Exception) as e:   # crash-proof: record da o trong JSONL
            self.log(f"[FATAL] {type(e).__name__}: {e}")


def read_excel_col(path: str, sheet: str, col: str) -> list[str]:
    """Loi -> ValueError (KHONG SystemExit: worker GUI se khong bat duoc -> treo)."""
    if not path:
        raise ValueError("Chua chon file Excel")
    df = pd.read_excel(path, sheet_name=sheet or 0, dtype=str)
    df.columns = df.columns.str.strip()
    if col not in df.columns:
        raise ValueError(f"Khong thay cot '{col}'. Cot co san: {list(df.columns)}")
    return [norm_code(v) for v in df[col].dropna().tolist()]


def final_msg(eng: Engine) -> str:
    """Dong cuoi GUI: XONG chi khi chay tron; Stop -> noi ro chua xong."""
    if eng.stop_evt.is_set():
        return (f"[STOP] DA DUNG — {eng.rows} dong (chua xong). Bam EXPORT de xuat phan da co; "
                f"chay lai de tiep tuc.")
    return f"[SYS] XONG — {eng.rows} dong. Bam EXPORT de xuat Excel."


def stat_text(eng: Engine) -> str:
    """Nhan tien do GUI: discovery (total=0) -> 'Dang tim ...'; co work-list -> done/total."""
    if eng.total:
        return f"{eng.done}/{eng.total} SP | {eng.rows} dong"
    if eng.phase:
        return f"Dang tim: {eng.phase} | trang {eng.pages_seen} | {eng.found} SP tim thay"
    return f"0/0 | {eng.rows} dong"


# ================================================================ GUI
def run_gui() -> None:
    import sys
    import tkinter as tk
    from tkinter import messagebox, ttk

    # ha tang GUI chung o root (di nguoc len cay thu muc -> song sot khi doi nhom folder)
    root_dir = next(p for p in (HERE, *HERE.parents) if (p / "gui_common.py").exists())
    sys.path.insert(0, str(root_dir))
    from gui_common import ask_export_path, attach_excel_picker

    root = tk.Tk()
    root.title("TRUNG TAM THIET BI Crawler — trungtamthietbi.com")
    root.geometry("1080x700")
    ttk.Style().theme_use("clam")

    q: queue.Queue[str] = queue.Queue()
    stop_evt, pause_evt = threading.Event(), threading.Event()
    state: dict[str, Any] = {"running": False, "eng": None}

    left = ttk.Frame(root, padding=10)
    left.pack(side="left", fill="y")
    right = ttk.Frame(root, padding=(0, 10, 10, 10))
    right.pack(side="right", fill="both", expand=True)

    mode = tk.StringVar(value="code")
    ttk.Label(left, text="Pipeline", font=("Segoe UI", 10, "bold")).pack(anchor="w")
    for val, txt in [("code", "1. Theo MA (ma he thong / ma hang)"),
                     ("keyword", "2. Theo TU KHOA"),
                     ("link", "3. Theo LINK (moi dong 1 link)"),
                     ("excel", "4. Theo cot EXCEL (ma hoac link)"),
                     ("all", "5. Crawl ALL (~26.6k SP, ~1.5 gio)")]:
        ttk.Radiobutton(left, text=txt, value=val, variable=mode).pack(anchor="w")

    ttk.Label(left, text="\nMa / Tu khoa:").pack(anchor="w")
    ent = ttk.Entry(left, width=36)
    ent.pack(anchor="w")
    ttk.Label(left, text="Link (moi dong 1 link):").pack(anchor="w")
    box = tk.Text(left, width=36, height=5)
    box.pack(anchor="w")

    xls, sheet, col = tk.StringVar(), tk.StringVar(), tk.StringVar()
    outdir = tk.StringVar(value="images")
    attach_excel_picker(left, xls, sheet, col, outdir)

    conc, delay, limit = tk.IntVar(value=16), tk.DoubleVar(value=0.05), tk.IntVar(value=0)
    for lbl, var in [("Concurrency:", conc), ("Delay (s):", delay), ("Limit (0=all):", limit)]:
        f3 = ttk.Frame(left)
        f3.pack(anchor="w", pady=1)
        ttk.Label(f3, text=lbl, width=13).pack(side="left")
        ttk.Entry(f3, textvariable=var, width=8).pack(side="left")

    console = tk.Text(right, bg="#10151b", fg="#d6e2ee", insertbackground="#fff", wrap="none")
    console.pack(fill="both", expand=True)
    pb = ttk.Progressbar(right, mode="determinate")
    pb.pack(fill="x", pady=4)
    stat = ttk.Label(right, text="0/0")
    stat.pack(anchor="w")

    def drain() -> None:
        for _ in range(500):
            if q.empty():
                break
            console.insert("end", q.get() + "\n")
        console.see("end")
        if int(console.index("end-1c").split(".")[0]) > 5000:      # console khong phinh RAM
            console.delete("1.0", "2000.0")
        eng: Engine | None = state["eng"]
        if eng and eng.total:
            if str(pb["mode"]) != "determinate":
                pb.stop()
                pb.config(mode="determinate")
            pb["maximum"], pb["value"] = eng.total, min(eng.done, eng.total)
        elif eng and eng.phase and str(pb["mode"]) != "indeterminate":
            pb.config(mode="indeterminate")          # discovery: chua biet tong -> thanh chay
            pb.start(60)
        elif eng and not eng.phase and str(pb["mode"]) == "indeterminate":
            pb.stop()                                # discovery xong ma khong co SP: thoi chay
            pb.config(mode="determinate")
        if eng:
            stat.config(text=stat_text(eng))
        root.after(100, drain)

    def worker() -> None:
        m = mode.get()
        target = box.get("1.0", "end").strip() if m == "link" else ent.get().strip()
        eng = Engine(conc.get(), delay.get(), q.put, stop_evt, pause_evt)
        state["eng"] = eng
        try:
            asyncio.run(eng.run(m, target, xls.get(), sheet.get(), col.get(), limit.get()))
        except BaseException as e:      # BaseException: SystemExit/KeyboardInterrupt cung phai bat
            q.put(f"[FATAL] {type(e).__name__}: {e}")
        finally:                        # finally: LUON nha co, khong khoa START vinh vien
            state["running"] = False
            q.put(final_msg(eng))

    def start() -> None:
        if state["running"]:
            messagebox.showinfo("Dang chay", "Dang crawl. Bam Stop truoc.")
            return
        m = mode.get()
        if m == "excel" and not (xls.get() and col.get()):
            messagebox.showwarning("Thieu du lieu", "Chon file Excel va cot truoc.")
            return
        if m in ("code", "keyword") and not ent.get().strip():
            messagebox.showwarning("Thieu du lieu", "Nhap ma / tu khoa truoc.")
            return
        if m == "link" and not box.get("1.0", "end").strip():
            messagebox.showwarning("Thieu du lieu", "Dan link vao o (moi dong 1 link).")
            return
        stop_evt.clear()
        pause_evt.clear()
        state["running"] = True
        console.delete("1.0", "end")
        threading.Thread(target=worker, daemon=True).start()

    def export() -> None:
        eng: Engine | None = state["eng"]
        if state["running"]:
            messagebox.showinfo("Export", "Dang crawl. Bam Stop (hoac doi xong) roi EXPORT.")
            return
        if not eng or not eng.rows:
            messagebox.showinfo("Export", "Chua co du lieu")
            return
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        p = ask_export_path(xls.get(), OUT_DIR, "trungtamthietbi.xlsx")
        if p:
            n = export_excel(eng.iter_rows(), pathlib.Path(p))
            eng.cleanup()               # da xuat xong -> xoa checkpoint (bi Stop thi GIU)
            messagebox.showinfo("Export", f"Da luu {n} dong:\n{p}")

    fb = ttk.Frame(left)
    fb.pack(anchor="w", pady=8)
    ttk.Button(fb, text="START", command=start).pack(side="left")
    ttk.Button(fb, text="Pause", command=lambda: (pause_evt.clear() if pause_evt.is_set()
                                                  else pause_evt.set())).pack(side="left")
    ttk.Button(fb, text="Stop", command=stop_evt.set).pack(side="left")
    ttk.Button(fb, text="EXPORT", command=export).pack(side="left")

    def on_close() -> None:
        if state["running"] and not messagebox.askyesno("Thoat", "Dang crawl. Thoat va tat luon?"):
            return
        stop_evt.set()
        try:
            root.destroy()
        finally:
            os._exit(0)      # kill cung: khong de thread asyncio/console treo

    root.protocol("WM_DELETE_WINDOW", on_close)
    drain()
    root.mainloop()


# ================================================================ CLI
def main() -> None:
    ap = argparse.ArgumentParser(description="Trung Tam Thiet Bi crawler (trungtamthietbi.com)")
    ap.add_argument("--mode", choices=["code", "keyword", "link", "excel", "all"])
    ap.add_argument("--code")
    ap.add_argument("--keyword")
    ap.add_argument("--link")
    ap.add_argument("--excel")
    ap.add_argument("--sheet", default="")
    ap.add_argument("--col", default="")
    ap.add_argument("--out", default=str(OUT_DIR / "trungtamthietbi.xlsx"))
    ap.add_argument("--conc", type=int, default=16)
    ap.add_argument("--delay", type=float, default=0.05)
    ap.add_argument("--limit", type=int, default=0)
    a = ap.parse_args()
    try:
        sys.stdout.reconfigure(line_buffering=True)   # `> file.log` thay tien do ngay, khong cho exit
    except (AttributeError, ValueError, OSError):
        pass

    if not a.mode:
        run_gui()
        return

    target = a.code or a.keyword or a.link or ""
    eng = Engine(a.conc, a.delay)
    t0 = time.time()
    try:
        asyncio.run(eng.run(a.mode, target, a.excel or "", a.sheet, a.col, a.limit))
    except KeyboardInterrupt:
        eng.stop_evt.set()               # Ctrl+C -> GIU checkpoint de chay lai la resume
        print("[STOP] Ctrl+C -> giu checkpoint, chay lai de resume")
    finally:
        if eng.rows:
            pathlib.Path(a.out).parent.mkdir(parents=True, exist_ok=True)
            n = export_excel(eng.iter_rows(), pathlib.Path(a.out))
            print(f"[SYS] {n} dong -> {a.out} ({time.time() - t0:.1f}s)")
        eng.cleanup()


if __name__ == "__main__":
    main()
