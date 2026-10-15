"""
super-mro.com — Sieu thi vat tu MRO (httpx + selectolax, KHONG browser, KHONG Selenium).

RECON (2026-07-13)
------------------
API      : KHONG co JSON API. Toan bo render server-side (search + category + detail deu la HTML).
robots   : "Disallow: /tim-kiem?key=" va "/admin/". Sitemap: /sitemap.xml (duoc phep).
           => Search VAN la duong chinh (user chon), nhung co FALLBACK sitemap-index khi bi chan.
sitemap  : /sitemap.xml -> sitemap-products-page-1..309.xml (200 URL/file) ~ 61.600 san pham
detail   : /san-pham/<slug>.html   (~587 KB/trang - NANG, crawl ALL ton ~36GB)
           - h1 trong div.product-summary__right
           - p.mb-17 dang "<b>Nhan: </b>Gia tri" -> Thuong hieu / Ma he thong / Model hang /
             Don vi / Xuat xu  (so luong nhan KHAC NHAU tuy SP)
           - div.price -> "Gia: 45,496d" hoac "Gia: Lien he" (+ gia gach ngang = gia niem yet)
           - div.product-content__technical > ul > li > span.title + span.content  = THONG SO
           - ld+json Product: sku (= Ma he thong), name, image
           - anh: img.super-mro.com/... (.webp)
search   : /tim-kiem?key=<ma>&page=N  -> HTML, link tuyet doi /san-pham/*.html, 20 KQ/trang.
           Tra duoc CA "Ma he thong" (HTP48 -> choi-than-makita-cb-303a) — ma nay KHONG nam
           trong slug, nen KHONG the doan URL tu ma => search la cach duy nhat tra ma nhanh.
fallback : search bi chan/403 -> tai sitemap (309 file, cache ra _sitemap_urls.txt) roi khop ma
           TRONG SLUG. Chi bat duoc ma co trong slug (vd 4R002-...), ma he thong kieu HTP48 se miss
           -> log ro, khong im lang.
gia      : CO gia VND cong khai (mot so SP la "Lien he").

Run
---
  super_mro.bat                                        -> GUI
  python super_mro.py --mode code --code 4R002
  python super_mro.py --mode excel --excel parts.xlsx --col Part
  python super_mro.py --mode keyword --keyword makita --limit 50
  python super_mro.py --mode link --link https://super-mro.com/san-pham/<slug>.html
  python super_mro.py --mode all [--limit 500]         -> quet sitemap (~61.6k SP, NANG)
  python super_mro.py --mode images --excel parts.xlsx --col Part --outdir out
  them --sitemap de ep di duong sitemap (bo qua search)
"""
from __future__ import annotations

import argparse
import asyncio
import gc
import json
import os
import pathlib
import queue
import re
import threading
import time
from typing import Any
from urllib.parse import quote, urljoin, urlparse

import httpx
import pandas as pd
from selectolax.lexbor import LexborHTMLParser

HERE = pathlib.Path(__file__).parent
BASE = "https://super-mro.com"
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/122.0 Safari/537.36")
SITEMAP_CACHE = HERE / "_sitemap_urls.txt"
SITEMAP_TTL_DAYS = 7      # cache cu hon -> tai lai, khong thi SP moi cua site khong bao gio thay

CIRCUIT_BREAK = 25        # N loi LIEN TIEP -> nghi bi chan IP -> dung han
MAX_ITEMS = 70000         # site co ~61.6k SP
MAX_PAGES = 600
MAX_REDIRECTS = 5
TIMEOUT = 40.0            # trang detail nang (~587KB)
GC_EVERY = 200

PROD_RE = re.compile(r"https://super-mro\.com/san-pham/[^\"'\s]+\.html")


# ================================================================ CORE (khong dung GUI)
def norm_code(s: Any) -> str:
    v = re.sub(r"\s+", "", str(s or ""))
    return v if v.lower().startswith("http") else v.upper()


def search_url(key: str, page: int = 1) -> str:
    return f"{BASE}/tim-kiem?key={quote(key)}" + (f"&page={page}" if page > 1 else "")


def parse_links(html: str) -> list[str]:
    """Link san pham trong trang search/category — giu thu tu, bo trung."""
    seen: set[str] = set()
    out: list[str] = []
    for u in PROD_RE.findall(html):
        if u not in seen:
            seen.add(u)
            out.append(u)
    return out


def _price(t: LexborHTMLParser) -> tuple[str, str]:
    """(gia ban, gia niem yet). Gia ban o <span class="fs-26"> ngay sau 'Gia:' ('Lien he' neu
    khong niem yet); gia gach ngang nam trong <p class="price"> -> KHONG duoc tron 2 cai."""
    box = t.css_first("div.price")
    if box is None:
        return "", ""
    sell_n = box.css_first("span.fs-26")
    listed_n = box.css_first("p.price span")
    clean = lambda n: re.sub(r"\s+", " ", n.text(strip=True)) if n else ""
    return clean(sell_n), clean(listed_n)


def parse_page(html: str, url: str) -> dict | None:
    t = LexborHTMLParser(html)
    right = t.css_first("div.product-summary__right")
    h1 = (right or t).css_first("h1")
    if h1 is None:
        return None
    name = re.sub(r"\s+", " ", h1.text(strip=True))

    # nhan "<b>Thuong hieu: </b>Makita" - so luong nhan khac nhau tuy SP
    info: dict[str, str] = {}
    for p in (right.css("p.mb-17") if right else []):
        b = p.css_first("b")
        if not b:
            continue
        k = b.text(strip=True).rstrip(":").strip()
        v = re.sub(r"\s+", " ", p.text(strip=True)[len(b.text(strip=True)):]).strip()
        if k and v:
            info[k] = v

    spec: dict[str, str] = {}
    for li in t.css("div.product-content__technical li"):
        k = li.css_first("span.title")
        v = li.css_first("span.content")
        if k and v and k.text(strip=True) and v.text(strip=True):
            spec[re.sub(r"\s+", " ", k.text(strip=True))] = re.sub(r"\s+", " ", v.text(strip=True))

    sku, images, cats = "", [], []
    for sc in t.css('script[type="application/ld+json"]'):
        try:
            d = json.loads(sc.text())
        except Exception:
            continue
        for blk in (d if isinstance(d, list) else [d]):
            if not isinstance(blk, dict):
                continue
            if blk.get("@type") == "Product":
                sku = sku or (blk.get("sku") or "")
                img = blk.get("image")
                for u in (img if isinstance(img, list) else [img]):
                    if u and u not in images:
                        images.append(u)
    # site KHONG co ld+json BreadcrumbList -> lay tu HTML (bo 'Trang chu' va ten SP o cuoi)
    crumbs = [a.text(strip=True) for a in t.css(".breadcrum a, .breadcrumb a")]
    cats = [c for c in crumbs[1:-1] if c]
    for im in t.css("div.product-summary__left img, .mySlides img"):
        u = im.attributes.get("src") or ""
        if u and urljoin(url, u) not in images:
            images.append(urljoin(url, u))

    sell, listed = _price(t)
    return {"Input": "", "Ma": sku or info.get("Mã hệ thống", ""), "Ten": name,
            "Model hang": info.get("Model hãng", ""), "Thuong hieu": info.get("Thương hiệu", ""),
            "Don vi": info.get("Đơn vị", ""), "Xuat xu": info.get("Xuất xứ", ""),
            "Gia": sell, "Gia niem yet": listed, "Status": "OK",
            "_cats": [c for c in cats if c], "_spec": spec,
            "Image": " | ".join(images[:5]), "URL": url}


def export_excel(results: list[dict], path: pathlib.Path) -> pathlib.Path:
    """Spec gop 1 COT 'Thong so' (moi dong 'Ten: Gia tri'). Xem AGENTS.md 4.6."""
    from openpyxl.utils import get_column_letter

    fixed = ("Input", "Ma", "Ten", "Model hang", "Thuong hieu", "Don vi", "Xuat xu",
             "Gia", "Gia niem yet", "Status")
    max_cat = max((len(r.get("_cats") or []) for r in results), default=0)
    rows = []
    for r in results:
        row = {k: r.get(k, "") for k in fixed}
        cats = r.get("_cats") or []
        for i in range(max_cat):
            row[f"Category {i + 1}"] = cats[i] if i < len(cats) else ""
        row["Thong so"] = "\n".join(f"{k}: {v}" for k, v in (r.get("_spec") or {}).items())
        row["Image"] = r.get("Image", "")
        row["URL"] = r.get("URL", "")
        rows.append(row)
    df = pd.DataFrame(rows)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, index=False)
        ws = xw.sheets["Sheet1"]
        col = get_column_letter(df.columns.get_loc("Thong so") + 1)
        ws.column_dimensions[col].width = 60
        for c in ws[col][1:]:
            c.alignment = c.alignment.copy(wrap_text=True, vertical="top")
    return path


class Engine:
    """Search la duong chinh (tra duoc Ma he thong). Search chan -> fallback sitemap-index."""

    def __init__(self, conc: int = 6, delay: float = 0.15, log=print, force_sitemap: bool = False,
                 stop_evt: threading.Event | None = None,
                 pause_evt: threading.Event | None = None) -> None:
        self.conc, self.delay, self.log = max(1, conc), max(0.0, delay), log
        self.stop_evt = stop_evt or threading.Event()
        self.pause_evt = pause_evt or threading.Event()
        self.search_ok = not force_sitemap
        self.fails = 0
        self.results: list[dict] = []
        self.seen_urls: set[str] = set()
        self.sitemap: list[str] = []
        self.done = 0
        self.total = 0
        self.key = ""
        self._ckpt_loaded = False

    # ---- resume: JSONL sidecar
    def ckpt(self, key: str) -> pathlib.Path:
        return HERE / f"_raw_{re.sub(r'[^A-Za-z0-9]+', '_', key)[:40]}.jsonl"

    def _load_ckpt(self, key: str) -> set[str]:
        p, seen = self.ckpt(key), set()
        if self._ckpt_loaded:                 # goi 2 lan (link + code) -> KHONG nap lai
            return set(self.seen_urls)
        self._ckpt_loaded = True
        if not p.exists():
            return seen
        for line in p.read_text("utf-8", errors="ignore").splitlines():
            try:
                r = json.loads(line)
            except Exception:
                continue
            self.results.append(r)
            seen.add(r.get("URL", ""))
            self.seen_urls.add(r.get("URL", ""))
        if seen:
            self.log(f"[SYS] Resume: {len(self.results)} record da co -> bo qua")
        return seen

    def add(self, recs: list[dict], key: str) -> None:
        fresh = [r for r in recs if r.get("URL") and r["URL"] not in self.seen_urls]
        for r in fresh:
            self.seen_urls.add(r["URL"])
        if fresh:
            self.results += fresh
            with self.ckpt(key).open("a", encoding="utf-8") as f:
                for r in fresh:
                    f.write(json.dumps(r, ensure_ascii=False) + "\n")
                f.flush()

    def cleanup(self) -> None:
        p = self.ckpt(self.key) if self.key else None
        if p and p.exists() and not self.stop_evt.is_set():
            p.unlink()

    async def _gate(self) -> bool:
        while self.pause_evt.is_set() and not self.stop_evt.is_set():
            await asyncio.sleep(0.2)
        return not self.stop_evt.is_set()

    async def fetch(self, cli: httpx.AsyncClient, url: str, tries: int = 4) -> httpx.Response | None:
        for i in range(tries):
            if not await self._gate():
                return None
            try:
                r = await cli.get(url)
                if r.status_code in (403, 429, 503):
                    if "/tim-kiem" in url and self.search_ok:   # search bi chan -> doi duong
                        self.log(f"[WARN] search tra {r.status_code} -> CHUYEN sang sitemap-index")
                        self.search_ok = False
                        return None
                    wait = min(2 ** i * 5, 60)
                    self.log(f"[WARN] {r.status_code} -> ngu {wait}s | {url[-55:]}")
                    await asyncio.sleep(wait)
                    continue
                if r.status_code == 404:
                    return None
                r.raise_for_status()
                self.fails = 0
                return r
            except Exception as e:
                if i < tries - 1:
                    await asyncio.sleep(min(2 ** i, 20))
                    continue
                self.fails += 1
                self.log(f"[ERR] {type(e).__name__} | {url[-55:]}")
                if self.fails >= CIRCUIT_BREAK:
                    self.log(f"[STOP] {CIRCUIT_BREAK} loi LIEN TIEP - nghi bi chan IP. DUNG.")
                    self.stop_evt.set()
                return None
        return None

    # ---- sitemap: 309 file, cache ra dia (dung cho mode all va fallback tra ma)
    async def load_sitemap(self, cli: httpx.AsyncClient) -> list[str]:
        if self.sitemap:
            return self.sitemap
        if SITEMAP_CACHE.exists():
            age = (time.time() - SITEMAP_CACHE.stat().st_mtime) / 86400
            if age < SITEMAP_TTL_DAYS:
                self.sitemap = SITEMAP_CACHE.read_text("utf-8").split()
                self.log(f"[SYS] Sitemap cache ({age:.1f} ngay): {len(self.sitemap)} URL")
                return self.sitemap
            self.log(f"[SYS] Cache sitemap cu {age:.0f} ngay (>{SITEMAP_TTL_DAYS}) -> tai lai")
        idx = await self.fetch(cli, f"{BASE}/sitemap.xml")
        if idx is None:
            return []
        sms = [u for u in re.findall(r"<loc>([^<]+)</loc>", idx.text) if "sitemap-products-page" in u]
        self.log(f"[SYS] {len(sms)} file sitemap san pham -> dang tai...")
        sem = asyncio.Semaphore(self.conc)

        async def one(sm: str) -> list[str] | None:
            async with sem:
                if not await self._gate():
                    return None
                r = await self.fetch(cli, sm)
                return re.findall(r"<loc>([^<]+)</loc>", r.text) if r else None

        parts = await asyncio.gather(*(one(s) for s in sms))
        urls = [u for part in parts if part for u in part if "/san-pham/" in u]
        self.sitemap = list(dict.fromkeys(urls))
        n_fail = sum(p is None for p in parts)
        if n_fail:
            # tai khong het (loi/timeout/nguoi dung bam Stop) -> KHONG luu cache,
            # nhung van tra ket qua tam thoi cho lan chay nay (giong ingco.py)
            self.log(f"[WARN] Sitemap-index thieu {n_fail}/{len(sms)} file (loi hoac bi dung) - "
                     f"KHONG luu cache, lan sau se tai lai. Tam thoi dung {len(self.sitemap)} URL.")
        else:
            SITEMAP_CACHE.write_text("\n".join(self.sitemap), "utf-8")
            self.log(f"[SYS] Sitemap: {len(self.sitemap)} URL (cache -> {SITEMAP_CACHE.name})")
        return self.sitemap

    async def _resolve_code(self, cli: httpx.AsyncClient, code: str) -> dict | None:
        """Ma -> RECORD. Search co the tra ket qua gan dung -> phai KHOP MA CHINH XAC
        (AGENTS.md 4.3): mo tung hit, so 'Ma he thong' that; lech het thi lay hit dau
        nhung DANH DAU ro trong Status de nguoi doc biet ma khong khop."""
        urls: list[str] = []
        if self.search_ok:
            r = await self.fetch(cli, search_url(code))
            urls = (parse_links(r.text) if r else [])[:5]
        if not self.search_ok:                      # fallback: chi bat duoc ma NAM TRONG slug
            urls = [u for u in await self.load_sitemap(cli)
                    if re.search(rf"/san-pham/{re.escape(code)}[-.]", u, re.I)][:1]
            if not urls:
                self.log(f"[MISS] fallback sitemap khong khop '{code}' "
                         f"(ma he thong khong nam trong slug)")
        first: dict | None = None
        for u in urls:
            r = await self.fetch(cli, u)
            rec = parse_page(r.text, str(r.url)) if r else None
            if rec is None:
                continue
            if norm_code(rec["Ma"]) == norm_code(code) or norm_code(rec["Model hang"]) == norm_code(code):
                rec["Input"] = code
                return rec
            first = first or rec
        if first:
            first["Input"] = code
            first["Status"] = f"OK (ma lech: trang tra ve '{first['Ma']}')"
            self.log(f"[WARN] '{code}' khong khop chinh xac -> lay '{first['Ma']}' (da danh dau)")
        return first

    async def _crawl_urls(self, cli: httpx.AsyncClient, urls: list[str], key: str,
                          inputs: dict[str, str] | None = None) -> None:
        seen = self._load_ckpt(key)
        urls = [u for u in dict.fromkeys(urls) if u not in seen][:MAX_ITEMS]
        self.total, self.done = len(urls) + len(seen), len(seen)
        sem = asyncio.Semaphore(self.conc)

        async def one(u: str) -> None:
            async with sem:
                if not await self._gate():
                    return
                r = await self.fetch(cli, u)
                await asyncio.sleep(self.delay)
                rec = parse_page(r.text, str(r.url)) if r else None
                if rec:
                    rec["Input"] = (inputs or {}).get(u, "") or rec["Ma"] or u
                    self.add([rec], key)
                self.done += 1
                self.log(f"[{'OK ' if rec else 'MISS'}] {self.done}/{self.total} | {u[-52:]}")
                if self.done % GC_EVERY == 0:
                    gc.collect()

        try:
            await asyncio.gather(*(one(u) for u in urls))
        except (KeyboardInterrupt, Exception) as e:
            self.log(f"[FATAL] {type(e).__name__}: {e}")

    async def run(self, mode: str, target: str = "", excel: str = "", sheet: str = "",
                  col: str = "", limit: int = 0, outdir: str = "out") -> list[dict]:
        key = self.key = f"{mode}_{target or excel}_{col}"
        headers = {"User-Agent": UA, "Accept-Language": "vi-VN,vi;q=0.9,en;q=0.8"}
        async with httpx.AsyncClient(headers=headers, timeout=TIMEOUT, follow_redirects=True,
                                     max_redirects=MAX_REDIRECTS,
                                     limits=httpx.Limits(max_connections=self.conc + 4)) as cli:
            if mode == "all":
                urls = (await self.load_sitemap(cli))[:limit or MAX_ITEMS]
                self.log(f"[SYS] Crawl ALL: {len(urls)} trang (moi trang ~587KB - se ton bang thong)")
                await self._crawl_urls(cli, urls, key)

            elif mode in ("code", "excel"):
                raw = [norm_code(target)] if mode == "code" else read_excel_col(excel, sheet, col)
                raw = [c for c in dict.fromkeys(raw) if c][:MAX_ITEMS]
                links = [c for c in raw if c.lower().startswith("http")]
                codes = [c for c in raw if c not in links]
                if links:
                    await self._crawl_urls(cli, links, key, inputs={u: u for u in links})
                self._load_ckpt(key)
                self.total, self.done = len(codes) + self.done, self.done
                sem = asyncio.Semaphore(self.conc)

                async def resolve(c: str) -> None:
                    async with sem:
                        if not await self._gate():
                            return
                        rec = await self._resolve_code(cli, c)   # da parse san -> khong tai lai
                        await asyncio.sleep(self.delay)
                        self.done += 1
                        if rec:
                            self.add([rec], key)
                            self.log(f"[OK ] {self.done}/{self.total} {c} -> {rec['Ten'][:45]}")
                        else:
                            self.log(f"[MISS] {self.done}/{self.total} khong tra ra ma: {c}")

                await asyncio.gather(*(resolve(c) for c in codes))
                # doi chieu ca Input/Ma/Model hang: 1 SP co the da vao ket qua qua dong LINK
                # (dedup theo URL) -> neu chi so Input se bao "Not Found" oan cho ma cua no
                have = {norm_code(v) for r in self.results
                        for v in (r.get("Input"), r.get("Ma"), r.get("Model hang")) if v}
                for c in codes:                    # ma miss van co dong Not Found
                    if c not in have:
                        self.results.append({"Input": c, "Ma": c, "Ten": "", "Model hang": "",
                                             "Thuong hieu": "", "Don vi": "", "Xuat xu": "",
                                             "Gia": "", "Gia niem yet": "", "Status": "Not Found",
                                             "_cats": [], "_spec": {}, "Image": "", "URL": ""})

            elif mode == "keyword":
                urls, page = [], 1
                while page <= MAX_PAGES and await self._gate():
                    r = await self.fetch(cli, search_url(target, page))
                    new = [u for u in (parse_links(r.text) if r else []) if u not in urls]
                    if not new:                    # stuck-page = het ket qua
                        break
                    urls += new
                    self.log(f"[SYS] '{target}' trang {page} -> {len(urls)} URL")
                    if limit and len(urls) >= limit:
                        break
                    page += 1
                    await asyncio.sleep(self.delay)
                if not urls and not self.search_ok:
                    self.log("[ERR] Search bi chan -> mode keyword khong chay duoc. Dung 'code'/'all'.")
                await self._crawl_urls(cli, urls[:limit or MAX_ITEMS], key)

            elif mode == "link":
                await self._crawl_urls(
                    cli, [u.strip() for u in target.splitlines() if u.strip().startswith("http")], key)

            elif mode == "images":
                await self._download_images(cli, excel, sheet, col, outdir)

            else:
                self.log(f"[ERR] mode khong hop le: {mode}")
        return self.results

    async def _download_images(self, cli: httpx.AsyncClient, excel: str, sheet: str,
                               col: str, outdir: str) -> None:
        root = pathlib.Path(outdir) if os.path.isabs(outdir) else HERE / outdir
        root.mkdir(parents=True, exist_ok=True)
        keys = read_excel_col(excel, sheet, col)[:MAX_ITEMS]
        self.total = len(keys)
        sem = asyncio.Semaphore(self.conc)

        async def one(k: str) -> None:
            async with sem:
                if not await self._gate():
                    return
                if k.lower().startswith("http"):
                    r = await self.fetch(cli, k)
                    rec = parse_page(r.text, str(r.url)) if r else None
                else:
                    rec = await self._resolve_code(cli, k)
                self.done += 1
                if not rec or not rec["Image"]:
                    self.log(f"[MISS] {self.done}/{self.total} khong ra anh: {k}")
                    return
                name = re.sub(r'[\\/*?:"<>|]', "_", rec["Ma"] or k)[:60]
                folder = root / name
                folder.mkdir(parents=True, exist_ok=True)
                ok = 0
                for i, iu in enumerate(rec["Image"].split(" | "), 1):
                    try:
                        rr = await cli.get(iu, headers={"Referer": rec["URL"]})
                        rr.raise_for_status()
                        ext = pathlib.Path(urlparse(iu).path).suffix or ".jpg"
                        (folder / f"{name}_{i}{ext}").write_bytes(rr.content)
                        ok += 1
                    except Exception as e:
                        self.log(f"[ERR] tai anh {iu[-45:]}: {type(e).__name__}")
                self.log(f"[OK ] {self.done}/{self.total} {name} -> {ok} anh")

        await asyncio.gather(*(one(k) for k in keys))


def read_excel_col(path: str, sheet: str, col: str) -> list[str]:
    if not path:
        raise ValueError("Chua chon file Excel")
    df = pd.read_excel(path, sheet_name=sheet or 0, dtype=str)
    df.columns = df.columns.str.strip()
    if col not in df.columns:
        # ValueError chu KHONG SystemExit: SystemExit la BaseException, worker GUI bat
        # `except Exception` se KHONG bat duoc -> thread chet cam, GUI treo running=True
        raise ValueError(f"Khong thay cot '{col}'. Cot co san: {list(df.columns)}")
    return [norm_code(v) for v in df[col].dropna().tolist()]


# ================================================================ GUI
def run_gui() -> None:
    import sys
    import tkinter as tk
    from tkinter import messagebox, ttk

    # ha tang GUI chung o root (di nguoc len cay thu muc -> song sot khi doi nhom folder)
    root_dir = next(p for p in (HERE, *HERE.parents) if (p / "gui_common.py").exists())
    sys.path.insert(0, str(root_dir))
    from gui_common import ask_export_path, attach_excel_picker

    root = tk.Tk()
    root.title("SUPER-MRO Crawler — super-mro.com")
    root.geometry("1080x680")
    ttk.Style().theme_use("clam")

    q: queue.Queue[str] = queue.Queue()
    stop_evt, pause_evt = threading.Event(), threading.Event()
    state: dict[str, Any] = {"running": False, "results": [], "eng": None}

    left = ttk.Frame(root, padding=10)
    left.pack(side="left", fill="y")
    right = ttk.Frame(root, padding=(0, 10, 10, 10))
    right.pack(side="right", fill="both", expand=True)

    mode = tk.StringVar(value="excel")
    ttk.Label(left, text="Pipeline", font=("Segoe UI", 10, "bold")).pack(anchor="w")
    for val, txt in [("code", "1. Theo MA"), ("keyword", "2. Theo TU KHOA"),
                     ("link", "3. Theo LINK (moi dong 1 link)"), ("excel", "4. Theo cot EXCEL"),
                     ("all", "5. Crawl ALL (~61.6k SP - NANG)"), ("images", "6. Tai ANH")]:
        ttk.Radiobutton(left, text=txt, value=val, variable=mode).pack(anchor="w")

    ttk.Label(left, text="\nMa / Tu khoa:").pack(anchor="w")
    ent = ttk.Entry(left, width=36)
    ent.pack(anchor="w")
    ttk.Label(left, text="Link (moi dong 1 link):").pack(anchor="w")
    box = tk.Text(left, width=36, height=5)
    box.pack(anchor="w")

    xls, sheet, col = tk.StringVar(), tk.StringVar(), tk.StringVar()
    outdir = tk.StringVar(value="out")
    attach_excel_picker(left, xls, sheet, col, outdir)

    conc, delay, limit = tk.IntVar(value=6), tk.DoubleVar(value=0.15), tk.IntVar(value=0)
    for lbl, var in [("Concurrency:", conc), ("Delay (s):", delay), ("Limit (0=all):", limit)]:
        f3 = ttk.Frame(left)
        f3.pack(anchor="w", pady=1)
        ttk.Label(f3, text=lbl, width=13).pack(side="left")
        ttk.Entry(f3, textvariable=var, width=8).pack(side="left")
    force_sm = tk.BooleanVar(value=False)
    ttk.Checkbutton(left, text="Ep di sitemap (bo qua search)", variable=force_sm).pack(anchor="w")

    console = tk.Text(right, bg="#12161c", fg="#d6e2ee", insertbackground="#fff", wrap="none")
    console.pack(fill="both", expand=True)
    pb = ttk.Progressbar(right, mode="determinate")
    pb.pack(fill="x", pady=4)
    stat = ttk.Label(right, text="0/0")
    stat.pack(anchor="w")

    def drain() -> None:
        while not q.empty():
            console.insert("end", q.get() + "\n")
            console.see("end")
        eng: Engine | None = state["eng"]
        if eng and eng.total:
            pb["maximum"], pb["value"] = eng.total, eng.done
            stat.config(text=f"{eng.done}/{eng.total} | {len(eng.results)} record")
        root.after(100, drain)

    def worker() -> None:
        m = mode.get()
        target = box.get("1.0", "end").strip() if m == "link" else ent.get().strip()
        eng = Engine(conc.get(), delay.get(), q.put, force_sm.get(), stop_evt, pause_evt)
        state["eng"] = eng
        try:
            asyncio.run(eng.run(m, target, xls.get(), sheet.get(), col.get(),
                                limit.get(), outdir.get()))
            eng.cleanup()
        except BaseException as e:     # ke ca SystemExit -> khong de thread chet cam
            q.put(f"[FATAL] {type(e).__name__}: {e}")
        finally:                       # LUON nha co, khong khoa START vinh vien
            state["results"] = eng.results
            state["running"] = False
            q.put(f"[SYS] XONG — {len(eng.results)} record. Bam EXPORT de xuat Excel.")

    def start() -> None:
        if state["running"]:
            return
        stop_evt.clear()
        pause_evt.clear()
        state["running"] = True
        console.delete("1.0", "end")
        threading.Thread(target=worker, daemon=True).start()

    def export() -> None:
        """Mac dinh luu CANH file Excel dau vao, ten <ten_goc>_ketqua.xlsx."""
        if not state["results"]:
            messagebox.showinfo("Export", "Chua co du lieu")
            return
        p = ask_export_path(xls.get(), HERE, "super_mro.xlsx")
        if p:
            export_excel(state["results"], pathlib.Path(p))
            messagebox.showinfo("Export", f"Da luu:\n{p}")

    fb = ttk.Frame(left)
    fb.pack(anchor="w", pady=8)
    ttk.Button(fb, text="START", command=start).pack(side="left")
    ttk.Button(fb, text="Pause", command=lambda: (pause_evt.clear() if pause_evt.is_set()
                                                  else pause_evt.set())).pack(side="left")
    ttk.Button(fb, text="Stop", command=stop_evt.set).pack(side="left")
    ttk.Button(fb, text="EXPORT", command=export).pack(side="left")

    def on_close() -> None:
        if state["running"] and not messagebox.askyesno("Thoat", "Dang crawl. Thoat va tat luon?"):
            return
        stop_evt.set()
        try:
            root.destroy()
        finally:
            os._exit(0)

    root.protocol("WM_DELETE_WINDOW", on_close)
    drain()
    root.mainloop()


# ================================================================ CLI
def main() -> None:
    ap = argparse.ArgumentParser(description="super-mro.com crawler")
    ap.add_argument("--mode", choices=["code", "keyword", "link", "excel", "all", "images"])
    ap.add_argument("--code")
    ap.add_argument("--keyword")
    ap.add_argument("--link")
    ap.add_argument("--excel")
    ap.add_argument("--sheet", default="")
    ap.add_argument("--col", default="Part")
    ap.add_argument("--outdir", default="out")
    ap.add_argument("--out", default="super_mro.xlsx")
    ap.add_argument("--sitemap", action="store_true", help="ep di sitemap, bo qua search")
    ap.add_argument("--conc", type=int, default=6)
    ap.add_argument("--delay", type=float, default=0.15)
    ap.add_argument("--limit", type=int, default=0)
    a = ap.parse_args()

    if not a.mode:
        run_gui()
        return

    eng = Engine(a.conc, a.delay, print, a.sitemap)
    t0 = time.time()
    try:
        asyncio.run(eng.run(a.mode, a.code or a.keyword or a.link or "", a.excel or "",
                            a.sheet, a.col, a.limit, a.outdir))
    finally:
        if eng.results and a.mode != "images":
            p = export_excel(eng.results, pathlib.Path(a.out))
            print(f"[SYS] {len(eng.results)} record -> {p} ({time.time() - t0:.1f}s)")
        eng.cleanup()


if __name__ == "__main__":
    main()
