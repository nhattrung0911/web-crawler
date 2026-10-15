"""
sieuthicongcu.com - Sapo crawler (httpx + selectolax, no browser).

RECON (2026-08-11)
------------------
- Nen tang: Sapo/Haravan kieu Shopify. Footer va robots cho thay public catalog JSON.
- API list: /products.json?limit=250&page=N -> public JSON, page 74 con du lieu, 75 rong.
- API detail: /products/<alias>.js -> public JSON 1 san pham, co variants/images/content.
- Search: /search?q=<term> tra HTML SSR; lay link trong .search-main .item_product_main.
- Sitemap: /sitemap.xml -> sitemap_products_1.xml, sitemap_products_2.xml; doi chieu URL.
- Spec: trong content HTML (table 2 cot, ul/li) va summary HTML (<strong>Ten:</strong> Gia tri).
- Variants: 1 PDP co the chua nhieu SKU -> xuat 1 dong / variant.
"""
from __future__ import annotations

import argparse
import asyncio
import html
import json
import os
import pathlib
import queue
import re
import sys
import threading
import time
from typing import Any
from urllib.parse import quote

import httpx
import pandas as pd
from selectolax.lexbor import LexborHTMLParser

HERE = pathlib.Path(__file__).parent
BASE = "https://www.sieuthicongcu.com"
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")
HEADERS = {"User-Agent": UA, "Accept-Language": "vi-VN,vi;q=0.9,en;q=0.8"}

PAGE_SIZE = 250
SEARCH_PER_PAGE = 24
CIRCUIT_BREAK = 25
MAX_ITEMS = 25000
MAX_PAGES = 100
MAX_REDIRECTS = 5
TIMEOUT = 45.0
BUILD_INDEX_FROM = 40

LOC_RE = re.compile(r"<loc>([^<]+)</loc>")


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", html.unescape(str(value or "")).replace("\xa0", " ")).strip()


def norm_code(value: Any) -> str:
    v = re.sub(r"\s+", "", str(value or ""))
    return v if v.lower().startswith("http") else v.upper()


def alias_from_url(url: str) -> str:
    path = re.sub(r"^https?://[^/]+", "", str(url or "").strip()).split("?", 1)[0].split("#", 1)[0]
    path = path.strip("/")
    if path.startswith("products/"):
        path = path[len("products/"):]
    return re.sub(r"\.(js|json)$", "", path)


def detail_url(alias: str) -> str:
    return f"{BASE}/products/{quote(alias)}.js"


def search_url(term: str, page: int = 1) -> str:
    return f"{BASE}/search?q={quote(term)}&page={page}"


def _spec_table(tree: LexborHTMLParser) -> dict[str, str]:
    spec: dict[str, str] = {}
    for tr in tree.css("table tr"):
        cells = tr.css("th,td")
        if len(cells) < 2:
            continue
        key = clean_text(cells[0].text()).rstrip(":")
        val = clean_text(cells[1].text())
        if key and val and key not in spec:
            spec[key] = val
    return spec


def _spec_list(tree: LexborHTMLParser) -> dict[str, str]:
    spec: dict[str, str] = {}
    for li in tree.css("li"):
        txt = clean_text(li.text(separator=" ", strip=True))
        if not txt or ":" not in txt:
            continue
        key, val = txt.split(":", 1)
        key = clean_text(key).rstrip(":")
        val = clean_text(val)
        if key and val and key not in spec:
            spec[key] = val
    return spec


def _spec_summary(summary_html: str) -> dict[str, str]:
    tree = LexborHTMLParser(summary_html or "")
    spec: dict[str, str] = {}
    for p in tree.css("p,li,div"):
        raw = clean_text(p.text(separator=" ", strip=True))
        if not raw or ":" not in raw:
            continue
        key, val = raw.split(":", 1)
        key = clean_text(key).rstrip(":")
        val = clean_text(val)
        if key and val and key not in spec:
            spec[key] = val
    return spec


def build_spec(product: dict) -> dict[str, str]:
    tree = LexborHTMLParser(product.get("content") or "")
    spec = _spec_table(tree)
    for key, val in _spec_list(tree).items():
        spec.setdefault(key, val)
    for key, val in _spec_summary(product.get("summary") or "").items():
        if norm_code(key) == "MPN":
            continue
        spec.setdefault(key, val)
    return spec


def _content_images(product: dict) -> list[str]:
    out = [clean_text(img) for img in (product.get("images") or []) if clean_text(img)]
    if not out and product.get("featured_image"):
        out.append(clean_text(product.get("featured_image")))
    if out:
        return list(dict.fromkeys(out))
    tree = LexborHTMLParser(product.get("content") or "")
    for node in tree.css("img"):
        src = clean_text(node.attributes.get("src"))
        if src:
            out.append(src if src.startswith("http") else f"https:{src}" if src.startswith("//") else src)
    return list(dict.fromkeys(out))


def parse_product(product: dict, input_value: str = "") -> list[dict]:
    if not isinstance(product, dict) or not product.get("alias"):
        return []
    spec = build_spec(product)
    summary_spec = _spec_summary(product.get("summary") or "")
    images = _content_images(product)
    categories = [clean_text(product.get("vendor")), clean_text(product.get("product_type"))]
    categories = [x for x in categories if x]
    variants = product.get("variants") or [{}]
    mpn = summary_spec.get("MPN") or spec.get("MPN") or ""
    rows: list[dict] = []
    base_url = f"{BASE}/{product['alias']}"
    for variant in variants:
        variant_image = variant.get("featured_image")
        if isinstance(variant_image, dict):
            variant_image = variant_image.get("src") or variant_image.get("url")
        img_list = [clean_text(variant_image)] if clean_text(variant_image) else []
        img_list.extend(images)
        variant_url = f"{base_url}?variant={variant['id']}" if variant.get("id") and len(variants) > 1 else base_url
        rows.append({
            "Input": input_value,
            "Item No": clean_text(variant.get("sku")),
            "Page": clean_text(product.get("name")),
            "Status": "OK",
            "_cats": categories,
            "_spec": spec,
            "Bien the": clean_text(variant.get("title")) if len(variants) > 1 else "",
            "MPN": mpn,
            "Barcode": clean_text(variant.get("barcode")),
            "Gia (VND)": int(variant.get("price") or product.get("price") or 0) or "",
            "Ton kho": variant.get("inventory_quantity", ""),
            "Con hang": "Co" if variant.get("available", product.get("available")) else "Het",
            "Image": " | ".join(x for x in dict.fromkeys(img_list) if x),
            "URL": variant_url,
        })
    return rows


def match_code(rows: list[dict], code: str) -> list[dict]:
    c = norm_code(code)
    exact = [row for row in rows if c in {
        norm_code(row.get("Item No")),
        norm_code(row.get("Barcode")),
        norm_code(row.get("MPN")),
    }]
    if exact:
        return exact
    patt = re.compile(rf"(?<![A-Z0-9]){re.escape(c)}(?![A-Z0-9])")
    out = []
    for row in rows:
        hay = " ".join(norm_code(row.get(k)) for k in ("Item No", "Barcode", "MPN", "Page", "URL"))
        if patt.search(hay):
            out.append(row)
    if out:
        for row in out:
            row["Status"] = "OK (khop mot phan)"
    return out


def not_found(input_value: str, status: str = "Not Found") -> dict:
    return {
        "Input": input_value,
        "Item No": input_value,
        "Page": "",
        "Status": status,
        "_cats": [],
        "_spec": {},
        "Bien the": "",
        "MPN": "",
        "Barcode": "",
        "Gia (VND)": "",
        "Ton kho": "",
        "Con hang": "",
        "Image": "",
        "URL": "",
    }


def export_excel(rows: list[dict], path: pathlib.Path) -> pathlib.Path:
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    fixed_head = ["Input", "Page", "Item No", "Status"]
    max_cat = max((len(row.get("_cats") or []) for row in rows), default=0)
    out = []
    for row in rows:
        item = {k: row.get(k, "") for k in fixed_head}
        cats = row.get("_cats") or []
        for i in range(max_cat):
            item[f"Category {i + 1}"] = cats[i] if i < len(cats) else ""
        for key in ("Bien the", "MPN", "Barcode", "Gia (VND)", "Ton kho", "Con hang"):
            item[key] = row.get(key, "")
        item["Thong so"] = "\n".join(f"{k}: {v}" for k, v in (row.get("_spec") or {}).items())
        item["Image"] = row.get("Image", "")
        item["URL"] = row.get("URL", "")
        out.append(item)
    df = pd.DataFrame(out)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False)
        ws = writer.sheets["Sheet1"]
        if "Thong so" in df.columns:
            col = get_column_letter(df.columns.get_loc("Thong so") + 1)
            ws.column_dimensions[col].width = 60
            align = Alignment(wrap_text=True, vertical="top")
            for cell in ws[col][1:]:
                cell.alignment = align
    return path


def read_excel_col(path: str, sheet: str, col: str) -> list[str]:
    if not path:
        raise ValueError("Chua chon file Excel")
    df = pd.read_excel(path, sheet_name=sheet or 0, dtype=str)
    df.columns = df.columns.map(str)
    if col not in df.columns:
        raise ValueError(f"Khong thay cot '{col}'")
    return [norm_code(v) for v in df[col].dropna().tolist()]


class Engine:
    def __init__(
        self,
        conc: int = 8,
        delay: float = 0.05,
        log=print,
        stop_evt: threading.Event | None = None,
        pause_evt: threading.Event | None = None,
    ) -> None:
        self.conc = max(1, conc)
        self.delay = max(0.0, delay)
        self.log = log
        self.stop_evt = stop_evt or threading.Event()
        self.pause_evt = pause_evt or threading.Event()
        self.results: list[dict] = []
        self.done = 0
        self.total = 0
        self.fails = 0
        self.key = ""
        self._ckpt_loaded = False
        self._detail_cache: dict[str, dict | None] = {}

    def ckpt(self, key: str) -> pathlib.Path:
        safe = re.sub(r"[^A-Za-z0-9]+", "_", key)[:50]
        return HERE / f"_raw_{safe}.jsonl"

    def _load_ckpt(self, key: str) -> set[str]:
        if self._ckpt_loaded:
            return {alias_from_url(row.get("URL", "")) for row in self.results if row.get("URL")}
        self._ckpt_loaded = True
        seen: set[str] = set()
        p = self.ckpt(key)
        if not p.exists():
            return seen
        for line in p.read_text("utf-8", errors="ignore").splitlines():
            try:
                row = json.loads(line)
            except Exception:
                continue
            if row.get("Status") != "OK":
                continue
            self.results.append(row)
            if row.get("URL"):
                seen.add(alias_from_url(row["URL"]))
        return seen

    def _save(self, key: str, rows: list[dict]) -> None:
        with self.ckpt(key).open("a", encoding="utf-8") as fh:
            for row in rows:
                fh.write(json.dumps(row, ensure_ascii=False) + "\n")
            fh.flush()

    def cleanup(self) -> None:
        p = self.ckpt(self.key) if self.key else None
        if p and p.exists() and not self.stop_evt.is_set():
            p.unlink()

    async def _gate(self) -> bool:
        while self.pause_evt.is_set() and not self.stop_evt.is_set():
            await asyncio.sleep(0.2)
        return not self.stop_evt.is_set()

    async def fetch(self, client: httpx.AsyncClient, url: str, *, json_mode: bool = False) -> Any:
        for i in range(4):
            if not await self._gate():
                return None
            try:
                resp = await client.get(url)
                if resp.status_code in (403, 429, 503):
                    wait = min(5 * (2 ** i), 60)
                    self.log(f"[WARN] {resp.status_code} -> ngu {wait}s | {url[-70:]}")
                    await asyncio.sleep(wait)
                    continue
                if resp.status_code == 404:
                    return None
                resp.raise_for_status()
                self.fails = 0
                return resp.json() if json_mode else resp.text
            except Exception as exc:
                if i < 3:
                    await asyncio.sleep(min(2 ** i, 20))
                    continue
                self.fails += 1
                self.log(f"[ERR] {type(exc).__name__} | {url[-70:]}")
                if self.fails >= CIRCUIT_BREAK:
                    self.log(f"[STOP] {CIRCUIT_BREAK} loi lien tiep - dung.")
                    self.stop_evt.set()
                return None
        return None

    async def fetch_detail(self, client: httpx.AsyncClient, alias: str) -> dict | None:
        if alias in self._detail_cache:
            return self._detail_cache[alias]
        data = await self.fetch(client, detail_url(alias), json_mode=True)
        self._detail_cache[alias] = data
        return data

    async def search_aliases(self, client: httpx.AsyncClient, term: str, limit: int = 0) -> list[str]:
        out: list[str] = []
        seen: set[str] = set()
        page = 1
        while page <= MAX_PAGES:
            html_text = await self.fetch(client, search_url(term, page))
            if not html_text:
                break
            tree = LexborHTMLParser(html_text)
            aliases = []
            for node in tree.css(".search-main .item_product_main a.image_thumb"):
                href = node.attributes.get("href", "")
                alias = alias_from_url(href)
                if alias and alias not in seen:
                    seen.add(alias)
                    aliases.append(alias)
            if not aliases:
                break
            out.extend(aliases)
            if limit and len(out) >= limit:
                return out[:limit]
            page += 1
            if self.delay:
                await asyncio.sleep(self.delay)
        return out

    async def all_products(self, client: httpx.AsyncClient, limit: int = 0) -> list[dict]:
        out: list[dict] = []
        seen_ids: set[int] = set()
        page = 1
        while page <= MAX_PAGES:
            data = await self.fetch(client, f"{BASE}/products.json?limit={PAGE_SIZE}&page={page}", json_mode=True)
            items = (data or {}).get("products") or []
            if not items:
                break
            new = [item for item in items if item.get("id") not in seen_ids]
            if not new:
                break
            for item in new:
                if item.get("id") is not None:
                    seen_ids.add(item["id"])
            out.extend(new)
            if limit and len(out) >= limit:
                return out[:limit]
            page += 1
            if self.delay:
                await asyncio.sleep(self.delay)
        return out

    async def sitemap_aliases(self, client: httpx.AsyncClient) -> list[str]:
        idx = await self.fetch(client, f"{BASE}/sitemap.xml")
        if not idx:
            return []
        submaps = [u for u in LOC_RE.findall(idx) if "sitemap_products" in u]
        aliases: list[str] = []
        for url in submaps:
            xml = await self.fetch(client, url)
            if not xml:
                continue
            aliases.extend(alias_from_url(u) for u in LOC_RE.findall(xml))
        return [a for a in dict.fromkeys(aliases) if a]

    async def _crawl_aliases(self, client: httpx.AsyncClient, aliases: list[str], inputs: dict[str, str]) -> None:
        todo = [a for a in dict.fromkeys(aliases) if a][:MAX_ITEMS]
        self.total += len(todo)
        seen_inputs = {norm_code(row.get("Input")) for row in self.results if row.get("Input")}
        sem = asyncio.Semaphore(self.conc)

        async def one(alias: str) -> None:
            async with sem:
                if not await self._gate():
                    return
                data = await self.fetch_detail(client, alias)
                rows = parse_product(data, inputs.get(alias, "")) if data else []
                if rows:
                    self.results.extend(rows)
                    self._save(self.key, rows)
                else:
                    inp = inputs.get(alias, alias)
                    if norm_code(inp) not in seen_inputs:
                        row = not_found(inp)
                        self.results.append(row)
                        self._save(self.key, [row])
                self.done += 1
                self.log(f"[{'OK ' if rows else 'MISS'}] {self.done}/{self.total} | {alias[:60]}")

        await asyncio.gather(*(one(alias) for alias in todo))

    async def _run_codes(self, client: httpx.AsyncClient, codes: list[str]) -> None:
        done_inputs = {norm_code(row.get("Input")) for row in self.results if row.get("Input")}
        todo = [code for code in codes if norm_code(code) not in done_inputs][:MAX_ITEMS]
        self.total += len(todo)
        if len(todo) >= BUILD_INDEX_FROM:
            products = await self.all_products(client)
            code_map: dict[str, list[dict]] = {}
            for product in products:
                for row in parse_product(product):
                    for key in filter(None, [row.get("Item No"), row.get("Barcode"), row.get("MPN")]):
                        code_map.setdefault(norm_code(key), []).append(row)
            for code in todo:
                rows = [dict(row, Input=code) for row in code_map.get(norm_code(code), [])]
                if rows:
                    self.results.extend(rows)
                    self._save(self.key, rows)
                else:
                    row = not_found(code)
                    self.results.append(row)
                    self._save(self.key, [row])
                self.done += 1
                self.log(f"[{'OK ' if rows else 'MISS'}] {self.done}/{self.total} | {code}")
            return
        sem = asyncio.Semaphore(self.conc)

        async def one(code: str) -> None:
            async with sem:
                aliases = await self.search_aliases(client, code, limit=SEARCH_PER_PAGE)
                rows: list[dict] = []
                for alias in aliases:
                    data = await self.fetch_detail(client, alias)
                    if not data:
                        continue
                    rows.extend(match_code(parse_product(data, code), code))
                if rows:
                    self.results.extend(rows)
                    self._save(self.key, rows)
                else:
                    row = not_found(code)
                    self.results.append(row)
                    self._save(self.key, [row])
                self.done += 1
                self.log(f"[{'OK ' if rows else 'MISS'}] {self.done}/{self.total} | {code}")

        await asyncio.gather(*(one(code) for code in todo))

    async def _run_all(self, client: httpx.AsyncClient, limit: int) -> None:
        seen = self._load_ckpt(self.key)  # da_ckpt_loaded roi nen chi tra lai set alias da co, khong doc lai file
        products = await self.all_products(client, limit=limit)
        if limit and len(products) >= limit:
            self.total = limit
            for product in products[:limit]:
                if product.get("alias") in seen:
                    continue
                rows = parse_product(product, "")
                if rows:
                    self.results.extend(rows)
                    self._save(self.key, rows)
                self.done += 1
            return
        sitemap_alias_set = set(await self.sitemap_aliases(client))
        products_by_alias = {p.get("alias"): p for p in products if p.get("alias")}
        aliases = list(products_by_alias)
        for alias in sitemap_alias_set:
            if alias not in products_by_alias:
                data = await self.fetch_detail(client, alias)
                if data and data.get("alias"):
                    products_by_alias[alias] = data
                    aliases.append(alias)
        todo = [a for a in aliases[:limit or None] if a not in seen]
        self.total = len(todo)
        for alias in todo:
            rows = parse_product(products_by_alias[alias], "")
            if rows:
                self.results.extend(rows)
                self._save(self.key, rows)
            self.done += 1

    async def run(
        self,
        mode: str,
        target: str = "",
        excel: str = "",
        sheet: str = "",
        col: str = "",
        limit: int = 0,
    ) -> list[dict]:
        self.key = f"{mode}_{target or excel}_{col}"
        self._load_ckpt(self.key)
        async with httpx.AsyncClient(
            headers=HEADERS,
            timeout=TIMEOUT,
            follow_redirects=True,
            max_redirects=MAX_REDIRECTS,
            limits=httpx.Limits(max_connections=self.conc + 4),
        ) as client:
            if mode == "all":
                await self._run_all(client, limit)
            elif mode in {"code", "excel"}:
                raw = [norm_code(target)] if mode == "code" else read_excel_col(excel, sheet, col)
                links = [v for v in raw if v.lower().startswith("http")]
                codes = [v for v in raw if v not in links]
                if links:
                    aliases = [alias_from_url(v) for v in links]
                    await self._crawl_aliases(client, aliases, {a: u for a, u in zip(aliases, links)})
                if codes:
                    await self._run_codes(client, codes)
            elif mode == "keyword":
                aliases = await self.search_aliases(client, target, limit=limit)
                await self._crawl_aliases(client, aliases, {a: target for a in aliases})
            elif mode == "link":
                links = [line.strip() for line in target.splitlines() if line.strip()]
                aliases = [alias_from_url(u) for u in links]
                await self._crawl_aliases(client, aliases, {a: u for a, u in zip(aliases, links)})
            else:
                raise ValueError(f"Mode khong hop le: {mode}")
        return self.results


def run_gui() -> None:
    import tkinter as tk
    from tkinter import messagebox, ttk

    root_dir = next(p for p in (HERE, *HERE.parents) if (p / "gui_common.py").exists())
    sys.path.insert(0, str(root_dir))
    from gui_common import ask_export_path, attach_excel_picker

    root = tk.Tk()
    root.title("Sieu Thi Cong Cu Crawler")
    root.geometry("1080x700")
    ttk.Style().theme_use("clam")

    q: queue.Queue[str] = queue.Queue()
    stop_evt = threading.Event()
    pause_evt = threading.Event()
    state: dict[str, Any] = {"running": False, "rows": [], "eng": None}

    left = ttk.Frame(root, padding=10)
    left.pack(side="left", fill="y")
    right = ttk.Frame(root, padding=(0, 10, 10, 10))
    right.pack(side="right", fill="both", expand=True)

    mode = tk.StringVar(value="code")
    ttk.Label(left, text="Pipeline", font=("Segoe UI", 10, "bold")).pack(anchor="w")
    for val, label in [
        ("code", "1. Theo MA"),
        ("keyword", "2. Theo TU KHOA"),
        ("link", "3. Theo LINK"),
        ("excel", "4. Theo EXCEL"),
        ("all", "5. Crawl ALL"),
    ]:
        ttk.Radiobutton(left, text=label, value=val, variable=mode).pack(anchor="w")

    ttk.Label(left, text="\nMa / Tu khoa:").pack(anchor="w")
    ent = ttk.Entry(left, width=36)
    ent.pack(anchor="w")
    ttk.Label(left, text="Link (moi dong 1 link):").pack(anchor="w")
    box = tk.Text(left, width=36, height=5)
    box.pack(anchor="w")

    xls, sheet, col = tk.StringVar(), tk.StringVar(), tk.StringVar()
    outdir = tk.StringVar(value="images")
    attach_excel_picker(left, xls, sheet, col, outdir)

    conc, delay, limit = tk.IntVar(value=8), tk.DoubleVar(value=0.05), tk.IntVar(value=0)
    for label, var in [("Concurrency:", conc), ("Delay (s):", delay), ("Limit (0=all):", limit)]:
        fr = ttk.Frame(left)
        fr.pack(anchor="w", pady=1)
        ttk.Label(fr, text=label, width=13).pack(side="left")
        ttk.Entry(fr, textvariable=var, width=8).pack(side="left")

    console = tk.Text(right, bg="#12161c", fg="#d6e2ee", insertbackground="#fff", wrap="none")
    console.pack(fill="both", expand=True)
    pb = ttk.Progressbar(right, mode="determinate")
    pb.pack(fill="x", pady=4)
    stat = ttk.Label(right, text="0/0")
    stat.pack(anchor="w")

    def drain() -> None:
        while not q.empty():
            console.insert("end", q.get() + "\n")
            console.see("end")
        eng: Engine | None = state["eng"]
        if eng and eng.total:
            pb["maximum"] = eng.total
            pb["value"] = eng.done
            stat.config(text=f"{eng.done}/{eng.total} | {len(eng.results)} dong")
        root.after(100, drain)

    def worker() -> None:
        target = box.get("1.0", "end").strip() if mode.get() == "link" else ent.get().strip()
        eng = Engine(conc.get(), delay.get(), q.put, stop_evt, pause_evt)
        state["eng"] = eng
        try:
            state["rows"] = asyncio.run(eng.run(mode.get(), target, xls.get(), sheet.get(), col.get(), limit.get()))
            eng.cleanup()
        except BaseException as exc:
            q.put(f"[FATAL] {type(exc).__name__}: {exc}")
        finally:
            state["running"] = False
            q.put(f"[SYS] XONG - {len(state['rows'])} dong. Bam EXPORT de xuat Excel.")

    def start() -> None:
        if state["running"]:
            return
        if mode.get() in {"code", "keyword"} and not ent.get().strip():
            messagebox.showwarning("Thieu du lieu", "Nhap ma hoac tu khoa truoc.")
            return
        if mode.get() == "link" and not box.get("1.0", "end").strip():
            messagebox.showwarning("Thieu du lieu", "Dan link vao o moi dong 1 link.")
            return
        if mode.get() == "excel" and not (xls.get() and col.get()):
            messagebox.showwarning("Thieu du lieu", "Chon file Excel va cot.")
            return
        stop_evt.clear()
        pause_evt.clear()
        state["running"] = True
        console.delete("1.0", "end")
        threading.Thread(target=worker, daemon=True).start()

    def export() -> None:
        if not state["rows"]:
            messagebox.showinfo("Export", "Chua co du lieu")
            return
        path = ask_export_path(xls.get(), HERE, "sieuthicongcu.xlsx")
        if path:
            export_excel(state["rows"], pathlib.Path(path))
            messagebox.showinfo("Export", f"Da luu:\n{path}")

    fr_btn = ttk.Frame(left)
    fr_btn.pack(anchor="w", pady=8)
    ttk.Button(fr_btn, text="START", command=start).pack(side="left")
    ttk.Button(fr_btn, text="Pause", command=lambda: (pause_evt.clear() if pause_evt.is_set() else pause_evt.set())).pack(side="left")
    ttk.Button(fr_btn, text="Stop", command=stop_evt.set).pack(side="left")
    ttk.Button(fr_btn, text="EXPORT", command=export).pack(side="left")

    def on_close() -> None:
        if state["running"] and not messagebox.askyesno("Thoat", "Dang crawl. Thoat va tat luon?"):
            return
        stop_evt.set()
        try:
            root.destroy()
        finally:
            os._exit(0)

    root.protocol("WM_DELETE_WINDOW", on_close)
    drain()
    root.mainloop()


def main() -> None:
    ap = argparse.ArgumentParser(description="Crawler sieuthicongcu.com")
    ap.add_argument("--mode", choices=["code", "keyword", "link", "excel", "all"])
    ap.add_argument("--code")
    ap.add_argument("--keyword")
    ap.add_argument("--link")
    ap.add_argument("--excel")
    ap.add_argument("--sheet", default="")
    ap.add_argument("--col", default="")
    ap.add_argument("--out", default="sieuthicongcu.xlsx")
    ap.add_argument("--conc", type=int, default=8)
    ap.add_argument("--delay", type=float, default=0.05)
    ap.add_argument("--limit", type=int, default=0)
    args = ap.parse_args()

    if not args.mode:
        run_gui()
        return

    target = args.code or args.keyword or args.link or ""
    engine = Engine(args.conc, args.delay)
    t0 = time.time()
    try:
        rows = asyncio.run(engine.run(args.mode, target, args.excel or "", args.sheet, args.col, args.limit))
    finally:
        if engine.results:
            export_excel(engine.results, pathlib.Path(args.out))
        engine.cleanup()
    print(f"[SYS] {len(rows)} dong -> {args.out} ({time.time() - t0:.1f}s)")


if __name__ == "__main__":
    main()
