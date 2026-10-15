"""Hoffmann Group UK crawler.

Recon facts:
- SAP Commerce storefront with public exact resolver:
  /GB/en/houk/fastproductsearch/v1/check/product?productCode=<code>
- PDP URL pattern: /GB/en/houk/p/<article>-<variant>
- Technical data: #technicalData > table:nth-child(2).
"""

from __future__ import annotations

import html as html_lib
import json
import sqlite3
import asyncio
import time
import argparse
import threading
from pathlib import Path
from collections.abc import Callable
from urllib.parse import quote, urljoin, urlparse

import httpx
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from selectolax.lexbor import LexborHTMLParser


BASE_URL = "https://www.hoffmann-group.com/GB/en/houk"
USER_AGENT = "Mozilla/5.0 (compatible; HoffmannCrawler/1.0)"
MAX_REDIRECTS = 5
CIRCUIT_BREAK = 25
MAX_ITEMS = 5000
MAX_BFS_PAGES = 1500
MAX_PAGES = 600
HERE = Path(__file__).resolve().parent
PDP_PREFIX = urlparse(BASE_URL).path + "/p/"


def norm_input(value: str) -> str:
    """Normalize an article code without corrupting a case-sensitive URL."""
    normalized = str(value).strip()
    if normalized.lower().startswith(("http://", "https://")):
        return normalized
    return normalized.upper()


def parse_resolver(payload: dict) -> str | None:
    """Return the PDP path from Hoffmann's public exact-code response."""
    path = payload.get("pdpLink")
    return path if isinstance(path, str) and path else None


def is_block_page(page_html: str) -> bool:
    """Detect Hoffmann's HTTP-200 WAF interstitial instead of parsing it as empty search."""
    lowered = page_html.lower()
    return "blockingpage" in lowered and "blocking-page" in lowered


def _text(node: object) -> str:
    value = node.text(strip=True)  # type: ignore[attr-defined]
    return " ".join(value.replace("\xa0", " ").split())


def parse_category_urls(page_html: str) -> list[str]:
    """Return ordered, unique PDP paths from a rendered category/search page."""
    tree = LexborHTMLParser(page_html)
    seen: set[str] = set()
    paths: list[str] = []
    for anchor in tree.css("a[href]"):
        href = anchor.attributes.get("href", "")
        path = urlparse(href).path
        if not path.startswith(PDP_PREFIX) or path in seen:
            continue
        seen.add(path)
        paths.append(path)
    return paths


def _product_payload(tree: LexborHTMLParser) -> dict:
    for node in tree.css("[data-product]"):
        raw = node.attributes.get("data-product", "")
        if not raw:
            continue
        try:
            value = json.loads(html_lib.unescape(raw))
        except json.JSONDecodeError:
            continue
        if isinstance(value, dict):
            return value
    return {}


def _table_values(tree: LexborHTMLParser) -> tuple[dict[str, str], list[str]]:
    identifiers: dict[str, str] = {}
    specs: list[str] = []
    id_keys = {
        "article number": "Item No",
        "ean / gtin": "EAN / GTIN",
        "manufacturer number": "Manufacturer No",
    }
    for row in tree.css("#technicalData > table:nth-child(2) tr"):
        cells = row.css("td")
        if len(cells) < 2:
            continue
        name = _text(cells[0])
        value = _text(cells[-1])
        if not name or not value:
            continue
        normalized = name.casefold()
        matched = next((field for key, field in id_keys.items() if key in normalized), None)
        if matched:
            identifiers[matched] = value
        elif "manufacturer/responsible" not in normalized:
            specs.append(f"{name}: {value}")
    return identifiers, specs


def _description(tree: LexborHTMLParser) -> str:
    section = tree.css_first("#description")
    if section is None:
        return ""
    parts: list[str] = []
    for heading in section.css("h3"):
        label = _text(heading).rstrip(":")
        sibling = heading.next
        while sibling is not None and sibling.tag != "h3":
            value = _text(sibling)
            if value:
                parts.append(f"{label}: {value}" if label else value)
                break
            sibling = sibling.next
    return "\n".join(parts)


def parse_pdp(page_html: str, url: str, input_value: str = "") -> list[dict[str, object]]:
    """Flatten the public PDP into one exact-variant export row."""
    tree = LexborHTMLParser(page_html)
    product = _product_payload(tree)
    identifiers, specs = _table_values(tree)
    title = _text(tree.css_first("h1")) if tree.css_first("h1") else ""
    images = product.get("images") if isinstance(product.get("images"), list) else []
    image = next(
        (item.get("url", "") for item in images if isinstance(item, dict) and item.get("url")),
        "",
    )
    datasheet = next(
        (anchor.attributes.get("href", "") for anchor in tree.css("a[href]")
         if anchor.attributes.get("href", "").lower().split("?")[0].endswith(".pdf")),
        "",
    )
    price = product.get("price") if isinstance(product.get("price"), dict) else {}
    stock = product.get("stock") if isinstance(product.get("stock"), dict) else {}
    item_no = identifiers.get("Item No") or str(product.get("code") or "")
    row: dict[str, object] = {
        "Input": input_value,
        "Page/Name": title or str(product.get("name") or ""),
        "Item No": item_no,
        "Status": "OK" if item_no else "Partial",
        "Brand": str(product.get("manufacturer") or ""),
        "EAN / GTIN": identifiers.get("EAN / GTIN", ""),
        "Manufacturer No": identifiers.get("Manufacturer No", ""),
        "Thong so": "\n".join(specs),
        "Description": _description(tree),
        "Price": price.get("value", ""),
        "Currency": price.get("currencyIso", ""),
        "Stock": stock.get("stockLevel", ""),
        "Data Sheet": datasheet,
        "Image": image,
        "URL": url,
    }
    return [row]


def parse_variant_urls(page_html: str) -> list[str]:
    """Extract ordered variant PDP paths from the structured PDP payload."""
    product = _product_payload(LexborHTMLParser(page_html))
    options = product.get("variantOptions")
    if not isinstance(options, list):
        return []
    paths: list[str] = []
    seen: set[str] = set()
    for option in options:
        path = option.get("url") if isinstance(option, dict) else None
        if isinstance(path, str) and "/p/" in path and path not in seen:
            seen.add(path)
            paths.append(path)
    return paths


class CheckpointStore:
    """Small SQLite work queue that survives a stop or process crash."""

    def __init__(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.conn = sqlite3.connect(path)
        self.conn.execute("PRAGMA journal_mode=WAL")
        self.conn.execute("PRAGMA synchronous=NORMAL")
        self.conn.execute(
            "CREATE TABLE IF NOT EXISTS queue ("
            "url TEXT PRIMARY KEY, source TEXT NOT NULL, done INTEGER NOT NULL DEFAULT 0)"
        )
        self.conn.execute(
            "CREATE TABLE IF NOT EXISTS products (url TEXT PRIMARY KEY, payload TEXT NOT NULL)"
        )
        self.conn.execute(
            "CREATE TABLE IF NOT EXISTS code_resolution (code TEXT PRIMARY KEY, url TEXT NOT NULL)"
        )
        self.conn.commit()

    def enqueue(self, url: str, source: str) -> bool:
        cursor = self.conn.execute(
            "INSERT OR IGNORE INTO queue(url, source) VALUES (?, ?)", (url, source)
        )
        self.conn.commit()
        return cursor.rowcount == 1

    def next_pending(self, limit: int) -> list[str]:
        rows = self.conn.execute(
            "SELECT url FROM queue WHERE done = 0 ORDER BY rowid LIMIT ?", (limit,)
        ).fetchall()
        return [row[0] for row in rows]

    def mark_done(self, url: str) -> None:
        self.conn.execute("UPDATE queue SET done = 1 WHERE url = ?", (url,))
        self.conn.commit()

    def resolved_url(self, code: str) -> str | None:
        row = self.conn.execute("SELECT url FROM code_resolution WHERE code = ?", (code,)).fetchone()
        return str(row[0]) if row else None

    def save_resolution(self, code: str, url: str) -> None:
        self.conn.execute(
            "INSERT INTO code_resolution(code, url) VALUES (?, ?) "
            "ON CONFLICT(code) DO UPDATE SET url = excluded.url",
            (code, url),
        )
        self.conn.commit()

    def upsert_product(self, row: dict[str, object]) -> None:
        url = str(row["URL"])
        payload = json.dumps(row, ensure_ascii=False, default=str)
        self.conn.execute(
            "INSERT INTO products(url, payload) VALUES (?, ?) "
            "ON CONFLICT(url) DO UPDATE SET payload = excluded.payload",
            (url, payload),
        )
        self.conn.commit()

    def products(self) -> list[dict[str, object]]:
        rows = self.conn.execute("SELECT payload FROM products ORDER BY rowid").fetchall()
        return [json.loads(row[0]) for row in rows]


class RateLimiter:
    """Token bucket shared by all concurrent Hoffmann requests."""

    def __init__(self, per_second: float) -> None:
        self.per_second = per_second
        self.tokens = per_second
        self.updated_at = time.monotonic()
        self.lock = asyncio.Lock()

    async def acquire(self) -> None:
        async with self.lock:
            now = time.monotonic()
            self.tokens = min(self.per_second, self.tokens + (now - self.updated_at) * self.per_second)
            self.updated_at = now
            if self.tokens < 1:
                await asyncio.sleep((1 - self.tokens) / self.per_second)
                self.tokens = 0
            else:
                self.tokens -= 1


class HoffmannEngine:
    """Async API-first crawler engine, independent of GUI and CLI."""

    def __init__(
        self,
        store: CheckpointStore,
        *,
        base_url: str = BASE_URL,
        transport: httpx.AsyncBaseTransport | None = None,
        concurrency: int = 1,
        rate_per_second: float = 1.0,
        on_event: Callable[[str], None] | None = None,
    ) -> None:
        self.store = store
        self.base_url = base_url.rstrip("/")
        self.transport = transport
        self.semaphore = asyncio.Semaphore(concurrency)
        self.limiter = RateLimiter(rate_per_second)
        self.pause_event = asyncio.Event()
        self.pause_event.set()
        self.stop_event = asyncio.Event()
        self.consecutive_errors = 0
        self.last_failure = ""
        self.client: httpx.AsyncClient | None = None
        self.background: set[asyncio.Task[None]] = set()
        self.on_event = on_event or (lambda _: None)

    def emit(self, message: str) -> None:
        self.on_event(message)

    async def __aenter__(self) -> "HoffmannEngine":
        self.client = httpx.AsyncClient(
            http2=True,
            follow_redirects=True,
            max_redirects=MAX_REDIRECTS,
            timeout=httpx.Timeout(20.0, connect=10.0),
            limits=httpx.Limits(max_connections=8, max_keepalive_connections=4),
            headers={"User-Agent": USER_AGENT, "Accept-Language": "en-GB,en;q=0.9"},
            transport=self.transport,
        )
        return self

    async def __aexit__(self, *_: object) -> None:
        if self.client is not None:
            await self.client.aclose()
            self.client = None

    async def _resume_after_waf_pause(self) -> None:
        await asyncio.sleep(60)
        self.pause_event.set()

    def _start_waf_pause(self, status: int | str) -> None:
        self.last_failure = "Rate Limited"
        self.emit(f"[WARN] {status} rate-limited; pausing 60s")
        if self.pause_event.is_set():
            self.pause_event.clear()
            task = asyncio.create_task(self._resume_after_waf_pause())
            self.background.add(task)
            task.add_done_callback(self.background.discard)

    async def fetch(self, url: str) -> httpx.Response | None:
        if self.client is None:
            raise RuntimeError("HoffmannEngine must be used inside 'async with'")
        self.last_failure = ""
        for attempt in range(5):
            if self.stop_event.is_set():
                return None
            await self.pause_event.wait()
            await self.limiter.acquire()
            try:
                response = await self.client.get(url)
                if response.status_code == 404:
                    self.emit(f"[MISS] 404 {url}")
                    return None
                if response.status_code in {403, 429, 503, 520, 521, 522}:
                    self._start_waf_pause(response.status_code)
                    return None
                if is_block_page(response.text):
                    self._start_waf_pause("HTTP 200 blocking page")
                    return None
                response.raise_for_status()
                self.consecutive_errors = 0
                self.last_failure = ""
                return response
            except (httpx.HTTPError, asyncio.TimeoutError) as exc:
                if not self.last_failure:
                    self.last_failure = type(exc).__name__
                self.consecutive_errors += 1
                self.emit(f"[WARN] {type(exc).__name__} on attempt {attempt + 1}/5")
                if self.consecutive_errors >= CIRCUIT_BREAK:
                    self.stop_event.set()
                    return None
                if attempt == 4:
                    return None
                await asyncio.sleep(min(2**attempt, 20))
        return None

    def failure_row(self, input_value: str, url: str = "") -> dict[str, object]:
        if not self.last_failure:
            return {"Input": input_value, "Item No": "", "Status": "Not Found", "URL": url}
        status = "Rate Limited" if self.last_failure == "Rate Limited" else "Network Error"
        return {
            "Input": input_value,
            "Item No": "",
            "Status": status,
            "Error": self.last_failure,
            "URL": url,
        }

    def _absolute(self, path_or_url: str) -> str:
        return urljoin(self.base_url + "/", path_or_url)

    async def resolve_code(self, code: str) -> str | None:
        normalized = norm_input(code)
        cached_url = self.store.resolved_url(normalized)
        if cached_url:
            self.emit(f"[CACHE] Resolved {normalized}")
            return cached_url
        self.emit(f"[INFO] Resolving {normalized}")
        response = await self.fetch(
            f"{self.base_url}/fastproductsearch/v1/check/product?productCode={quote(normalized, safe='')}"
        )
        if response is None:
            return None
        try:
            path = parse_resolver(response.json())
        except json.JSONDecodeError:
            return None
        url = self._absolute(path) if path else None
        if url:
            self.store.save_resolution(normalized, url)
        return url

    async def crawl_url(self, url: str, input_value: str = "") -> list[dict[str, object]]:
        self.emit(f"[INFO] Loading PDP {url}")
        response = await self.fetch(url)
        if response is None:
            return [self.failure_row(input_value, url)]
        rows = parse_pdp(response.text, str(response.url), input_value)
        for path in parse_variant_urls(response.text):
            variant_url = self._absolute(path)
            if variant_url != str(response.url):
                self.store.enqueue(variant_url, "variant")
        for row in rows:
            self.store.upsert_product(row)
            self.emit(f"[OK] Saved {row.get('Item No') or input_value}")
        return rows

    async def crawl_codes(self, codes: list[str]) -> list[dict[str, object]]:
        results: list[dict[str, object]] = []
        async with self:
            for raw in codes[:MAX_ITEMS]:
                if self.stop_event.is_set():
                    break
                value = norm_input(raw)
                url = value if value.lower().startswith("http") else await self.resolve_code(value)
                if not url:
                    results.append(self.failure_row(raw))
                    continue
                results.extend(await self.crawl_url(url, raw))
        return results

    async def discover_category(self, category_url: str, max_pages: int = MAX_PAGES) -> int:
        """Persist PDP URLs from every non-stuck server-side category page."""
        if self.client is None:
            raise RuntimeError("HoffmannEngine must be used inside 'async with'")
        added = 0
        empty_pages = 0
        for page in range(1, min(max_pages, MAX_PAGES) + 1):
            separator = "&" if "?" in category_url else "?"
            response = await self.fetch(f"{category_url}{separator}page={page}")
            if response is None:
                break
            urls = parse_category_urls(response.text)
            new_on_page = 0
            for path in urls:
                if self.store.enqueue(self._absolute(path), f"category:{category_url}"):
                    new_on_page += 1
            added += new_on_page
            empty_pages = empty_pages + 1 if new_on_page == 0 else 0
            if not urls or empty_pages >= 2 or self.stop_event.is_set():
                break
        return added

    async def _crawl_pending_open(self, limit: int = 0) -> list[dict[str, object]]:
        results: list[dict[str, object]] = []
        remaining = limit or MAX_ITEMS
        while remaining > 0 and not self.stop_event.is_set():
            urls = self.store.next_pending(min(remaining, max(1, self.semaphore._value * 4)))
            if not urls:
                break

            async def crawl_one(item_url: str) -> list[dict[str, object]]:
                async with self.semaphore:
                    try:
                        return await self.crawl_url(item_url, item_url)
                    finally:
                        self.store.mark_done(item_url)

            batches = await asyncio.gather(*(crawl_one(item_url) for item_url in urls))
            for batch in batches:
                results.extend(batch)
            remaining -= len(urls)
        return results

    async def crawl_pending(self, limit: int = 0) -> list[dict[str, object]]:
        async with self:
            return await self._crawl_pending_open(limit)

    async def crawl_keyword(self, keyword: str, limit: int = MAX_ITEMS) -> list[dict[str, object]]:
        async with self:
            response = await self.fetch(f"{self.base_url}/search?text={quote(keyword)}")
            if response is None:
                return [self.failure_row(keyword)]
            urls = parse_category_urls(response.text)[:limit]
            if not urls:
                self.emit(f"[MISS] No product matches keyword: {keyword}")
                return [self.failure_row(keyword)]
            return [row for url in urls for row in await self.crawl_url(self._absolute(url), keyword)]

    async def crawl_all(self, limit: int = 0) -> list[dict[str, object]]:
        roots = [
            "/Machining/c/10-00-00-00-00", "/Clamping-technology/c/30-00-00-00-00",
            "/Measurement-technology/c/994", "/Grinding-cutting-technology/c/50-00-00-00-00",
            "/Hand-tools/c/60-00-00-00-00", "/Workstations-and-storage-and-factory-equipment/c/90-00-00-00-00",
            "/Power-tools-and-pneumatic-tools/c/07-00-00-00-00", "/Chemical-technical-products/c/06-00-00-00-00",
            "/Personal-Protective-Equipment/c/PSA",
        ]
        async with self:
            for root in roots:
                if self.stop_event.is_set():
                    break
                await self.discover_category(self._absolute(root))
                if limit and len(self.store.next_pending(limit)) >= limit:
                    break
            return await self._crawl_pending_open(limit)


def _scalar(value: object) -> str | int | float | bool | None:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, dict):
        return "; ".join(f"{key}: {_scalar(item) or ''}" for key, item in value.items())
    if isinstance(value, (list, tuple, set)):
        return ", ".join(str(_scalar(item) or "") for item in value)
    return str(value)


def export_excel(rows: list[dict[str, object]], path: Path) -> None:
    """Export flattened rows; technical specifications always occupy one column."""
    path.parent.mkdir(parents=True, exist_ok=True)
    base_headers = [
        "Input", "Page/Name", "Item No", "Status", "Brand", "EAN / GTIN",
        "Manufacturer No", "Thong so", "Description", "Price", "Currency", "Stock",
        "Data Sheet", "Image", "URL",
    ]
    categories = sorted({key for row in rows for key in row if key.startswith("Category ")})
    headers = base_headers[:7] + categories + base_headers[7:]
    book = Workbook()
    sheet = book.active
    sheet.title = "Hoffmann"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top")
    for row in rows:
        sheet.append([_scalar(row.get(header, "")) for header in headers])
    spec_col = headers.index("Thong so") + 1
    sheet.column_dimensions["A"].width = 20
    for column in range(2, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 22
    sheet.column_dimensions[get_column_letter(spec_col)].width = 60
    for cells in sheet.iter_rows(min_row=2):
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column == spec_col)
    sheet.freeze_panes = "A2"
    book.save(path)


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Hoffmann Group UK crawler")
    parser.add_argument("--mode", choices=["code", "keyword", "link", "excel", "all"], default="code")
    parser.add_argument("--code", action="append", default=[], help="Article code; may be passed repeatedly")
    parser.add_argument("--keyword", default="")
    parser.add_argument("--links", default="", help="One or more PDP URLs separated by newlines")
    parser.add_argument("--excel", default="")
    parser.add_argument("--sheet", default="")
    parser.add_argument("--col", default="")
    parser.add_argument("--output", default=str(HERE / "hoffmann_ketqua.xlsx"))
    parser.add_argument("--db", default=str(HERE / ".hoffmann_state.sqlite3"))
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--concurrency", type=int, default=1)
    parser.add_argument("--rate", type=float, default=1.0)
    return parser


def _excel_values(path: str, sheet: str, column: str) -> list[str]:
    frame = pd.read_excel(path, sheet_name=sheet or 0)
    selected = column if column in frame.columns else str(frame.columns[0])
    return [str(value).strip() for value in frame[selected].dropna() if str(value).strip()]


def result_summary(rows: list[dict[str, object]]) -> str:
    counts: dict[str, int] = {}
    for row in rows:
        status = str(row.get("Status") or "Unknown")
        counts[status] = counts.get(status, 0) + 1
    return " | ".join(f"{count} {status}" for status, count in counts.items()) or "0 rows"


async def run_args(
    args: argparse.Namespace,
    on_event: Callable[[str], None] | None = None,
) -> list[dict[str, object]]:
    store = CheckpointStore(Path(args.db))
    engine = HoffmannEngine(
        store,
        concurrency=max(1, args.concurrency),
        rate_per_second=max(0.5, args.rate),
        on_event=on_event,
    )
    if args.mode == "code":
        rows = await engine.crawl_codes(args.code)
    elif args.mode == "link":
        rows = await engine.crawl_codes([line.strip() for line in args.links.splitlines() if line.strip()])
    elif args.mode == "keyword":
        rows = await engine.crawl_keyword(args.keyword, args.limit or MAX_ITEMS)
    elif args.mode == "excel":
        rows = await engine.crawl_codes(_excel_values(args.excel, args.sheet, args.col))
    else:
        rows = await engine.crawl_all(args.limit)
    export_excel(rows or store.products(), Path(args.output))
    return rows


class HoffmannApp:
    """Small Tk surface; the engine remains testable and headless."""

    def __init__(self) -> None:
        import tkinter as tk
        from tkinter import ttk

        root_dir = next(path for path in (HERE, *HERE.parents) if (path / "gui_common.py").exists())
        import sys
        sys.path.insert(0, str(root_dir))
        from gui_common import attach_excel_picker

        self.tk, self.ttk = tk, ttk
        self.root = tk.Tk()
        self.root.title("Hoffmann Group Crawler")
        self.root.geometry("900x620")
        self.mode = tk.StringVar(value="code")
        self.value = tk.StringVar()
        self.xlsx, self.sheet, self.col, self.outdir = tk.StringVar(), tk.StringVar(), tk.StringVar(), tk.StringVar()
        self.output = tk.StringVar(value=str(HERE / "hoffmann_ketqua.xlsx"))
        self.status = tk.StringVar(value="San sang")
        left, right = ttk.Frame(self.root, padding=12), ttk.Frame(self.root, padding=12)
        left.pack(side="left", fill="y"); right.pack(side="right", fill="both", expand=True)
        for mode, label in [("code", "Theo ma"), ("keyword", "Theo keyword"), ("link", "Theo link"), ("excel", "Theo Excel"), ("all", "Crawl ALL")]:
            ttk.Radiobutton(left, text=label, variable=self.mode, value=mode).pack(anchor="w")
        ttk.Label(left, text="Ma / keyword / link (moi dong mot gia tri):").pack(anchor="w", pady=(12, 0))
        self.text = tk.Text(left, width=42, height=9); self.text.pack()
        attach_excel_picker(left, self.xlsx, self.sheet, self.col, self.outdir)
        ttk.Label(left, text="File ket qua:").pack(anchor="w", pady=(10, 0))
        ttk.Entry(left, textvariable=self.output, width=42).pack()
        ttk.Button(left, text="START", command=self.start).pack(fill="x", pady=(10, 2))
        ttk.Label(left, textvariable=self.status).pack(anchor="w")
        self.console = tk.Text(right, bg="#1e1e1e", fg="#dddddd", insertbackground="white")
        self.console.pack(fill="both", expand=True)
        self.root.protocol("WM_DELETE_WINDOW", self.close)

    def log(self, message: str) -> None:
        self.root.after(0, lambda: (self.console.insert("end", message + "\n"), self.console.see("end")))

    def start(self) -> None:
        mode, raw = self.mode.get(), self.text.get("1.0", "end").strip()
        self.status.set("Dang crawl...")

        def worker() -> None:
            args = build_arg_parser().parse_args(["--mode", mode, "--output", self.output.get()])
            if mode == "code": args.code = [line for line in raw.splitlines() if line.strip()]
            elif mode == "keyword": args.keyword = raw
            elif mode == "link": args.links = raw
            elif mode == "excel": args.excel, args.sheet, args.col = self.xlsx.get(), self.sheet.get(), self.col.get()
            try:
                rows = asyncio.run(run_args(args, self.log))
                self.log(f"[DONE] {result_summary(rows)} -> {args.output}")
                self.root.after(0, lambda: self.status.set("Xong"))
            except BaseException as exc:
                self.log(f"[ERR] {exc}"); self.root.after(0, lambda: self.status.set("Loi"))
        threading.Thread(target=worker, daemon=True).start()

    def close(self) -> None:
        self.root.destroy()

    def run(self) -> None:
        self.root.mainloop()


def main() -> int:
    import sys
    if len(sys.argv) == 1:
        HoffmannApp().run()
        return 0
    args = build_arg_parser().parse_args()
    rows = asyncio.run(run_args(args, print))
    print(f"[DONE] {result_summary(rows)} -> {args.output}")
    return 0 if rows and all(row.get("Status") == "OK" for row in rows) else 1


if __name__ == "__main__":
    raise SystemExit(main())
