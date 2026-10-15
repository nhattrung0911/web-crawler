"""
totalbusiness.com (TOTAL Tools) — crawler catalog (httpx thuan, KHONG browser, KHONG Selenium).

RECON (2026-07-13)
------------------
CMS      : SPA React/umi 3.4 NHUNG co **SSR** -> moi trang tra san JSON trong HTML:
             <script>window.g_initialProps = { seo, langData, layoutData, config, data }</script>
           => chi can httpx + json, KHONG can Playwright (dung browser o day la phi).
API      : XHR rieng chi la lat vat (webcenterapi.totalbusiness.com: /home/ip, /website-contact/want-to).
           Du lieu san pham nam het trong g_initialProps cua trang -> dung thang.
list     : /products?pageNum=N        -> data.total=919, data.list[] (40 SP/trang) => 23 trang
           /products/<category-slug>?pageNum=N   (11 category, vd power-tools, hand-tools)
search   : /products?keyWord=<ma>     -> khop chinh xac ('TKLI1201'->1 KQ, 'drill'->64, 'ZZZZZ'->0)
           (LUU Y: ?keyword= / ?searchKey= bi BO QUA, tra ve ca 919 SP -> phai dung dung 'keyWord')
detail   : /product/<pathReplacement>/<productNo>  -> data:
             productNo, productName, categoryName, keyData1..3,
             parameter (THONG SO, cac dong ngan cach bang <br>, dang 'Ten:Gia tri'),
             picture[] (list URL tuyet doi), video, relevant[], annex[], seo
gia      : KHONG co gia (catalog B2B, chi hoi hang) -> khong crawl gia.
robots   : /robots.txt tra ve chinh HTML SPA (khong co robots that) -> khong co luat cam.

Run
---
  totalbusiness.bat                                  -> GUI
  python totalbusiness.py --mode all [--limit 100]   -> quet sach 919 SP
  python totalbusiness.py --mode all --category hand-tools
  python totalbusiness.py --mode code --code TKLI1201
  python totalbusiness.py --mode keyword --keyword drill
  python totalbusiness.py --mode link --link https://www.totalbusiness.com/product/cordless-drill/TDLI2002
  python totalbusiness.py --mode excel --excel parts.xlsx --col Part
  python totalbusiness.py --mode images --excel parts.xlsx --col Part --outdir out
"""
from __future__ import annotations

import argparse
import asyncio
import gc
import json
import os
import pathlib
import queue
import re
import threading
import time
from typing import Any
from urllib.parse import quote, urlparse

import httpx
import pandas as pd

HERE = pathlib.Path(__file__).parent
BASE = "https://www.totalbusiness.com"
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/122.0 Safari/537.36")

PER_PAGE = 40             # SSR tra 40 SP/trang, khong doi duoc
CIRCUIT_BREAK = 25        # N loi LIEN TIEP -> nghi bi chan IP -> dung han
MAX_ITEMS = 5000
MAX_PAGES = 200
MAX_REDIRECTS = 5
TIMEOUT = 30.0
GC_EVERY = 300

PROPS_RE = re.compile(r"window\.g_initialProps\s*=\s*(\{.*?\});?\s*</script>", re.S)


# ================================================================ CORE (khong dung GUI)
def norm_code(s: Any) -> str:
    v = re.sub(r"\s+", "", str(s or ""))
    return v if v.lower().startswith("http") else v.upper()


def parse_props(html: str) -> dict | None:
    """Lay JSON SSR window.g_initialProps ra khoi HTML."""
    m = PROPS_RE.search(html)
    if not m:
        return None
    try:
        return json.loads(m.group(1))
    except Exception:
        return None


def list_url(page: int = 1, category: str = "", keyword: str = "") -> str:
    path = f"/products/{category}" if category else "/products"
    q = f"?pageNum={page}"
    if keyword:
        q += f"&keyWord={quote(keyword)}"
    return f"{BASE}{path}{q}"


def detail_url(item: dict) -> str:
    return f"{BASE}/product/{item.get('pathReplacement', '')}/{item.get('productNo', '')}"


def parse_spec(parameter: str) -> tuple[dict[str, str], str]:
    """`parameter` = cac dong ngan cach <br>. 'Ten:Gia tri' -> spec; dong khong co ':' -> ghi chu."""
    spec: dict[str, str] = {}
    notes: list[str] = []
    for line in re.split(r"<br\s*/?>", parameter or ""):
        line = re.sub(r"<[^>]+>", "", line).replace("\xa0", " ").strip()
        if not line:
            continue
        k, sep, v = line.partition(":")
        k, v = k.strip(), v.strip()
        if sep and k and v and len(k) < 60:
            spec[k] = v if k not in spec else f"{spec[k]}; {v}"
        else:
            notes.append(line)
    return spec, " | ".join(notes)


def rec_from_list(it: dict, url: str) -> dict:
    """Fallback khi trang detail bi 403 (site chan san 1 so SP): dung du lieu tu trang list."""
    spec = {}
    for i in (1, 2, 3):
        line = (it.get(f"keyData{i}") or "").strip()
        k, sep, v = line.partition(":")
        if sep and k.strip() and v.strip():
            spec[k.strip()] = v.strip()
    return {"Input": "", "Item No": it.get("productNo") or "", "Ten": it.get("productName") or "",
            "Status": "OK (list only - detail 403)", "Ghi chu": "", "Video": "",
            "_cats": [], "_spec": spec, "Image": it.get("picture") or "", "URL": url}


def rec_from_detail(d: dict, url: str) -> dict:
    spec, notes = parse_spec(d.get("parameter") or "")
    pics = d.get("picture") or []
    if isinstance(pics, str):
        pics = [pics]
    return {"Input": "", "Item No": d.get("productNo") or "", "Ten": d.get("productName") or "",
            "Status": "OK", "Ghi chu": notes, "Video": d.get("video") or "",
            "_cats": [c for c in [d.get("categoryName")] if c], "_spec": spec,
            "Image": " | ".join(pics[:5]), "URL": url}


def export_excel(results: list[dict], path: pathlib.Path) -> pathlib.Path:
    """1 cot 'Thong so' duy nhat, moi dong 'Ten: Gia tri' (TOTAL dat ten spec rat tuy hung ->
    trai ra cot dong se thanh 500+ cot rong)."""
    max_cat = max((len(r.get("_cats") or []) for r in results), default=0)
    rows = []
    for r in results:
        row = {k: r.get(k, "") for k in ("Input", "Item No", "Ten", "Status")}
        cats = r.get("_cats") or []
        for i in range(max_cat):
            row[f"Category {i + 1}"] = cats[i] if i < len(cats) else ""
        lines = [f"{k}: {v}" for k, v in (r.get("_spec") or {}).items()]
        lines += [n for n in (r.get("Ghi chu") or "").split(" | ") if n]
        row["Thong so"] = "\n".join(lines)
        row["Image"] = r.get("Image", "")
        row["Video"] = r.get("Video", "")
        row["URL"] = r.get("URL", "")
        rows.append(row)
    from openpyxl.utils import get_column_letter

    df = pd.DataFrame(rows)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, index=False)
        ws = xw.sheets["Sheet1"]
        col = get_column_letter(df.columns.get_loc("Thong so") + 1)
        ws.column_dimensions[col].width = 60
        for c in ws[col][1:]:                       # xuong dong trong o -> hien thi duoc
            c.alignment = c.alignment.copy(wrap_text=True, vertical="top")
    return path


class Engine:
    def __init__(self, conc: int = 6, delay: float = 0.1, log=print,
                 stop_evt: threading.Event | None = None,
                 pause_evt: threading.Event | None = None) -> None:
        self.conc, self.delay, self.log = max(1, conc), max(0.0, delay), log
        self.stop_evt = stop_evt or threading.Event()
        self.pause_evt = pause_evt or threading.Event()
        self.fails = 0
        self.results: list[dict] = []
        self.seen_urls: set[str] = set()
        self.done = 0
        self.total = 0
        self.key = ""

    # ---- resume: JSONL sidecar
    def ckpt(self, key: str) -> pathlib.Path:
        return HERE / f"_raw_{re.sub(r'[^A-Za-z0-9]+', '_', key)[:40]}.jsonl"

    def _load_ckpt(self, key: str) -> set[str]:
        p, seen = self.ckpt(key), set()
        if not p.exists():
            return seen
        for line in p.read_text("utf-8", errors="ignore").splitlines():
            try:
                r = json.loads(line)
            except Exception:
                continue
            self.results.append(r)
            seen.add(r.get("URL", ""))
            self.seen_urls.add(r.get("URL", ""))
        if seen:
            self.log(f"[SYS] Resume: {len(self.results)} record da co -> bo qua")
        return seen

    def add(self, recs: list[dict], key: str) -> None:
        fresh = [r for r in recs if r.get("URL") and r["URL"] not in self.seen_urls]
        for r in fresh:
            self.seen_urls.add(r["URL"])
        if fresh:
            self.results += fresh
            with self.ckpt(key).open("a", encoding="utf-8") as f:
                for r in fresh:
                    f.write(json.dumps(r, ensure_ascii=False) + "\n")
                f.flush()

    def cleanup(self) -> None:
        p = self.ckpt(self.key) if self.key else None
        if p and p.exists() and not self.stop_evt.is_set():
            p.unlink()

    async def _gate(self) -> bool:
        while self.pause_evt.is_set() and not self.stop_evt.is_set():
            await asyncio.sleep(0.2)
        return not self.stop_evt.is_set()

    async def fetch(self, cli: httpx.AsyncClient, url: str, tries: int = 4) -> httpx.Response | None:
        for i in range(tries):
            if not await self._gate():
                return None
            try:
                r = await cli.get(url)
                if r.status_code in (403, 429, 503):
                    wait = min(2 ** i * 5, 60)
                    self.log(f"[WARN] {r.status_code} -> ngu {wait}s | {url[-55:]}")
                    await asyncio.sleep(wait)
                    continue
                if r.status_code == 404:
                    return None
                r.raise_for_status()
                self.fails = 0
                return r
            except Exception as e:
                if i < tries - 1:
                    await asyncio.sleep(min(2 ** i, 20))
                    continue
                self.fails += 1
                self.log(f"[ERR] {type(e).__name__} | {url[-55:]}")
                if self.fails >= CIRCUIT_BREAK:
                    self.log(f"[STOP] {CIRCUIT_BREAK} loi LIEN TIEP - nghi bi chan IP. DUNG.")
                    self.stop_evt.set()
                return None
        return None

    async def props(self, cli: httpx.AsyncClient, url: str) -> dict | None:
        r = await self.fetch(cli, url)
        return parse_props(r.text) if r else None

    async def _collect_list(self, cli: httpx.AsyncClient, category: str = "",
                            keyword: str = "", limit: int = 0) -> list[dict]:
        """Duyet trang list SSR -> list item (co productNo + pathReplacement)."""
        first = await self.props(cli, list_url(1, category, keyword))
        data = (first or {}).get("data") or {}
        total = int(data.get("total") or 0)
        items = list(data.get("list") or [])
        if not total:
            self.log("[MISS] 0 ket qua")
            return []
        pages = min((total + PER_PAGE - 1) // PER_PAGE, MAX_PAGES)
        self.log(f"[SYS] {total} SP / {pages} trang list")
        sem = asyncio.Semaphore(self.conc)

        async def one(pg: int) -> list[dict]:
            async with sem:
                if not await self._gate():
                    return []
                d = await self.props(cli, list_url(pg, category, keyword))
                await asyncio.sleep(self.delay)
                return list(((d or {}).get("data") or {}).get("list") or [])

        for part in await asyncio.gather(*(one(p) for p in range(2, pages + 1))):
            items += part
        uniq = list({i.get("productNo"): i for i in items if i.get("productNo")}.values())
        return uniq[:limit or MAX_ITEMS]

    async def _crawl_details(self, cli: httpx.AsyncClient, urls: list[str], key: str,
                             inputs: dict[str, str] | None = None,
                             items: dict[str, dict] | None = None) -> None:
        """items[url] = ban ghi tu trang list -> dung lam fallback khi detail 403/hong."""
        seen = self._load_ckpt(key)
        urls = [u for u in dict.fromkeys(urls) if u not in seen][:MAX_ITEMS]
        self.total, self.done = len(urls) + len(seen), len(seen)
        sem = asyncio.Semaphore(self.conc)

        async def one(u: str) -> None:
            async with sem:
                if not await self._gate():
                    return
                d = await self.props(cli, u)
                await asyncio.sleep(self.delay)
                data = (d or {}).get("data") or {}
                rec = rec_from_detail(data, u) if data.get("productNo") else None
                if rec is None and items and u in items:
                    rec = rec_from_list(items[u], u)      # detail 403 -> cuu bang du lieu list
                if rec:
                    # Input chi co nghia khi crawl theo ma/Excel; mode all -> de trong
                    rec["Input"] = (inputs or {}).get(norm_code(rec["Item No"]), "")
                    self.add([rec], key)
                self.done += 1
                self.log(f"[{'OK ' if rec else 'MISS'}] {self.done}/{self.total} | {u[-50:]}")
                if self.done % GC_EVERY == 0:
                    gc.collect()

        try:
            await asyncio.gather(*(one(u) for u in urls))
        except (KeyboardInterrupt, Exception) as e:
            self.log(f"[FATAL] {type(e).__name__}: {e}")

    async def _urls_for_codes(self, cli: httpx.AsyncClient, codes: list[str]) -> tuple[list[str], dict]:
        """Ma -> URL detail qua search ?keyWord= (khop CHINH XAC productNo, khong lay bua KQ dau)."""
        urls, inputs = [], {}
        sem = asyncio.Semaphore(self.conc)

        async def one(c: str) -> None:
            async with sem:
                if not await self._gate():
                    return
                d = await self.props(cli, list_url(1, keyword=c))
                await asyncio.sleep(self.delay)
                lst = ((d or {}).get("data") or {}).get("list") or []
                hit = next((i for i in lst if norm_code(i.get("productNo")) == norm_code(c)), None)
                if hit:
                    u = detail_url(hit)
                    urls.append(u)
                    inputs[norm_code(c)] = c
                else:
                    self.log(f"[MISS] search khong ra ma: {c}")

        await asyncio.gather(*(one(c) for c in codes))
        return urls, inputs

    async def run(self, mode: str, target: str = "", excel: str = "", sheet: str = "",
                  col: str = "", limit: int = 0, category: str = "", outdir: str = "out") -> list[dict]:
        key = self.key = f"{mode}_{target or category or excel}_{col}"
        headers = {"User-Agent": UA, "Accept-Language": "en-US,en;q=0.9"}
        async with httpx.AsyncClient(headers=headers, timeout=TIMEOUT, follow_redirects=True,
                                     max_redirects=MAX_REDIRECTS,
                                     limits=httpx.Limits(max_connections=self.conc + 4)) as cli:
            if mode in ("all", "keyword"):
                items = await self._collect_list(cli, category, target if mode == "keyword" else "",
                                                 limit)
                by_url = {detail_url(i): i for i in items}
                await self._crawl_details(cli, list(by_url), key, items=by_url)

            elif mode in ("code", "excel"):
                raw = [norm_code(target)] if mode == "code" else read_excel_col(excel, sheet, col)
                raw = [c for c in dict.fromkeys(raw) if c][:MAX_ITEMS]
                links = [c for c in raw if c.lower().startswith("http")]
                codes = [c for c in raw if c not in links]
                urls, inputs = await self._urls_for_codes(cli, codes)
                await self._crawl_details(cli, links + urls, key, inputs)
                have = {norm_code(r["Item No"]) for r in self.results}
                for c in codes:                       # ma miss van co dong Not Found
                    if c not in have:
                        self.results.append({"Input": c, "Item No": c, "Ten": "",
                                             "Status": "Not Found", "Ghi chu": "", "Video": "",
                                             "_cats": [], "_spec": {}, "Image": "", "URL": ""})

            elif mode == "link":
                await self._crawl_details(
                    cli, [u.strip() for u in target.splitlines() if u.strip().startswith("http")], key)

            elif mode == "images":
                await self._download_images(cli, excel, sheet, col, outdir)

            else:
                self.log(f"[ERR] mode khong hop le: {mode}")
        return self.results

    async def _download_images(self, cli: httpx.AsyncClient, excel: str, sheet: str,
                               col: str, outdir: str) -> None:
        root = pathlib.Path(outdir) if os.path.isabs(outdir) else HERE / outdir
        root.mkdir(parents=True, exist_ok=True)
        keys = read_excel_col(excel, sheet, col)[:MAX_ITEMS]
        self.total = len(keys)
        sem = asyncio.Semaphore(self.conc)

        async def one(k: str) -> None:
            async with sem:
                if not await self._gate():
                    return
                if k.lower().startswith("http"):
                    url = k
                else:
                    urls, _ = await self._urls_for_codes(cli, [k])
                    url = urls[0] if urls else ""
                rec = None
                if url:
                    d = await self.props(cli, url)
                    data = (d or {}).get("data") or {}
                    rec = rec_from_detail(data, url) if data.get("productNo") else None
                self.done += 1
                if not rec or not rec["Image"]:
                    self.log(f"[MISS] {self.done}/{self.total} khong ra anh: {k}")
                    return
                name = re.sub(r'[\\/*?:"<>|]', "_", rec["Item No"] or k)[:60]
                folder = root / name
                folder.mkdir(parents=True, exist_ok=True)
                ok = 0
                for i, iu in enumerate(rec["Image"].split(" | "), 1):
                    try:
                        rr = await cli.get(iu, headers={"Referer": rec["URL"]})
                        rr.raise_for_status()
                        (folder / f"{name}_{i}{pathlib.Path(urlparse(iu).path).suffix or '.jpg'}"
                         ).write_bytes(rr.content)
                        ok += 1
                    except Exception as e:
                        self.log(f"[ERR] tai anh {iu[-45:]}: {type(e).__name__}")
                self.log(f"[OK ] {self.done}/{self.total} {name} -> {ok} anh")

        await asyncio.gather(*(one(k) for k in keys))


def read_excel_col(path: str, sheet: str, col: str) -> list[str]:
    if not path:
        raise ValueError("Chua chon file Excel")
    df = pd.read_excel(path, sheet_name=sheet or 0, dtype=str)
    df.columns = df.columns.str.strip()
    if col not in df.columns:
        # ValueError chu KHONG SystemExit: SystemExit la BaseException, worker GUI bat
        # `except Exception` se KHONG bat duoc -> thread chet cam, GUI treo running=True
        raise ValueError(f"Khong thay cot '{col}'. Cot co san: {list(df.columns)}")
    return [norm_code(v) for v in df[col].dropna().tolist()]


# ================================================================ GUI
def run_gui() -> None:
    import sys
    import tkinter as tk
    from tkinter import messagebox, ttk

    # ha tang GUI chung o root (di nguoc len cay thu muc -> song sot khi doi nhom folder)
    root_dir = next(p for p in (HERE, *HERE.parents) if (p / "gui_common.py").exists())
    sys.path.insert(0, str(root_dir))
    from gui_common import ask_export_path, attach_excel_picker

    root = tk.Tk()
    root.title("TOTAL Crawler — totalbusiness.com")
    root.geometry("1080x680")
    ttk.Style().theme_use("clam")

    q: queue.Queue[str] = queue.Queue()
    stop_evt, pause_evt = threading.Event(), threading.Event()
    state: dict[str, Any] = {"running": False, "results": [], "eng": None}

    left = ttk.Frame(root, padding=10)
    left.pack(side="left", fill="y")
    right = ttk.Frame(root, padding=(0, 10, 10, 10))
    right.pack(side="right", fill="both", expand=True)

    mode = tk.StringVar(value="code")
    ttk.Label(left, text="Pipeline", font=("Segoe UI", 10, "bold")).pack(anchor="w")
    for val, txt in [("code", "1. Theo MA"), ("keyword", "2. Theo TU KHOA"),
                     ("link", "3. Theo LINK (moi dong 1 link)"), ("excel", "4. Theo cot EXCEL"),
                     ("all", "5. Crawl ALL (919 SP)"), ("images", "6. Tai ANH")]:
        ttk.Radiobutton(left, text=txt, value=val, variable=mode).pack(anchor="w")

    ttk.Label(left, text="\nMa / Tu khoa:").pack(anchor="w")
    ent = ttk.Entry(left, width=36)
    ent.pack(anchor="w")
    ttk.Label(left, text="Link (moi dong 1 link):").pack(anchor="w")
    box = tk.Text(left, width=36, height=5)
    box.pack(anchor="w")

    cat = tk.StringVar()
    ttk.Label(left, text="Category (rong = tat ca):").pack(anchor="w")
    ttk.Combobox(left, textvariable=cat, width=33, values=[
        "", "power-tools", "power-tools-accessories", "generators", "welding-machines",
        "small-construction-equipment", "air-tools", "water-pumps", "hand-tools",
        "measuring-tools", "safety-products", "garden-tools"]).pack(anchor="w")

    xls, sheet, col = tk.StringVar(), tk.StringVar(), tk.StringVar()
    outdir = tk.StringVar(value="out")
    attach_excel_picker(left, xls, sheet, col, outdir)

    conc, delay, limit = tk.IntVar(value=6), tk.DoubleVar(value=0.1), tk.IntVar(value=0)
    for lbl, var in [("Concurrency:", conc), ("Delay (s):", delay), ("Limit (0=all):", limit)]:
        f3 = ttk.Frame(left)
        f3.pack(anchor="w", pady=1)
        ttk.Label(f3, text=lbl, width=13).pack(side="left")
        ttk.Entry(f3, textvariable=var, width=8).pack(side="left")

    console = tk.Text(right, bg="#12161c", fg="#d6e2ee", insertbackground="#fff", wrap="none")
    console.pack(fill="both", expand=True)
    pb = ttk.Progressbar(right, mode="determinate")
    pb.pack(fill="x", pady=4)
    stat = ttk.Label(right, text="0/0")
    stat.pack(anchor="w")

    def drain() -> None:
        while not q.empty():
            console.insert("end", q.get() + "\n")
            console.see("end")
        eng: Engine | None = state["eng"]
        if eng and eng.total:
            pb["maximum"], pb["value"] = eng.total, eng.done
            stat.config(text=f"{eng.done}/{eng.total} | {len(eng.results)} record")
        root.after(100, drain)

    def worker() -> None:
        m = mode.get()
        target = box.get("1.0", "end").strip() if m == "link" else ent.get().strip()
        eng = Engine(conc.get(), delay.get(), q.put, stop_evt, pause_evt)
        state["eng"] = eng
        try:
            asyncio.run(eng.run(m, target, xls.get(), sheet.get(), col.get(),
                                limit.get(), cat.get(), outdir.get()))
            eng.cleanup()
        except BaseException as e:     # ke ca SystemExit -> khong de thread chet cam
            q.put(f"[FATAL] {type(e).__name__}: {e}")
        finally:                       # LUON nha co, khong khoa START vinh vien
            state["results"] = eng.results
            state["running"] = False
            q.put(f"[SYS] XONG — {len(eng.results)} record. Bam EXPORT de xuat Excel.")

    def start() -> None:
        if state["running"]:
            return
        stop_evt.clear()
        pause_evt.clear()
        state["running"] = True
        console.delete("1.0", "end")
        threading.Thread(target=worker, daemon=True).start()

    def export() -> None:
        if not state["results"]:
            messagebox.showinfo("Export", "Chua co du lieu")
            return
        p = ask_export_path(xls.get(), HERE, "totalbusiness.xlsx")   # mac dinh canh file Excel dau vao
        if p:
            export_excel(state["results"], pathlib.Path(p))
            messagebox.showinfo("Export", f"Da luu:\n{p}")

    fb = ttk.Frame(left)
    fb.pack(anchor="w", pady=8)
    ttk.Button(fb, text="START", command=start).pack(side="left")
    ttk.Button(fb, text="Pause", command=lambda: (pause_evt.clear() if pause_evt.is_set()
                                                  else pause_evt.set())).pack(side="left")
    ttk.Button(fb, text="Stop", command=stop_evt.set).pack(side="left")
    ttk.Button(fb, text="EXPORT", command=export).pack(side="left")

    def on_close() -> None:
        if state["running"] and not messagebox.askyesno("Thoat", "Dang crawl. Thoat va tat luon?"):
            return
        stop_evt.set()
        try:
            root.destroy()
        finally:
            os._exit(0)

    root.protocol("WM_DELETE_WINDOW", on_close)
    drain()
    root.mainloop()


# ================================================================ CLI
def main() -> None:
    ap = argparse.ArgumentParser(description="totalbusiness.com (TOTAL Tools) crawler")
    ap.add_argument("--mode", choices=["code", "keyword", "link", "excel", "all", "images"])
    ap.add_argument("--code")
    ap.add_argument("--keyword")
    ap.add_argument("--link")
    ap.add_argument("--category", default="", help="slug category, vd hand-tools")
    ap.add_argument("--excel")
    ap.add_argument("--sheet", default="")
    ap.add_argument("--col", default="Part")
    ap.add_argument("--outdir", default="out")
    ap.add_argument("--out", default="totalbusiness.xlsx")
    ap.add_argument("--conc", type=int, default=6)
    ap.add_argument("--delay", type=float, default=0.1)
    ap.add_argument("--limit", type=int, default=0)
    a = ap.parse_args()

    if not a.mode:
        run_gui()
        return

    eng = Engine(a.conc, a.delay)
    t0 = time.time()
    try:
        asyncio.run(eng.run(a.mode, a.code or a.keyword or a.link or "", a.excel or "",
                            a.sheet, a.col, a.limit, a.category, a.outdir))
    finally:
        if eng.results and a.mode != "images":
            p = export_excel(eng.results, pathlib.Path(a.out))
            print(f"[SYS] {len(eng.results)} record -> {p} ({time.time() - t0:.1f}s)")
        eng.cleanup()


if __name__ == "__main__":
    main()
